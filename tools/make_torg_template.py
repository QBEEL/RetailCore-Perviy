"""Делает шаблон акта ТОРГ-2 из заполненного образца.

Образец — реальный акт склада: в нём и текст формы, и данные конкретной
поставки, а жёлтым отмечено то, что склад вписывает руками. Шаблону нужен
только текст формы. Скрипт очищает все поля, которые приложение заполняет
само, убирает жёлтую заливку и сохраняет книгу в `app/ui/assets`.

Запуск:  .venv/Scripts/python.exe tools/make_torg_template.py "Для Торг-2/образец.xlsx"

Нужен один раз и при смене бланка. Список очищаемых ячеек — единственное, что
связывает его с `app/core/claims/torg.py`; если в бланке появится новое поле,
добавьте его в оба места, иначе данные образца попадут в чужие акты (этот случай
ловит тест `test_в_шаблоне_нет_данных_образца`).
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "app" / "ui" / "assets" / "torg2_template.xlsx"

# Первая страница: шапка, документы, даты.
FIRST = [
    "A6", "A9", "CI11", "CI12", "CP12", "CW12", "CA16", "AY17", "BK17", "CK18",
    "CD20", "CI20", "CS20", "Z23", "BL24", "BS24", "CS24", "AL25", "U33", "N37",
    "AT42", "BN42", "BU42", "CG42", "U43", "AO43", "AV43", "BH43", "T48", "CP48",
    "AQ60", "BP60", "CT60",
]
# Вторая страница: дата вскрытия и строки товаров. Строки 40–42 — товары и итог.
SECOND = ["EB20", "EI20", "FA20"]
SECOND_COLUMNS = ["A", "AB", "AL", "BA", "BK", "BS", "BY", "CD", "CQ", "DD", "DK",
                  "DR", "DV", "DY", "EF", "EM", "ET", "FO", "FW"]
SECOND_ROWS = (40, 41, 42)
# Третья страница: описание, заключение, комиссия.
THIRD = ["AW20", "A21", "A22", "A23", "Y24", "A25", "A26", "AC30", "CA30",
         "AC33", "CA33", "AC35", "CA35", "AC37", "CA37"]


def main(source: str) -> None:
    book = openpyxl.load_workbook(source)
    for address in FIRST:
        book["стр1"][address].value = None
    sheet = book["стр2"]
    for address in SECOND:
        sheet[address].value = None
    for row in SECOND_ROWS:
        for column in SECOND_COLUMNS:
            sheet[f"{column}{row}"].value = None
    for address in THIRD:
        book["стр3"][address].value = None

    nothing = PatternFill(fill_type=None)
    for ws in book:
        for row in ws.iter_rows():
            for cell in row:
                fill = cell.fill
                if fill and fill.fill_type and str(fill.fgColor.rgb) == "FFFFFF00":
                    cell.fill = nothing
    # Автор и «кем изменён» — люди, правившие образец. В репозиторий они не
    # нужны, а в каждый акт переходили бы вместе с шаблоном.
    book.properties.creator = "RetailCore"
    book.properties.lastModifiedBy = None
    book.properties.lastPrinted = None
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    book.save(TARGET)
    print(f"Шаблон сохранён: {TARGET}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
