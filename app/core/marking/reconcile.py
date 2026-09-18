"""Сверка приехавшего товара с кодами из УПД.

Работа выглядит так: кладовщик берёт вещь, пикает сканером наклейку и кладёт в
сторону. Программа должна ответить мгновенно и одним словом — эта вещь в
документе или нет. Всё остальное (сколько осталось, по каким строкам недобор,
что лишнее) считается заодно, но человек смотрит на товар, а не в экран, и
узнаёт о расхождении в тот момент, когда вещь ещё у него в руках.

Сравнивать «как есть» нельзя. В документе лежит код идентификации — товар и
серийный номер, 31 знак. Сканер отдаёт то, что напечатано на этикетке: тот же
код плюс криптохвост, а разделители полей теряются по дороге — их срезает и
драйвер сканера, и поле ввода. Поэтому сравнение идёт по разобранному коду:
пара «GTIN и серийный номер» у одной и той же вещи совпадёт всегда, как бы её
ни сняли. Разбором занимается `codes`, здесь только сведение двух списков.

Отмена последнего скана — не украшение. Сканер стреляет по соседней коробке
чаще, чем хотелось бы, и без отмены единственный выход из ошибки — начать
сверку заново.
"""
from __future__ import annotations

import os
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import codes as codes_module
from .upd import Document, Line, Mark

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
TROUBLE_FILL = PatternFill("solid", fgColor="FDE8E8")
DONE_FILL = PatternFill("solid", fgColor="EAF7EE")


class Verdict(Enum):
    """Что программа отвечает на один скан."""

    MATCHED = ("matched", "Есть в документе")
    REPEAT = ("repeat", "Уже сверяли")
    UNKNOWN = ("unknown", "Нет в документе")
    BROKEN = ("broken", "Код не разобран")

    def __init__(self, value: str, title: str) -> None:
        self._value_ = value
        self.title = title

    @property
    def good(self) -> bool:
        return self is Verdict.MATCHED


@dataclass(slots=True)
class Scan:
    """Один скан вместе с тем, чем он оказался."""

    raw: str
    key: str
    verdict: Verdict
    line: Line | None = None
    note: str = ""
    at: datetime = field(default_factory=datetime.now)

    @property
    def title(self) -> str:
        """Код в читаемом виде: разобранный — товаром и номером, иначе как есть."""
        code = codes_module.parse(self.raw)
        if code.gtin and code.serial:
            return f"{code.gtin} · {code.serial}"
        return self.raw[:48]


@dataclass(slots=True)
class LineProgress:
    """Сколько по строке ожидалось и сколько уже сверено."""

    line: Line
    scanned: int = 0

    @property
    def expected(self) -> int:
        return len(self.line.marks)

    @property
    def left(self) -> int:
        return max(0, self.expected - self.scanned)

    @property
    def done(self) -> bool:
        return self.scanned >= self.expected

    @property
    def state(self) -> str:
        if self.done:
            return "Сверено"
        if self.scanned:
            return f"Не хватает {self.left}"
        return "Не начато"


def key_of(text: str) -> str:
    """По чему один и тот же экземпляр опознаётся в документе и на этикетке.

    Для товара это GTIN и серийный номер — криптохвост у одной и той же вещи
    разный в документе и на наклейке, и сравнение по целой строке развалилось бы
    на первом же коде. Для упаковки — её собственный номер (SSCC). Всё
    остальное сравнивается как есть: разобрать не вышло, но человеку показать
    что-то надо.
    """
    code = codes_module.parse(text)
    if code.gtin and code.serial:
        return f"01{code.gtin}21{code.serial}"
    if sscc := code.fields.get("00"):
        return f"00{sscc}"
    return codes_module.normalize(text)


class Reconciliation:
    """Ход сверки одного документа."""

    def __init__(self, document: Document) -> None:
        self.document = document
        self.started_at = datetime.now()
        self.expected: dict[str, tuple[Line, Mark]] = {}
        # Коды, записанные в документе дважды. Это ошибка поставщика, и заметна
        # она до сканирования: сверить такую строку до конца невозможно в
        # принципе, потому что второй такой вещи в коробке нет.
        self.repeated_in_document: list[str] = []
        for line, mark in document.marks:
            key = key_of(mark.value)
            if key in self.expected:
                self.repeated_in_document.append(mark.value)
                continue
            self.expected[key] = (line, mark)
        self.seen: dict[str, Scan] = {}
        self.history: list[Scan] = []

    # --- сканирование ---------------------------------------------------------

    def scan(self, raw: str) -> Scan:
        """Принимает один код и говорит, чем он оказался."""
        text = (raw or "").strip()
        key = key_of(text)
        if not text:
            return Scan(raw=text, key=key, verdict=Verdict.BROKEN,
                        note="пустая строка")
        if key in self.expected:
            line, _ = self.expected[key]
            if key in self.seen:
                scan = Scan(raw=text, key=key, verdict=Verdict.REPEAT, line=line,
                            note="этот экземпляр уже сверяли")
            else:
                scan = Scan(raw=text, key=key, verdict=Verdict.MATCHED, line=line)
                self.seen[key] = scan
            self.history.append(scan)
            return scan
        code = codes_module.parse(text)
        if not code.valid:
            reason = code.problems[0] if code.problems else "код не разобран"
            scan = Scan(raw=text, key=key, verdict=Verdict.BROKEN, note=reason)
        else:
            scan = Scan(raw=text, key=key, verdict=Verdict.UNKNOWN,
                        note=self._why_unknown(code))
        self.history.append(scan)
        return scan

    def _why_unknown(self, code: codes_module.Code) -> str:
        """Почему кода нет в документе — насколько это видно отсюда.

        Разница между «привезли не тот товар» и «привезли лишнюю единицу того же»
        видна по GTIN, и для разговора с поставщиком это разные разговоры.
        """
        for line in self.document.lines:
            if line.gtin and code.gtin.endswith(line.gtin.lstrip("0")):
                return f"товар из строки {line.number}, но этот экземпляр в документе не указан"
        return "такого товара в документе нет"

    def undo(self) -> Scan | None:
        """Отменяет последний скан. Возвращает отменённое или ничего."""
        if not self.history:
            return None
        scan = self.history.pop()
        if scan.verdict is Verdict.MATCHED:
            self.seen.pop(scan.key, None)
        return scan

    def reset(self) -> None:
        """Начинает сверку того же документа заново."""
        self.seen.clear()
        self.history.clear()
        self.started_at = datetime.now()

    # --- что получилось --------------------------------------------------------

    @property
    def progress(self) -> list[LineProgress]:
        counted: dict[str, int] = {}
        for key in self.seen:
            line, _ = self.expected[key]
            counted[line.number] = counted.get(line.number, 0) + 1
        return [LineProgress(line=line, scanned=counted.get(line.number, 0))
                for line in self.document.marked_lines]

    @property
    def total(self) -> int:
        return len(self.expected)

    @property
    def done(self) -> int:
        return len(self.seen)

    @property
    def left(self) -> int:
        return self.total - self.done

    @property
    def extra(self) -> list[Scan]:
        """Коды, которых в документе нет. Повторы сюда не попадают."""
        return [scan for scan in self.history
                if scan.verdict in (Verdict.UNKNOWN, Verdict.BROKEN)]

    @property
    def repeats(self) -> int:
        return sum(1 for scan in self.history if scan.verdict is Verdict.REPEAT)

    @property
    def missing(self) -> list[tuple[Line, Mark]]:
        """Коды документа, которых на товаре не нашлось."""
        return [pair for key, pair in self.expected.items() if key not in self.seen]

    @property
    def complete(self) -> bool:
        """Сверка сошлась: всё из документа найдено и ничего лишнего."""
        return not self.left and not self.extra

    @property
    def summary(self) -> str:
        if not self.history:
            return (f"К сверке {self.total} кодов по {len(self.document.marked_lines)} "
                    "строкам. Сканируйте первую вещь.")
        parts = [f"сверено {self.done} из {self.total}"]
        if self.left:
            parts.append(f"осталось {self.left}")
        if extra := len(self.extra):
            parts.append(f"лишних {extra}")
        if self.repeats:
            parts.append(f"повторов {self.repeats}")
        if self.complete:
            parts.append("расхождений нет")
        return " · ".join(parts)


# --- акт сверки -------------------------------------------------------------------

def save_report(session: Reconciliation, destination: str) -> str:
    """Записывает результат сверки в xlsx. Возвращает путь к готовому файлу.

    Файл нужен для разговора с поставщиком и для приложения к претензии,
    поэтому в нём есть и то, чего не хватило, и то, что приехало сверх
    документа, — а не только итоговое «не сошлось».
    """
    if not destination:
        raise ValueError("Не указан файл для сохранения")
    book = Workbook()
    _lines_sheet(book.active, session)
    _codes_sheet(book.create_sheet("Коды"), session)
    if directory := os.path.dirname(destination):
        os.makedirs(directory, exist_ok=True)
    book.save(destination)
    return destination


def default_name(session: Reconciliation, folder: str) -> str:
    """Имя по умолчанию: «Сверка УПД 3349 от 09.09.2026.xlsx»."""
    number = session.document.number or "без номера"
    date = f" от {session.document.date}" if session.document.date else ""
    name = _safe(f"Сверка УПД {number}{date}")
    return os.path.join(folder, f"{name}.xlsx")


def _safe(name: str) -> str:
    cleaned = "".join(" " if char in '\\/:*?"<>|' else char for char in name)
    return " ".join(cleaned.split()) or "Сверка"


def _lines_sheet(sheet, session: Reconciliation) -> None:
    sheet.title = "Сверка"
    document = session.document
    head = [
        ("Акт сверки кодов маркировки", ""),
        ("Документ", document.title),
        ("Поставщик", f"{document.seller} ИНН {document.seller_inn}".strip()),
        ("Покупатель", f"{document.buyer} ИНН {document.buyer_inn}".strip()),
        ("Сверка начата", f"{session.started_at:%d.%m.%Y %H:%M}"),
        ("Итог", session.summary),
    ]
    for row, (label, value) in enumerate(head, start=1):
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, size=12 if row == 1 else 11)
        sheet.cell(row=row, column=2, value=value)
    row = len(head) + 2

    columns = ("№", "Товар", "GTIN", "Ожидалось", "Сверено", "Не найдено", "Состояние")
    for column, title in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1

    for item in session.progress:
        values = (item.line.number, item.line.title, item.line.gtin, item.expected,
                  item.scanned, item.left, item.state)
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.fill = DONE_FILL if item.done else TROUBLE_FILL
        row += 1

    sheet.freeze_panes = sheet.cell(row=len(head) + 3, column=1)
    _widths(sheet, (6, 46, 16, 12, 10, 12, 18))


def _codes_sheet(sheet, session: Reconciliation) -> None:
    columns = ("Строка", "Товар", "Код маркировки", "Состояние", "Замечание")
    for column, title in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    row = 2

    for line, mark in session.document.marks:
        found = key_of(mark.value) in session.seen
        values = (line.number, line.title, mark.value,
                  "Сверен" if found else "Не найден на товаре",
                  "" if found else "код есть в документе, вещь не отсканирована")
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.fill = DONE_FILL if found else TROUBLE_FILL
        row += 1

    for scan in session.extra:
        values = ("", "", scan.raw, "Лишний — нет в документе", scan.note)
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.fill = TROUBLE_FILL
        row += 1

    sheet.freeze_panes = "A2"
    _widths(sheet, (8, 40, 44, 24, 42))


def _widths(sheet, widths: tuple[int, ...]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
