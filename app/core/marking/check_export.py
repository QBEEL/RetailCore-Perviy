"""Выгрузка результата проверки кодов в Excel.

Нужна для разговора с поставщиком: коды без разрешительного документа (РД)
исправляет только тот, кто вводил их в оборот, и ему надо передать список
кодов, а не снимок экрана. Выгружается то, что видно в таблице, — с учётом
отбора, — чтобы поставщику не пришлось искать свои коды среди исправных.
"""
from __future__ import annotations

import os
from datetime import datetime
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .. import workbook as workbook_module
from .models import CodeInfo
from .reconcile import HEADER_FILL, TROUBLE_FILL

COLUMNS = ("Код маркировки", "GTIN", "Товар", "Состояние", "Владелец",
           "ИНН владельца", "РД", "Разрешительные документы", "Замечания")
WIDTHS = (36, 16, 46, 18, 30, 15, 18, 60, 50)


def save(items: Sequence[CodeInfo], destination: str) -> str:
    """Записывает коды в xlsx. Возвращает путь к готовому файлу."""
    if not destination:
        raise ValueError("Не указан файл для сохранения")
    if not items:
        raise ValueError("Нечего выгружать: в таблице нет кодов")
    book = Workbook()
    sheet = book.active
    sheet.title = "Проверка кодов"
    for column, title in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row, item in enumerate(items, start=2):
        remarks = item.remarks
        values = (item.code, item.gtin, item.product_name, item.state.title,
                  item.owner_name, item.owner_inn, item.permit_state.title,
                  "\n".join(doc.title for doc in item.permits or []),
                  "; ".join(remarks))
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            # Код, GTIN и ИНН — текст: числом Excel съел бы ведущие нули GTIN,
            # а код маркировки бывает похож на формулу.
            cell.number_format = "@"
            workbook_module.as_text(cell)
            cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 8, 9))
            if remarks:
                cell.fill = TROUBLE_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    if directory := os.path.dirname(destination):
        os.makedirs(directory, exist_ok=True)
    book.save(destination)
    return destination


def default_name(folder: str, *, only_permits: bool = False) -> str:
    """Имя по умолчанию: «Коды без РД 20.10.2026 16-40.xlsx»."""
    title = "Коды без РД" if only_permits else "Проверка кодов"
    return os.path.join(folder, f"{title} {datetime.now():%d.%m.%Y %H-%M}.xlsx")
