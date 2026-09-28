"""Оформление листов отчёта в Excel — одно на все разборы выгрузок 1С.

Ведомость по складам и аналитика по чекам сохраняются одинаково: лист с
выводами, затем по листу на вопрос, шапка тёмная, строки через одну. Держать
это в каждом отчёте своим — значит через полгода получить два разных вида у
файлов, которые бухгалтер кладёт рядом.

Форматы колонок задаются строкой по буквам, по одной на колонку:
`t` — текст, `n` — штуки, `m` — деньги, `%` — доля.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FONT = "Calibri"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="DCE3F0")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
STRIPE = PatternFill("solid", fgColor="F4F6FB")

HAIR = Side(style="hair", color="B4BFD4")
GRID = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)

QTY = "#,##0"
MONEY = "#,##0.00"
SHARE = "0.0%"

MIN_WIDTH, MAX_WIDTH = 9, 56

_NUMBERS = {"n": QTY, "m": MONEY, "%": SHARE}


def title(sheet: Worksheet, row: int, text: str, width: int) -> int:
    cell = sheet.cell(row, 1, text)
    cell.font = Font(FONT, size=13, bold=True, color="1F3864")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    return row + 1


def note(sheet: Worksheet, row: int, text: str, width: int) -> int:
    cell = sheet.cell(row, 1, text)
    cell.font = Font(FONT, size=9, italic=True, color="5A6A85")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    sheet.row_dimensions[row].height = 26
    return row + 1


def head(sheet: Worksheet, row: int, titles: list[str]) -> int:
    for column, text in enumerate(titles, 1):
        cell = sheet.cell(row, column, text)
        cell.font = Font(FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    sheet.row_dimensions[row].height = 30
    sheet.freeze_panes = sheet.cell(row + 1, 1)
    return row + 1


def row(sheet: Worksheet, number: int, values: list, *, formats: str = "",
        stripe: bool = False, warn: bool = False) -> None:
    for column, value in enumerate(values, 1):
        kind = formats[column - 1] if column - 1 < len(formats) else "t"
        if kind == "m" and isinstance(value, float):
            # Сумма сотен строк несёт «хвост» вида 5933279.670000216 — в
            # ячейке его не видно, но он всплывает при копировании значения.
            value = round(value, 2)
        cell = sheet.cell(number, column, value)
        cell.font = Font(FONT, size=10)
        cell.border = GRID
        if kind in _NUMBERS:
            cell.number_format = _NUMBERS[kind]
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        if warn:
            cell.fill = WARN_FILL
        elif stripe:
            cell.fill = STRIPE


def widths(sheet: Worksheet, head_row: int) -> None:
    for column in range(1, sheet.max_column + 1):
        longest = 0
        for number in range(head_row, sheet.max_row + 1):
            value = sheet.cell(number, column).value
            if value is None:
                continue
            longest = max(longest, len(str(value)) if isinstance(value, str)
                          else len(f"{value:,.0f}"))
        sheet.column_dimensions[get_column_letter(column)].width = min(
            MAX_WIDTH, max(MIN_WIDTH, longest + 2))
