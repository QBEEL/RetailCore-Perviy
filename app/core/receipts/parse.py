"""Разбор выгрузки «Аналитика по чекам» из 1С.

Над таблицей 1С печатает параметры отчёта: период, подразделения и отбор. Из
них нужен прежде всего период — без него «выручка» не отвечает на вопрос «за
что», а средняя выручка в день не считается вовсе.

Шапка ищется по подписям, а не по буквам колонок: отчёт в 1С настраивается,
и колонку «Бренд» или «Продавец» могут убрать или переставить. Обязательны
только документ, дата, магазин, товар, количество и сумма — на них стоят все
выводы. Остальные колонки желательны: без себестоимости нет прибыли, без
продавца — среза по продавцам, и такие выводы не показываются, а не
считаются нулём.

Строка «Итого» в конце таблицы — сверка: если сумма строк с ней не сходится,
шапку прочитали неверно, и сказать об этом важнее, чем показать отчёт.
"""
from __future__ import annotations

import os
import re
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from .models import Line, Receipts

DOCUMENT, KIND, DAY, STORE, GUEST, ITEM = "document", "kind", "day", "store", "guest", "item"
CATEGORY, BRAND, PRICE, PROMO, AUTO_PERCENT = "category", "brand", "price", "promo", "auto_percent"
SELLER, QUANTITY, COST, AUTO_DISCOUNT = "seller", "quantity", "cost", "auto_discount"
BONUS_DISCOUNT, AMOUNT, BONUS_ACCRUED = "bonus_discount", "amount", "bonus_accrued"

# Подписи колонок после `_key`: без регистра, знаков препинания и лишних
# пробелов, «ё» как «е».
LABELS = {
    "документ": DOCUMENT,
    "чек": DOCUMENT,
    "вид документа": KIND,
    "дата": DAY,
    "магазин": STORE,
    "подразделение": STORE,
    "гость": GUEST,
    "покупатель": GUEST,
    "клиент": GUEST,
    "номенклатура с характеристикой": ITEM,
    "номенклатура": ITEM,
    "номенклатура характеристика": ITEM,
    "товар": ITEM,
    "категория товара": CATEGORY,
    "категория": CATEGORY,
    "бренд": BRAND,
    "марка": BRAND,
    "цена": PRICE,
    "акция": PROMO,
    "% авт": AUTO_PERCENT,
    "процент автоматической скидки": AUTO_PERCENT,
    "продавец": SELLER,
    "количество": QUANTITY,
    "себестоимость": COST,
    "сумма авт": AUTO_DISCOUNT,
    "сумма автоматической скидки": AUTO_DISCOUNT,
    "скидка бонусами": BONUS_DISCOUNT,
    "оплачено бонусами": BONUS_DISCOUNT,
    "сумма": AMOUNT,
    "сумма продажи": AMOUNT,
    "начислено бонусов": BONUS_ACCRUED,
}

REQUIRED = (DOCUMENT, DAY, STORE, ITEM, QUANTITY, AMOUNT)

TITLES = {
    DOCUMENT: "Документ", KIND: "Вид документа", DAY: "Дата", STORE: "Магазин",
    GUEST: "Гость", ITEM: "Номенклатура", CATEGORY: "Категория товара",
    BRAND: "Бренд", PRICE: "Цена", PROMO: "Акция", AUTO_PERCENT: "% авт.",
    SELLER: "Продавец", QUANTITY: "Количество", COST: "Себестоимость",
    AUTO_DISCOUNT: "Сумма авт.", BONUS_DISCOUNT: "Скидка бонусами",
    AMOUNT: "Сумма", BONUS_ACCRUED: "Начислено бонусов",
}

# Колонки, которые сверяются со строкой «Итого».
CHECKED = (QUANTITY, COST, AUTO_DISCOUNT, BONUS_DISCOUNT, AMOUNT, BONUS_ACCRUED)

# Докуда искать шапку: над ней название отчёта и параметры, строк немного, но
# их число меняется вместе с настройками выгрузки.
SEARCH_DEPTH = 40

_DATE = r"(\d{1,2}\.\d{1,2}\.\d{4})"
PERIOD = re.compile(
    rf"период\s*:?\s*(?:с\s*)?{_DATE}\s*(?:-|—|–|по)\s*{_DATE}", re.IGNORECASE)
UNITS = re.compile(r"подразделени[ея]\s*:\s*(.+)", re.IGNORECASE)
SELECTION = re.compile(r"^отбор\s*:?$", re.IGNORECASE)
# «Чек ККМ IP00-071331 от 01.08.2026 9:34:26» — время есть только здесь.
MOMENT = re.compile(r"от\s+(\d{1,2}\.\d{1,2}\.\d{4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?")


class ParseError(Exception):
    """Файл не похож на аналитику по чекам. Текст пригоден для показа."""


def _text(value: Any) -> str:
    return " ".join(str(value).split()).strip() if value is not None else ""


def _key(value: Any) -> str:
    text = _text(value).lower().replace("ё", "е")
    for mark in ".,;:":
        text = text.replace(mark, " ")
    return " ".join(text.split())


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10], pattern).date()
        except ValueError:
            continue
    return None


def _moment(document: str, day: date | None) -> datetime | None:
    found = MOMENT.search(document)
    if not found:
        return None
    moment_day = _date(found.group(1)) or day
    if moment_day is None:
        return None
    return datetime(moment_day.year, moment_day.month, moment_day.day,
                    int(found.group(2)), int(found.group(3)), int(found.group(4) or 0))


def read(path: str) -> Receipts:
    """Читает одну выгрузку. Ошибка объясняет, чего в файле не нашлось."""
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 — openpyxl бросает что угодно
        raise ParseError(f"Не удалось открыть {os.path.basename(path)}: {error}") from None
    try:
        rows = [list(row) for row in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()
    return parse_rows(rows, source=os.path.basename(path))


def parse_rows(rows: list[list[Any]], source: str = "") -> Receipts:
    """Разбор уже прочитанных строк листа — отдельно, чтобы проверять без файла."""
    result = Receipts(sources=[source] if source else [])
    head_index, columns = _find_head(rows, source)
    _read_parameters(rows[:head_index], result)
    result.missing = [name for name in TITLES if name not in columns]

    total: list[Any] | None = None
    for raw in rows[head_index + 1:]:
        if raw and _key(raw[0]) == "итого":
            total = raw
            break
        if (line := _line(raw, columns, source)) is not None:
            result.lines.append(line)

    if not result.lines:
        raise ParseError(f"{source or 'Файл'}: шапка найдена, но строк чеков под ней нет")
    if total is not None:
        result.balanced = _balanced(result.lines, total, columns, result)
    _fill_period(result)
    return result


def _find_head(rows: list[list[Any]], source: str) -> tuple[int, dict[str, int]]:
    """Строка шапки — та, где узнано больше всего подписей."""
    best: tuple[int, dict[str, int]] = (-1, {})
    for index, raw in enumerate(rows[:SEARCH_DEPTH]):
        columns: dict[str, int] = {}
        for position, value in enumerate(raw):
            name = LABELS.get(_key(value))
            # Первая подходящая колонка побеждает: вторая «Сумма» правее —
            # это уже нестандартная настройка, и стоит там не то.
            if name and name not in columns:
                columns[name] = position
        if len(columns) > len(best[1]):
            best = (index, columns)
    index, columns = best
    lacking = [TITLES[name] for name in REQUIRED if name not in columns]
    if index < 0 or lacking:
        found = ", ".join(TITLES[name] for name in columns) or "ничего"
        raise ParseError(
            f"{source or 'Файл'} не похож на «Аналитику по чекам»: в шапке "
            f"нашлось — {found}; не хватает — {', '.join(lacking)}.")
    return index, columns


def _read_parameters(rows: list[list[Any]], result: Receipts) -> None:
    """Период, подразделения и отбор из строк над шапкой."""
    for raw in rows:
        cells = [_text(value) for value in raw if _text(value)]
        for position, text in enumerate(cells):
            if result.start is None and (found := PERIOD.search(text)):
                result.start, result.end = _date(found.group(1)), _date(found.group(2))
            if not result.units and (found := UNITS.search(text)):
                result.units = found.group(1).strip()
            if (not result.selection and SELECTION.match(text)
                    and position + 1 < len(cells)):
                result.selection = cells[position + 1]


def _line(raw: list[Any], columns: dict[str, int], source: str) -> Line | None:
    def get(name: str) -> Any:
        position = columns.get(name)
        return raw[position] if position is not None and position < len(raw) else None

    document = _text(get(DOCUMENT))
    item = _text(get(ITEM))
    # Пустые строки и разделители: без документа и товара это не чек.
    if not document and not item:
        return None
    day = _date(get(DAY))
    return Line(
        document=document,
        kind=_text(get(KIND)),
        day=day,
        moment=_moment(document, day),
        store=_text(get(STORE)),
        guest=_text(get(GUEST)),
        item=item,
        category=_text(get(CATEGORY)),
        brand=_text(get(BRAND)),
        price=_number(get(PRICE)),
        promo=_text(get(PROMO)),
        auto_percent=_number(get(AUTO_PERCENT)),
        seller=_text(get(SELLER)),
        quantity=_number(get(QUANTITY)),
        cost=_number(get(COST)),
        auto_discount=_number(get(AUTO_DISCOUNT)),
        bonus_discount=_number(get(BONUS_DISCOUNT)),
        amount=_number(get(AMOUNT)),
        bonus_accrued=_number(get(BONUS_ACCRUED)),
        source=source,
    )


def _balanced(lines: list[Line], total: list[Any], columns: dict[str, int],
              result: Receipts) -> bool:
    """Сверка со строкой «Итого».

    Допуск — рубль или стотысячная доля итога, что больше: итог себестоимости
    1С округляет иначе, чем строки, и на трёх миллионах расходится на пару
    рублей. Больше этого — уже не округление.
    """
    fine = True
    for name in CHECKED:
        position = columns.get(name)
        if position is None or position >= len(total) or total[position] in (None, ""):
            continue
        expected = _number(total[position])
        actual = sum(getattr(line, name) for line in lines)
        if abs(expected - actual) > max(1.0, abs(expected) * 1e-5):
            fine = False
            result.warnings.append(
                f"«{TITLES[name]}» не сходится с итогом файла: по строкам "
                + f"{actual:,.2f}, в итоге {expected:,.2f}".replace(",", " "))
    return fine


def _fill_period(result: Receipts) -> None:
    """Период из шапки; нет его — по датам строк, с предупреждением."""
    days = [line.day for line in result.lines if line.day]
    if result.start is None and days:
        result.start, result.end = min(days), max(days)
        result.warnings.append("Период в параметрах отчёта не найден — взят по датам чеков")
    if result.start and result.end and days and (
            min(days) < result.start or max(days) > result.end):
        result.warnings.append(f"Есть чеки вне периода отчёта {result.period_title}")
