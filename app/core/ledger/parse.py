"""Разбор выгрузки «Ведомость по товарам на складах» из 1С.

Шапка ищется по подписям, а не по номерам колонок. Состав складов меняется от
выгрузки к выгрузке — магазин открыли, склад переименовали, — и отчёт, в котором
колонки посчитаны от буквы «H», сломается на первой же такой выгрузке молча:
числа встанут не под тем складом, и заметить это будет нечем.

Отчёт в 1С у каждого подразделения настроен по-своему, поэтому видов у него
несколько, и разбор узнаёт каждый:

- склады в колонках — на склад блок «Начальный остаток · Приход · Расход ·
  Конечный остаток», иногда вместе с блоком сумм;
- склад колонкой — плоская таблица, где склад указан в каждой строке;
- склады группами строк — строка склада, под ней его товары;
- один склад — шапка без складов, название берётся из отбора или имени файла.

Подписи в разных настройках разные — «Номенклатура» и «Номенклатура,
Характеристика», «Начальный остаток» и «Остаток на начало», — и сравниваются
по словарю синонимов.

Строку товара от строки группы отличает единица измерения. У групп её нет, а
числа в них — итоги вложенных товаров: сложить их вместе с товарами значило бы
посчитать всё дважды. Если колонки единицы нет, группу выдаёт уровень
группировки строк Excel.

Если файл не узнан, ошибка говорит, что именно в шапке нашлось и чего не
хватило: по такому сообщению видно, чем выгрузка отличается от ожидаемой.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .models import FIELDS as MOVE_FIELDS
from .models import Item, Ledger, Movement, Store

# Подписи полей движения. Сравниваются после `_key`: без регистра, точек,
# запятых и лишних пробелов, «ё» как «е».
FIELDS = {
    "начальный остаток": "opening",
    "остаток на начало": "opening",
    "остаток на начало периода": "opening",
    "нач остаток": "opening",
    "приход": "incoming",
    "поступление": "incoming",
    "расход": "outgoing",
    "выбытие": "outgoing",
    "конечный остаток": "closing",
    "остаток на конец": "closing",
    "остаток на конец периода": "closing",
    "кон остаток": "closing",
}

# Без расхода и конечного остатка разбирать нечего: на них стоят все выводы.
# Начальный остаток и приход желательны — без них нет доли «ушло от запаса».
REQUIRED = ("outgoing", "closing")

FIELD_TITLES = {
    "opening": "Начальный остаток",
    "incoming": "Приход",
    "outgoing": "Расход",
    "closing": "Конечный остаток",
}

# Подписи колонок товара. Точное совпадение после `_key`.
ARTICLE, CODE, NAME, UNIT, STORE = "article", "code", "name", "unit", "store"
LABELS = {
    "артикул": ARTICLE,
    "номенклатура артикул": ARTICLE,
    "код": CODE,
    "номенклатура код": CODE,
    "код номенклатуры": CODE,
    # Точки выброшены нормализацией: «Ед. изм.», «Ед.изм.» и «Ед.изм» — это
    # одна и та же колонка.
    "ед изм": UNIT,
    "ед": UNIT,
    "единица измерения": UNIT,
    "номенклатура единица измерения": UNIT,
    "номенклатура": NAME,
    "номенклатура характеристика": NAME,
    "номенклатура с характеристикой": NAME,
    "номенклатура наименование": NAME,
    "наименование": NAME,
    "товар": NAME,
    "склад": STORE,
    "магазин": STORE,
    "подразделение": STORE,
    "место хранения": STORE,
}

TOTAL = "итого"

# Строка показателей между складом и полями: «Количество» или «Сумма». Суммы
# в разбор не идут — вкладка считает штуки.
QUANTITY = "количеств"
MEASURES = (QUANTITY, "сумм", "стоимост")

# Признак склада «в пути»: товар отгружен, но не принят.
TRANSIT = "в пути"

# Докуда искать шапку. Над ней стоят название отчёта, параметры и отбор —
# строк немного, но их число меняется вместе с настройками выгрузки.
SEARCH_DEPTH = 40
# На сколько строк выше полей может начинаться шапка: название склада, строка
# «Количество» и объединённые ячейки подписей.
HEAD_DEPTH = 5

# «Склад Равно "Магазин Уссурийск"» в строке отбора — так 1С пишет выгрузку
# по одному складу.
STORE_FILTER = re.compile(r"склад\s+равно\s+[\"«]([^\"»]+)[\"»]", re.IGNORECASE)

LAYOUT_WIDE = "склады в колонках"
LAYOUT_FLAT = "склад колонкой"
LAYOUT_GROUPED = "склады группами строк"
LAYOUT_SINGLE = "один склад"


class ParseError(Exception):
    """Файл не похож на ведомость. Текст пригоден для показа пользователю."""


def _text(value: Any) -> str:
    return " ".join(str(value).split()).strip() if value is not None else ""


def _key(value: Any) -> str:
    """Подпись к сравнимому виду: без регистра, знаков препинания и лишних
    пробелов, «ё» как «е»."""
    text = _text(value).lower().replace("ё", "е")
    for mark in ".,;:":
        text = text.replace(mark, " ")
    return " ".join(text.split())


def _number(value: Any) -> float:
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def item_key(item: Item) -> tuple[str, ...]:
    """Ключ товара при сведении строк и файлов.

    Код 1С однозначен, если он есть. Без него — артикул, название и единица
    вместе: одно название бывает у разных фасовок, а артикулы в карточках
    пишут по-разному.
    """
    if item.code:
        return ("code", item.code)
    return ("item", _key(item.article), _key(item.name), _key(item.unit))


# --- чтение --------------------------------------------------------------------

def read(path: str) -> Ledger:
    """Читает ведомость. Значения формул берутся посчитанными."""
    book = load_workbook(path, data_only=True, read_only=False)
    try:
        stem = os.path.splitext(os.path.basename(path))[0]
        ledger = _sheet(book.worksheets[0], stem)
    finally:
        book.close()
    ledger.sources = [os.path.basename(path)]
    return ledger


@dataclass(slots=True)
class _Block:
    """Колонки одного склада в шапке."""

    name: str
    column: int
    fields: dict[str, int] = field(default_factory=dict)


@dataclass(slots=True)
class _Head:
    """Что нашлось в шапке."""

    top: int
    fields_row: int
    labels: dict[str, int]
    blocks: list[_Block]
    fields: frozenset[str]


def _sheet(sheet: Worksheet, stem: str = "") -> Ledger:
    head = _head(sheet)
    labels = head.labels

    if STORE in labels and labels[STORE] != labels.get(NAME):
        return _flat(sheet, head)
    named = [block for block in head.blocks if block.name]
    if named:
        return _wide(sheet, head)
    if STORE in labels:
        # «Склад» стоит в той же колонке, что и «Номенклатура»: строки сначала
        # сгруппированы по складам, внутри — товары.
        return _grouped(sheet, head)
    return _single(sheet, head, stem)


# --- шапка ---------------------------------------------------------------------

def _field_row(sheet: Worksheet) -> tuple[int, dict[int, str]]:
    """Строка с подписями полей движения. Лучшая из найденных — для ошибки."""
    best: dict[int, str] = {}
    for row in range(1, min(SEARCH_DEPTH, sheet.max_row) + 1):
        found = {cell.column: FIELDS[_key(cell.value)]
                 for cell in sheet[row] if _key(cell.value) in FIELDS}
        if all(name in found.values() for name in REQUIRED):
            return row, found
        if len(set(found.values())) > len(set(best.values())):
            best = found
    if best:
        present = ", ".join(f"«{FIELD_TITLES[name]}»"
                            for name in MOVE_FIELDS if name in best.values())
        missing = ", ".join(f"«{FIELD_TITLES[name]}»"
                            for name in REQUIRED if name not in best.values())
        raise ParseError(
            f"В шапке есть {present}, но нет {missing} — без них разбирать "
            "нечего. Добавьте поле в настройках отчёта в 1С")
    raise ParseError(
        "Не нашлись колонки «Расход» и «Конечный остаток» — нужна ведомость "
        "по товарам на складах, выгруженная из 1С в Excel. Подходят и подписи "
        "«Остаток на начало», «Поступление», «Остаток на конец»")


def _head(sheet: Worksheet) -> _Head:
    fields_row, found = _field_row(sheet)
    first_field = min(found)
    top = max(1, fields_row - HEAD_DEPTH)

    # Подписи товара ищутся только левее полей: справа стоят названия складов,
    # и склад «Магазин Номенклатура» не должен стать колонкой товара.
    labels: dict[str, int] = {}
    label_rows: list[int] = []
    for row in range(top, fields_row + 1):
        for column in range(1, first_field):
            role = LABELS.get(_key(sheet.cell(row, column).value))
            if role and role not in labels:
                labels[role] = column
                label_rows.append(row)
    if NAME not in labels and ARTICLE not in labels:
        seen = [_text(sheet.cell(row, column).value)
                for row in range(top, fields_row + 1)
                for column in range(1, first_field)
                if _text(sheet.cell(row, column).value)]
        shown = ", ".join(f"«{text}»" for text in seen[:8]) or "ничего"
        raise ParseError(
            "Поля движения нашлись, а колонка товара — нет. Ожидалась "
            "«Номенклатура», «Товар» или «Артикул»; левее полей в шапке: "
            f"{shown}")

    store_row, measure_row = _store_rows(sheet, fields_row, first_field,
                                         min(label_rows or [fields_row]))
    fields = _quantities(sheet, measure_row, found, first_field)
    blocks = _blocks(sheet, store_row, fields, first_field)
    present = frozenset(name for block in blocks for name in block.fields)
    return _Head(top=min(label_rows or [fields_row]), fields_row=fields_row,
                 labels=labels, blocks=blocks, fields=present)


def _store_rows(sheet: Worksheet, fields_row: int, first_field: int,
                label_top: int) -> tuple[int, int]:
    """Строки с названиями складов и с показателями над полями. 0 — нет такой.

    Идём вверх от полей: строка, где над полями одни «Количество» и «Сумма»,
    — показатели; первая строка с чем-то другим — склады.
    """
    store_row = measure_row = 0
    if label_top >= fields_row:
        # Подписи товара в одной строке с полями: складов над полями нет, но
        # строка показателей быть может.
        row = fields_row - 1
        texts = [_key(sheet.cell(row, column).value)
                 for column in range(first_field, sheet.max_column + 1)] if row else []
        texts = [text for text in texts if text]
        if texts and all(text.startswith(MEASURES) for text in texts):
            measure_row = row
        return 0, measure_row
    for row in range(fields_row - 1, label_top - 1, -1):
        texts = [_key(sheet.cell(row, column).value)
                 for column in range(first_field, sheet.max_column + 1)]
        texts = [text for text in texts if text]
        if not texts:
            continue
        if all(text.startswith(MEASURES) for text in texts):
            measure_row = measure_row or row
            continue
        store_row = row
        break
    return store_row, measure_row


def _quantities(sheet: Worksheet, measure_row: int, found: dict[int, str],
                first_field: int) -> dict[int, str]:
    """Поля под «Количеством». Суммы отбрасываются: вкладка считает штуки."""
    if not measure_row:
        return found
    kept = {column: name for column, name in found.items()
            if _key(_name_at(sheet, measure_row, column, first_field))
            .startswith(QUANTITY)}
    if not kept:
        raise ParseError(
            "В ведомости только суммы, без количества. Включите показатель "
            "«Количество» в настройках отчёта в 1С")
    return kept


def _blocks(sheet: Worksheet, store_row: int, fields: dict[int, str],
            first_field: int) -> list[_Block]:
    """Колонки полей, разложенные по складам.

    Название склада стоит в объединённой ячейке над первым полем блока. Если
    объединение при пересохранении потерялось, берём ближайшее название слева.
    Новый блок начинается, где сменилось название или повторилось поле, — так
    два склада с одинаковым названием подряд не слипаются в один.
    """
    blocks: list[_Block] = []
    for column in sorted(fields):
        name = _name_at(sheet, store_row, column, first_field) if store_row else ""
        # Своя ячейка с названием — начало склада, даже если название то же,
        # что у соседа слева.
        starts = bool(store_row and _text(sheet.cell(store_row, column).value))
        current = blocks[-1] if blocks else None
        if (current is None or starts or current.name != name
                or fields[column] in current.fields):
            blocks.append(_Block(name=name, column=column))
        blocks[-1].fields[fields[column]] = column
    return blocks


def _name_at(sheet: Worksheet, row: int, column: int, lower: int = 1) -> str:
    """Текст над колонкой. Ищется влево, но не левее `lower`."""
    for step in range(column, lower - 1, -1):
        name = _text(sheet.cell(row, step).value)
        if name:
            return name
    return ""


def _title(sheet: Worksheet, head_row: int) -> str:
    """Название отчёта — первая непустая строка над шапкой."""
    for row in range(1, head_row):
        for cell in sheet[row]:
            if text := _text(cell.value):
                return text
    return "Ведомость по товарам на складах"


# --- строки --------------------------------------------------------------------

class _Rows:
    """Строки данных под шапкой: что товар, что группа, что итог."""

    def __init__(self, sheet: Worksheet, head: _Head) -> None:
        self.sheet = sheet
        self.labels = head.labels
        self.first = head.fields_row + 1
        self.last = sheet.max_row
        # Без колонки единицы группу выдаёт уровень группировки: товары лежат
        # глубже всех. Если уровней нет вовсе, группами ничего не считается.
        levels = [self.level(row) for row in range(self.first, self.last + 1)
                  if not self.total(row)]
        self.deepest = max(levels, default=0)

    def text(self, row: int, role: str) -> str:
        column = self.labels.get(role)
        return _text(self.sheet.cell(row, column).value) if column else ""

    def level(self, row: int) -> int:
        return self.sheet.row_dimensions[row].outline_level or 0

    def total(self, row: int) -> bool:
        """Строка «Итого» внизу отчёта."""
        return any(_key(self.sheet.cell(row, column).value) == TOTAL
                   for column in self.labels.values())

    def group(self, row: int) -> bool:
        if UNIT in self.labels:
            return not self.text(row, UNIT)
        return self.level(row) < self.deepest

    def item(self, row: int) -> Item | None:
        """Товар в строке или None, если это группа, итог или пустая строка."""
        if self.total(row) or self.group(row):
            return None
        item = Item(article=self.text(row, ARTICLE), name=self.text(row, NAME),
                    unit=self.text(row, UNIT), code=self.text(row, CODE))
        if not (item.name or item.article):
            return None
        return item

    def move(self, row: int, block: _Block) -> Movement:
        move = Movement()
        for name, column in block.fields.items():
            setattr(move, name, _number(self.sheet.cell(row, column).value))
        return move

    def reported(self, blocks: list[_Block]) -> Movement | None:
        """Итог отчёта из строки «Итого», если она есть."""
        for row in range(self.last, self.first - 1, -1):
            if self.total(row):
                result = Movement()
                for block in blocks:
                    piece = self.move(row, block)
                    for name in MOVE_FIELDS:
                        setattr(result, name,
                                getattr(result, name) + getattr(piece, name))
                return result
        return None


def _ledger(sheet: Worksheet, head: _Head, layout: str) -> Ledger:
    ledger = Ledger(title=_title(sheet, head.top), layout=layout,
                    fields=head.fields)
    missing = [FIELD_TITLES[name] for name in MOVE_FIELDS
               if name not in head.fields]
    if missing:
        ledger.warnings.append(
            "В выгрузке нет полей: " + ", ".join(f"«{name}»" for name in missing)
            + " — доля «ушло от запаса» не считается")
    return ledger


def _store(name: str, title: str = "", column: int = 0) -> Store:
    return Store(name=name, title=title or name, column=column,
                 transit=TRANSIT in _key(name))


def _collect(items: dict[tuple, Item], item: Item, store: str,
             move: Movement) -> None:
    """Сводит строки одного товара: в плоской таблице он встречается на
    каждом складе отдельной строкой."""
    known = items.setdefault(item_key(item), item)
    if not move.empty:
        known.add(store, move)


# --- виды отчёта ---------------------------------------------------------------

def _wide(sheet: Worksheet, head: _Head) -> Ledger:
    """Склады в колонках: на каждый склад свой блок полей."""
    ledger = _ledger(sheet, head, LAYOUT_WIDE)
    stores: list[tuple[Store, _Block]] = []
    total_block: _Block | None = None
    seen: dict[str, int] = {}
    for block in head.blocks:
        if not block.name:
            continue
        if _key(block.name) == TOTAL:
            total_block = block
            continue
        seen[block.name] = seen.get(block.name, 0) + 1
        # Одноимённые склады различаются номером: без него данные одного
        # затирают данные другого при любой группировке по названию.
        title = (block.name if seen[block.name] == 1
                 else f"{block.name} ({seen[block.name]})")
        stores.append((_store(block.name, title, block.column), block))
    if not stores:
        raise ParseError("В шапке не нашлось ни одного склада")
    ledger.stores = [store for store, _ in stores]

    rows = _Rows(sheet, head)
    for row in range(rows.first, rows.last + 1):
        item = rows.item(row)
        if item is None:
            continue
        for store, block in stores:
            move = rows.move(row, block)
            if not move.empty:
                item.moves[store.title] = move
        ledger.items.append(item)

    if total_block is not None:
        # Итог считается по тем же строкам, что и сами товары: итог по группам
        # в файле продублирован, и суммировать подряд всё, что стоит в
        # колонке «Итого», значило бы получить кратно больше.
        ledger.reported = Movement()
        for row in range(rows.first, rows.last + 1):
            if rows.item(row) is not None:
                ledger.reported = _plus(ledger.reported,
                                        rows.move(row, total_block))
    else:
        ledger.reported = rows.reported([block for _, block in stores])
    return ledger


def _flat(sheet: Worksheet, head: _Head) -> Ledger:
    """Плоская таблица: склад указан в каждой строке."""
    ledger = _ledger(sheet, head, LAYOUT_FLAT)
    block = head.blocks[0]
    rows = _Rows(sheet, head)
    stores: dict[str, Store] = {}
    items: dict[tuple, Item] = {}
    for row in range(rows.first, rows.last + 1):
        item = rows.item(row)
        name = rows.text(row, STORE)
        if item is None or not name:
            continue
        store = stores.setdefault(name, _store(name))
        _collect(items, item, store.title, rows.move(row, block))
    if not stores:
        raise ParseError("В колонке склада не нашлось ни одного склада")
    ledger.stores = list(stores.values())
    ledger.items = list(items.values())
    ledger.reported = rows.reported([block])
    return ledger


def _grouped(sheet: Worksheet, head: _Head) -> Ledger:
    """Склады группами строк: строка склада, под ней его товары.

    Строку склада от строки группы товаров отличает уровень группировки:
    склады — самый верхний уровень среди групп. Если уровней в файле нет,
    различить их нечем, и гадать здесь нельзя.
    """
    ledger = _ledger(sheet, head, LAYOUT_GROUPED)
    block = head.blocks[0]
    rows = _Rows(sheet, head)
    groups = [row for row in range(rows.first, rows.last + 1)
              if not rows.total(row) and rows.group(row)
              and rows.text(row, NAME)]
    if not groups:
        raise ParseError(
            "В шапке указано, что строки сгруппированы по складам, но строк "
            "складов в файле нет")
    if rows.deepest == 0:
        raise ParseError(
            "Строки сгруппированы по складам, но в файле не сохранилась "
            "группировка строк — склад не отличить от группы товаров. "
            "Выгрузите отчёт из 1С заново, не пересохраняя его")
    top = min(rows.level(row) for row in groups)

    stores: dict[str, Store] = {}
    items: dict[tuple, Item] = {}
    current: Store | None = None
    for row in range(rows.first, rows.last + 1):
        if rows.total(row):
            continue
        if rows.group(row):
            if rows.level(row) == top and (name := rows.text(row, NAME)):
                current = stores.setdefault(name, _store(name))
            continue
        item = rows.item(row)
        if item is None or current is None:
            continue
        _collect(items, item, current.title, rows.move(row, block))
    ledger.stores = list(stores.values())
    ledger.items = list(items.values())
    ledger.reported = rows.reported([block])
    return ledger


def _single(sheet: Worksheet, head: _Head, stem: str) -> Ledger:
    """Один склад: названия в шапке нет, берём его из отбора или имени файла."""
    ledger = _ledger(sheet, head, LAYOUT_SINGLE)
    name = _filtered_store(sheet, head.top)
    if not name:
        name = stem or "Склад"
        ledger.warnings.append(
            f"Склад в файле не указан — назван по имени файла: «{name}»")
    block = head.blocks[0]
    store = _store(name, column=block.column)
    ledger.stores = [store]
    rows = _Rows(sheet, head)
    for row in range(rows.first, rows.last + 1):
        item = rows.item(row)
        if item is None:
            continue
        move = rows.move(row, block)
        if not move.empty:
            item.moves[store.title] = move
        ledger.items.append(item)
    ledger.reported = rows.reported([block])
    return ledger


def _filtered_store(sheet: Worksheet, head_row: int) -> str:
    """Склад из строки отбора: «Склад Равно "Магазин Уссурийск"».

    Обрезанное 1С название («Магазин Седанка Сити Перв...») не годится: по
    нему склад не узнать, лучше имя файла.
    """
    for row in range(1, head_row):
        for cell in sheet[row]:
            match = STORE_FILTER.search(_text(cell.value))
            if match and not match.group(1).rstrip().endswith("..."):
                return match.group(1).strip()
    return ""


def _plus(left: Movement, right: Movement) -> Movement:
    return Movement(*(getattr(left, name) + getattr(right, name)
                      for name in MOVE_FIELDS))
