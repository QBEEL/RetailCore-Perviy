"""Ведомость в понятном виде: несколько листов вместо одной широкой простыни.

Исходная выгрузка нечитаема не из-за оформления, а из-за формы: шестнадцать
складов по четыре колонки в строке, группы вперемешку с товарами и итог рядом с
данными. Поэтому здесь не «та же таблица, но покрашенная», а другая раскладка —
один лист на вопрос, который задают этой ведомости.

Первым идёт лист с выводами, а не свод: свод нужен, когда к цифре появились
вопросы, а начинают всегда с того, чего не хватило.
"""
from __future__ import annotations

import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .analysis import Line, Report

FONT = "Calibri"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="DCE3F0")
WARN_FILL = PatternFill("solid", fgColor="FCE4E4")
STRIPE = PatternFill("solid", fgColor="F4F6FB")

HAIR = Side(style="hair", color="B4BFD4")
GRID = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)

QTY = "#,##0"
SHARE = "0.0%"

MIN_WIDTH, MAX_WIDTH = 9, 56


def save(report: Report, destination: str) -> str:
    book = Workbook()
    book.remove(book.active)
    _summary(book.create_sheet("Главное"), report)
    _shortages(book.create_sheet("Не хватило"), report)
    _top(book.create_sheet("Топ товаров"), report)
    _stores(book.create_sheet("По магазинам"), report)
    _matrix(book.create_sheet("Свод"), report)
    book.save(destination)
    return destination


def default_name(folder: str) -> str:
    return os.path.join(folder, f"Ведомость — разбор {date.today():%Y-%m-%d}.xlsx")


# --- оформление ----------------------------------------------------------------

def _title(sheet: Worksheet, row: int, text: str, width: int) -> int:
    cell = sheet.cell(row, 1, text)
    cell.font = Font(FONT, size=13, bold=True, color="1F3864")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    return row + 1


def _note(sheet: Worksheet, row: int, text: str, width: int) -> int:
    cell = sheet.cell(row, 1, text)
    cell.font = Font(FONT, size=9, italic=True, color="5A6A85")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    sheet.row_dimensions[row].height = 26
    return row + 1


def _head(sheet: Worksheet, row: int, titles: list[str]) -> int:
    for column, text in enumerate(titles, 1):
        cell = sheet.cell(row, column, text)
        cell.font = Font(FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.border = GRID
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    sheet.row_dimensions[row].height = 30
    sheet.freeze_panes = sheet.cell(row + 1, 1)
    return row + 1


def _row(sheet: Worksheet, row: int, values: list, *, formats: str = "",
         stripe: bool = False, warn: bool = False) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(FONT, size=10)
        cell.border = GRID
        kind = formats[column - 1] if column - 1 < len(formats) else "t"
        if kind == "n":
            cell.number_format = QTY
            cell.alignment = Alignment(horizontal="right")
        elif kind == "%":
            cell.number_format = SHARE
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
        if warn:
            cell.fill = WARN_FILL
        elif stripe:
            cell.fill = STRIPE


def _widths(sheet: Worksheet, head_row: int) -> None:
    for column in range(1, sheet.max_column + 1):
        longest = 0
        for row in range(head_row, sheet.max_row + 1):
            value = sheet.cell(row, column).value
            if value is None:
                continue
            longest = max(longest, len(str(value)) if isinstance(value, str)
                          else len(f"{value:,.0f}"))
        sheet.column_dimensions[get_column_letter(column)].width = min(
            MAX_WIDTH, max(MIN_WIDTH, longest + 2))


# --- листы ---------------------------------------------------------------------

def _summary(sheet: Worksheet, report: Report) -> None:
    row = _title(sheet, 1, report.ledger.title, 4)
    row = _note(sheet, row,
                "«Расход» — всё, что ушло со склада: продажи, списания и "
                "перемещения. Разделить их в этой выгрузке нечем, поэтому "
                "столбец так и назван. Период в выгрузке не указан, поэтому "
                "всё в штуках и долях, без расчётов на дни.", 4)
    row += 1

    total = report.total
    rows = [
        ("Магазинов", len(report.ledger.shops)),
        ("Товаров в ведомости", len(report.ledger.items)),
        ("Начальный остаток", total.opening),
        ("Приход", total.incoming),
        ("Расход", total.outgoing),
        ("Конечный остаток", total.closing),
        ("Израсходовано от запаса", total.share),
        ("Кончилось, а расход был", len(report.shortages)),
        ("Остатка меньше, чем ушло", len(report.tight)),
        ("Отрицательный остаток", len(report.negative)),
        ("Лежит в пути", sum(line.move.closing for line in report.transit)),
    ]
    row = _head(sheet, row, ["Показатель", "Значение"])
    for number, (name, value) in enumerate(rows):
        fmt = "t%" if name == "Израсходовано от запаса" else "tn"
        _row(sheet, row, [name, value], formats=fmt, stripe=number % 2 == 1)
        row += 1

    if report.doubled:
        row += 1
        names = ", ".join(store.title for store in report.doubled)
        row = _note(sheet, row,
                    "В выгрузке есть склады с одинаковым названием — они "
                    f"разведены номером: {names}. Складывать их в один склад "
                    "нельзя: данные разные.", 4)
    _widths(sheet, 1)


def _lines_sheet(sheet: Worksheet, report: Report, lines: list[Line], *,
                 title: str, note: str, warn: bool = False) -> None:
    row = _title(sheet, 1, title, 6)
    row = _note(sheet, row, note, 6)
    row += 1
    head = row
    row = _head(sheet, row, ["Артикул", "Товар", "Магазин", "Расход",
                             "Конечный остаток", "Ушло от запаса"])
    for number, line in enumerate(lines):
        _row(sheet, row, [line.article, line.name, line.store,
                          line.move.outgoing, line.move.closing,
                          line.move.share],
             formats="tttnn%", stripe=number % 2 == 1, warn=warn)
        row += 1
    if not lines:
        _row(sheet, row, ["", "таких строк нет", "", None, None, None])
    _widths(sheet, head)


def _shortages(sheet: Worksheet, report: Report) -> None:
    _lines_sheet(
        sheet, report, report.shortages,
        title="Кончилось, а расход был",
        note="Расход был, а на конец периода не осталось ничего. Считать это "
             "потерянными продажами нельзя: дат в выгрузке нет, и товар мог "
             "кончиться в последний день. Но смотреть этот список нужно первым.",
        warn=True)

    # Отрицательные остатки — следом, на том же листе: это другая болезнь,
    # но узнать о ней нужно там же, где смотрят нехватку.
    row = sheet.max_row + 2
    row = _title(sheet, row, "Отрицательный остаток — это ошибка учёта", 6)
    row = _note(sheet, row,
                "Расход больше, чем было на складе. Завозом не лечится: ищите "
                "непроведённый приход или пересорт. Остальные числа по такой "
                "строке тоже под вопросом.", 6)
    row = _head(sheet, row, ["Артикул", "Товар", "Магазин", "Расход",
                             "Конечный остаток", "Ушло от запаса"])
    for number, line in enumerate(report.negative):
        _row(sheet, row, [line.article, line.name, line.store,
                          line.move.outgoing, line.move.closing,
                          line.move.share],
             formats="tttnn%", stripe=number % 2 == 1, warn=True)
        row += 1
    if not report.negative:
        _row(sheet, row, ["", "отрицательных остатков нет", "", None, None, None])


def _top(sheet: Worksheet, report: Report) -> None:
    row = _title(sheet, 1, "Топ товаров по расходу — все магазины", 8)
    row = _note(sheet, row,
                "Сумма по магазинам, склады «в пути» не считаются. «Магазинов» "
                "— в скольких был расход: товар с большим расходом в одном "
                "магазине и товар, который расходится везде, требуют разного.", 8)
    row += 1
    head = row
    row = _head(sheet, row, ["№", "Артикул", "Товар", "Ед.", "Расход",
                             "Доля расхода", "Конечный остаток", "Магазинов"])
    for number, item in enumerate(report.items, 1):
        _row(sheet, row, [number, item.article, item.name, item.item.unit,
                          item.move.outgoing, item.share, item.move.closing,
                          len(item.active)],
             formats="ntttn%nn", stripe=number % 2 == 0,
             warn=bool(item.ran_out))
        row += 1
    _widths(sheet, head)


def _stores(sheet: Worksheet, report: Report) -> None:
    row = _title(sheet, 1, "По магазинам", 7)
    row = _note(sheet, row,
                "Магазины по убыванию расхода, под каждым — что расходилось "
                "лучше всего именно в нём.", 7)
    row += 1
    head = row
    row = _head(sheet, row, ["Магазин", "Товар", "Расход", "Конечный остаток",
                             "Товаров с расходом", "Кончилось", "Ушло от запаса"])
    for total in report.stores:
        cell_row = [total.store.title, "— всего по магазину", total.move.outgoing,
                    total.move.closing, total.items, total.ran_out,
                    total.move.share]
        _row(sheet, row, cell_row, formats="ttnnnn%")
        for column in range(1, 8):
            sheet.cell(row, column).font = Font(FONT, size=10, bold=True)
            sheet.cell(row, column).fill = SUB_FILL
        row += 1
        for number, line in enumerate(total.top):
            _row(sheet, row, ["", line.name, line.move.outgoing,
                              line.move.closing, None, None, line.move.share],
                 formats="ttnntt%", stripe=number % 2 == 1,
                 warn=line.move.ran_out)
            row += 1
    _widths(sheet, head)


def _matrix(sheet: Worksheet, report: Report) -> None:
    """Свод: товары в строках, магазины в колонках. Расход и остаток рядом."""
    shops = report.ledger.shops
    row = _title(sheet, 1, "Свод: расход и остаток по магазинам", 3 + 2 * len(shops))
    row = _note(sheet, row,
                "На каждый магазин две колонки: расход и конечный остаток. "
                "Склады «в пути» в свод не входят — их остаток на листе "
                "«Главное».", 3 + 2 * len(shops))
    row += 1

    # Две строки шапки: магазин над своей парой колонок.
    top = row
    sheet.cell(top, 1, "Артикул")
    sheet.cell(top, 2, "Товар")
    sheet.cell(top, 3, "Ед.")
    for index, store in enumerate(shops):
        column = 4 + index * 2
        sheet.cell(top, column, store.title)
        sheet.merge_cells(start_row=top, start_column=column,
                          end_row=top, end_column=column + 1)
    for index in range(len(shops)):
        column = 4 + index * 2
        sheet.cell(top + 1, column, "расход")
        sheet.cell(top + 1, column + 1, "остаток")
    for column in range(1, 4):
        sheet.merge_cells(start_row=top, start_column=column,
                          end_row=top + 1, end_column=column)
    for line in (top, top + 1):
        for column in range(1, 4 + 2 * len(shops)):
            cell = sheet.cell(line, column)
            cell.font = Font(FONT, size=9, bold=True, color="FFFFFF")
            cell.fill = HEAD_FILL
            cell.border = GRID
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
    sheet.row_dimensions[top].height = 28
    sheet.freeze_panes = sheet.cell(top + 2, 4)
    row = top + 2

    for number, item in enumerate(report.items):
        values = [item.article, item.name, item.item.unit]
        for store in shops:
            move = item.item.move(store.title)
            values.extend([move.outgoing or None, move.closing or None])
        _row(sheet, row, values, formats="ttt" + "nn" * len(shops),
             stripe=number % 2 == 1)
        row += 1
    _widths(sheet, top)
