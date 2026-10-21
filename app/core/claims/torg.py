"""Акт ТОРГ-2 «об установленном расхождении по количеству и качеству при приёмке
товарно-материальных ценностей»: данные, расчёт и заполнение бланка Excel.

Бланк — не рисунок, а настоящая книга Excel с тремя листами и сотнями
объединённых ячеек: так его принимают поставщики и так его вёл склад. Поэтому
акт не рисуется заново, а заполняется по шаблону `torg2_template.xlsx`: в нём
остался только текст самой формы, а всё, что склад вписывал руками, пусто.

Модуль не зависит от Qt и сети — всё здесь проверяется обычным pytest.
"""
from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..payments.safe_files import safe_name

TEMPLATE_NAME = "torg2_template.xlsx"

MONTHS = ("января", "февраля", "марта", "апреля", "мая", "июня", "июля",
          "августа", "сентября", "октября", "ноября", "декабря")

# Сколько человек может быть в комиссии кроме председателя: на третьем листе
# бланка под них три строки, и четвёртой там нет.
MAX_MEMBERS = 3

# Сколько строк занимают описание дефектов и заключение комиссии на третьем
# листе и сколько знаков помещается в строке. Первая строка короче остальных:
# она начинается посреди надписи «…и мнение комиссии о причинах их образования».
# Знаки подобраны по печати в PDF (Times New Roman 11 пт) с запасом: текст,
# вылезший за край, в бланке просто обрезается, а заметить это можно только на
# бумаге.
DESCRIPTION_LINES = (46, 88, 88, 88)
CONCLUSION_LINES = (62, 88, 88)


# --- данные ----------------------------------------------------------------------

@dataclass(slots=True)
class TorgSettings:
    """Реквизиты организации и комиссия — общие для всех актов отдела."""

    organization: str = ""
    department: str = "Отдел Склад"
    director_title: str = ""
    director_name: str = ""
    place: str = ""
    chair_title: str = "Руководитель склада"
    chair_name: str = ""
    # Члены комиссии: (должность, «Фамилия И.»). Не больше трёх.
    members: list[tuple[str, str]] = field(default_factory=list)

    @classmethod
    def from_dict(cls, data: dict[str, Any] | None) -> TorgSettings:
        data = data or {}
        members = [(str(m[0]), str(m[1])) for m in (data.get("members") or [])
                   if len(m) >= 2 and (str(m[0]).strip() or str(m[1]).strip())]
        return cls(
            organization=str(data.get("organization") or ""),
            department=str(data.get("department") or "Отдел Склад"),
            director_title=str(data.get("director_title") or ""),
            director_name=str(data.get("director_name") or ""),
            place=str(data.get("place") or ""),
            chair_title=str(data.get("chair_title") or "Руководитель склада"),
            chair_name=str(data.get("chair_name") or ""),
            members=members[:MAX_MEMBERS],
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "organization": self.organization, "department": self.department,
            "director_title": self.director_title, "director_name": self.director_name,
            "place": self.place, "chair_title": self.chair_title,
            "chair_name": self.chair_name,
            "members": [list(m) for m in self.members[:MAX_MEMBERS]],
        }

    def problems(self) -> list[str]:
        """Что не заполнено в реквизитах — без этого бланк выйдет с пустыми местами."""
        missing = []
        for value, title in ((self.organization, "организация"),
                             (self.director_name, "руководитель"),
                             (self.place, "место приёмки"),
                             (self.chair_name, "председатель комиссии")):
            if not value.strip():
                missing.append(title)
        return missing


@dataclass(slots=True)
class TorgItem:
    """Строка товара в акте. Одна заявка Торг-2 — одна строка."""

    claim_id: int = 0
    name: str = ""
    barcode: str = ""
    unit: str = "шт"
    # Цена за единицу по документам поставщика.
    price: float = 0.0
    qty_doc: float = 0.0
    # Недостача — сколько не довезли. Брак — сколько привезли негодным. Излишек —
    # сколько привезли сверх документов.
    qty_short: float = 0.0
    qty_defect: float = 0.0
    qty_surplus: float = 0.0
    # Код маркировки экземпляра (DataMatrix), если товар маркируемый: уходит на
    # четвёртую страницу акта.
    mark_code: str = ""

    @property
    def qty_fact(self) -> float:
        """Фактически поступило. Бракованное тоже поступило — оно лежит на складе,
        поэтому из фактического количества не вычитается: так в образце,
        36 по документам, 36 фактически, 11 из них брак. Излишек поступил сверх
        документов и к фактическому количеству прибавляется."""
        return self.qty_doc - self.qty_short + self.qty_surplus

    @property
    def sum_doc(self) -> float:
        return money(self.qty_doc * self.price)

    @property
    def sum_fact(self) -> float:
        return money(self.qty_fact * self.price)

    @property
    def sum_short(self) -> float:
        return money(self.qty_short * self.price)

    @property
    def sum_defect(self) -> float:
        return money(self.qty_defect * self.price)

    @property
    def sum_surplus(self) -> float:
        return money(self.qty_surplus * self.price)

    def to_dict(self) -> dict[str, Any]:
        return {
            "claim_id": self.claim_id, "name": self.name, "barcode": self.barcode,
            "unit": self.unit, "price": self.price, "qty_doc": self.qty_doc,
            "qty_short": self.qty_short, "qty_defect": self.qty_defect,
            "qty_surplus": self.qty_surplus, "mark_code": self.mark_code,
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> TorgItem:
        return cls(
            claim_id=int(data.get("claim_id") or 0), name=str(data.get("name") or ""),
            barcode=str(data.get("barcode") or ""), unit=str(data.get("unit") or "шт"),
            price=float(data.get("price") or 0), qty_doc=float(data.get("qty_doc") or 0),
            qty_short=float(data.get("qty_short") or 0),
            qty_defect=float(data.get("qty_defect") or 0),
            qty_surplus=float(data.get("qty_surplus") or 0),
            mark_code=str(data.get("mark_code") or ""))


@dataclass(slots=True)
class TorgAct:
    id: int = 0
    number: int = 0
    act_date: date = field(default_factory=date.today)
    # Когда товар приняли на склад. Часто совпадает с датой акта, но акт
    # составляют и на следующий день после приёмки.
    receipt_date: date = field(default_factory=date.today)
    # Поставщик одной строкой: название, адрес, телефон — как в бланке.
    supplier: str = ""
    contract_number: str = ""
    contract_date: date | None = None
    invoice_number: str = ""
    invoice_date: date | None = None
    delivery: str = ""
    waybill: str = ""
    description: str = ""
    conclusion: str = ""
    items: list[TorgItem] = field(default_factory=list)

    @property
    def claim_ids(self) -> list[int]:
        return [item.claim_id for item in self.items if item.claim_id]

    def payload(self) -> dict[str, Any]:
        """Всё, кроме номера и даты акта: они лежат в базе отдельными полями."""
        return {
            "receipt_date": self.receipt_date.isoformat(),
            "contract_number": self.contract_number,
            "contract_date": self.contract_date.isoformat() if self.contract_date else "",
            "invoice_number": self.invoice_number,
            "invoice_date": self.invoice_date.isoformat() if self.invoice_date else "",
            "delivery": self.delivery, "waybill": self.waybill,
            "description": self.description, "conclusion": self.conclusion,
            "items": [item.to_dict() for item in self.items],
        }

    @classmethod
    def from_parts(cls, *, id: int, number: int, act_date: date, supplier: str,
                   payload: dict[str, Any]) -> TorgAct:
        return cls(
            id=id, number=number, act_date=act_date, supplier=supplier,
            receipt_date=_day(payload.get("receipt_date")) or act_date,
            contract_number=str(payload.get("contract_number") or ""),
            contract_date=_day(payload.get("contract_date")),
            invoice_number=str(payload.get("invoice_number") or ""),
            invoice_date=_day(payload.get("invoice_date")),
            delivery=str(payload.get("delivery") or ""),
            waybill=str(payload.get("waybill") or ""),
            description=str(payload.get("description") or ""),
            conclusion=str(payload.get("conclusion") or ""),
            items=[TorgItem.from_dict(i) for i in payload.get("items") or []])

    # --- расчёт ---------------------------------------------------------------

    def totals(self) -> dict[str, float]:
        items = self.items
        return {
            "qty_doc": sum(i.qty_doc for i in items),
            "sum_doc": money(sum(i.sum_doc for i in items)),
            "qty_fact": sum(i.qty_fact for i in items),
            "sum_fact": money(sum(i.sum_fact for i in items)),
            "qty_defect": sum(i.qty_defect for i in items),
            "sum_defect": money(sum(i.sum_defect for i in items)),
            "qty_short": sum(i.qty_short for i in items),
            "sum_short": money(sum(i.sum_short for i in items)),
            "qty_surplus": sum(i.qty_surplus for i in items),
            "sum_surplus": money(sum(i.sum_surplus for i in items)),
        }

    def problems(self) -> list[str]:
        """Что мешает составить акт. Пустой список — можно."""
        found: list[str] = []
        if not self.supplier.strip():
            found.append("Укажите поставщика")
        if not self.invoice_number.strip():
            found.append("Укажите номер счёта-фактуры")
        if self.invoice_date is None:
            found.append("Укажите дату счёта-фактуры")
        if not self.items:
            found.append("В акте нет товаров")
        for number, item in enumerate(self.items, 1):
            where = f"Строка {number}"
            # Излишек — товар, которого в документах могло не быть вовсе: тогда
            # количество по документам ноль, и требовать его нельзя.
            if item.qty_doc <= 0 and item.qty_surplus <= 0:
                found.append(f"{where}: укажите количество по документам")
            if item.price <= 0:
                found.append(f"{where}: укажите цену")
            if item.qty_short < 0 or item.qty_defect < 0 or item.qty_surplus < 0:
                found.append(f"{where}: количество не может быть отрицательным")
            if item.qty_short > item.qty_doc:
                found.append(f"{where}: недостача больше количества по документам")
            if item.qty_defect > item.qty_fact:
                found.append(f"{where}: брака больше, чем поступило")
            if item.qty_short <= 0 and item.qty_defect <= 0 and item.qty_surplus <= 0:
                found.append(f"{where}: укажите недостачу, брак или излишек")
        if overflow(self.description, DESCRIPTION_LINES):
            found.append("Описание не помещается в бланк — сократите его")
        if overflow(self.conclusion, CONCLUSION_LINES):
            found.append("Заключение не помещается в бланк — сократите его")
        return found

    def suggested_text(self) -> tuple[str, str]:
        """Описание и заключение по умолчанию — по тому, что в акте есть.

        Подсказка, а не готовый текст: человек её правит. Она нужна, чтобы не
        печатать по десять раз одни и те же слова про замятую упаковку.
        """
        shortage = any(i.qty_short > 0 for i in self.items)
        defect = any(i.qty_defect > 0 for i in self.items)
        surplus = any(i.qty_surplus > 0 for i in self.items)
        if surplus and (shortage or defect):
            found = [name for name, present in (
                ("недостача по количеству", shortage), ("брак товара", defect),
                ("излишек по количеству", surplus)) if present]
            return (f"Выявлено при приёмке: {', '.join(found)}.",
                    "выставить УКД на недопоставленную и бракованную продукцию, "
                    "уточнить у поставщика излишне поставленную"
                    if shortage and defect else
                    ("выставить УКД на недопоставленную продукцию, уточнить у "
                     "поставщика излишне поставленную" if shortage else
                     "выставить УКД на бракованную продукцию, уточнить у "
                     "поставщика излишне поставленную"))
        if surplus:
            return ("Фактическое количество больше указанного в сопроводительных "
                    "документах.",
                    "уточнить у поставщика излишне поставленную продукцию")
        if defect and shortage:
            return ("Недостача по количеству и брак товара, выявленные при приёмке.",
                    "выставить УКД на недопоставленную и бракованную продукцию")
        if defect:
            return ("Товар ненадлежащего качества, выявлено при приёмке.",
                    "выставить УКД на бракованную продукцию")
        return ("Фактическое количество меньше указанного в сопроводительных документах.",
                "выставить УКД на недопоставленную продукцию")


def act_from_claims(claims: list[Any], number: int, today: date | None = None) -> TorgAct:
    """Новый акт по выбранным заявкам: каждая заявка — строка товара.

    Из заявки берутся название, штрих-код, код маркировки и то, что случилось:
    недопоставка идёт в недостачу, брак — в брак, излишек — в излишки. Цену и количество по документам заявка не знает —
    их вносит склад по счёту-фактуре.
    """
    today = today or date.today()
    items = []
    for claim in claims:
        name = " ".join(part for part in (claim.product, claim.spec) if part)
        items.append(TorgItem(
            claim_id=claim.id, name=name or claim.brand or claim.barcode,
            barcode=claim.barcode,
            qty_short=claim.quantity if claim.kind == "shortage" else 0.0,
            qty_defect=claim.quantity if claim.kind == "defect" else 0.0,
            qty_surplus=claim.quantity if claim.kind == "surplus" else 0.0,
            mark_code=getattr(claim, "mark_code", "") or ""))
    return TorgAct(number=number, act_date=today, receipt_date=today, items=items)


# --- вспомогательное ---------------------------------------------------------------

def money(value: float) -> float:
    return round(float(value) + 1e-9, 2)


def _day(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def short_supplier(supplier: str) -> str:
    """Название поставщика без адреса — для имени файла: «ООО Ромашка»."""
    head = re.split(r",|\s\d{5,}", supplier, maxsplit=1)[0]
    return " ".join(head.replace('"', "").replace("«", "").replace("»", "").split())


def file_name(act: TorgAct) -> str:
    """Имя файла акта. Номер — в имени, чтобы вложение в заявке узнавалось сразу."""
    name = f"ТОРГ-2 №{act.number} {short_supplier(act.supplier)} {act.act_date:%d-%m-%Y}"
    return safe_name(name.strip() + ".xlsx", "ТОРГ-2.xlsx")


def _wrap(text: str, widths: tuple[int, ...]) -> list[str]:
    """Режет текст по словам на строки бланка, не отбрасывая лишнего.

    Явный перенос строки в тексте уважается: человек мог начать новый абзац.
    Слово длиннее строки остаётся в ней целиком — разрезать его было бы хуже,
    а о том, что оно не помещается, скажет `overflow`.
    """
    lines: list[str] = []
    for paragraph in text.replace("\r", "").split("\n"):
        current = ""
        for word in paragraph.split():
            limit = widths[min(len(lines), len(widths) - 1)]
            candidate = f"{current} {word}".strip()
            if len(candidate) <= limit or not current:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    while lines and not lines[-1]:
        lines.pop()
    return lines


def wrap_lines(text: str, widths: tuple[int, ...]) -> list[str]:
    """Строки для вписывания в бланк. Лишнее отбрасывается — заранее о нём
    предупреждает `overflow`, и до записи в файл дело не доходит."""
    return _wrap(text, widths)[:len(widths)]


def overflow(text: str, widths: tuple[int, ...]) -> bool:
    """Не помещается ли текст в строки бланка."""
    lines = _wrap(text, widths)
    return len(lines) > len(widths) or any(
        len(line) > widths[min(i, len(widths) - 1)] for i, line in enumerate(lines))


def template_path() -> Path:
    from ...ui.resources import asset

    return asset(TEMPLATE_NAME)


# --- заполнение бланка ------------------------------------------------------------

# Первая строка товаров на втором листе и сколько их в шаблоне.
ITEMS_FIRST_ROW = 40
ITEMS_IN_TEMPLATE = 2


def _two(number: int) -> str:
    return f"{number:02d}"


def build_workbook(act: TorgAct, settings: TorgSettings,
                   template: str | Path | None = None) -> openpyxl.Workbook:
    """Заполняет бланк. Шаблон не меняется: берётся копия."""
    book = openpyxl.load_workbook(template or template_path())
    first, second, third = book["стр1"], book["стр2"], book["стр3"]

    _fill_first_page(first, act, settings)
    _fill_items(second, act)
    _fill_third_page(third, act, settings)
    _fill_codes_page(book, act)
    _clear_yellow(book)
    # Свойства файла: автор и «кем изменён» остались бы от образца склада, а
    # акт уходит поставщику.
    book.properties.title = f"ТОРГ-2 №{act.number}"
    book.properties.creator = "RetailCore"
    book.properties.lastModifiedBy = None
    book.properties.lastPrinted = None
    return book


def save(act: TorgAct, settings: TorgSettings, path: str | Path,
         template: str | Path | None = None) -> Path:
    """Собирает акт и сохраняет в файл. Возвращает путь."""
    problems = settings.problems() + act.problems()
    if problems:
        raise ValueError("; ".join(problems))
    book = build_workbook(act, settings, template)
    # Во временный файл и переименование: оборванная запись не должна остаться
    # под именем целого акта.
    target = Path(path)
    partial = target.with_name(target.name + ".part")
    book.save(partial)
    partial.replace(target)
    return target


def _put(ws, address: str, value: Any) -> None:
    ws[address] = value


def _fill_first_page(ws, act: TorgAct, s: TorgSettings) -> None:
    d = act.act_date
    _put(ws, "A6", s.organization)
    _put(ws, "A9", s.department)
    _put(ws, "CI11", str(act.number))
    _put(ws, "CI12", _two(d.day))
    _put(ws, "CP12", _two(d.month))
    _put(ws, "CW12", str(d.year))
    _put(ws, "CA16", s.director_title)
    _put(ws, "AY17", str(act.number))
    _put(ws, "BK17", datetime(d.year, d.month, d.day))
    _put(ws, "CK18", s.director_name)
    _put(ws, "CD20", _two(d.day))
    _put(ws, "CI20", MONTHS[d.month - 1].capitalize())
    _put(ws, "CS20", str(d.year))
    _put(ws, "Z23", s.place)
    _put(ws, "BL24", _two(d.day))
    _put(ws, "BS24", MONTHS[d.month - 1])
    _put(ws, "CS24", str(d.year))
    invoice = f"счет-фактура {act.invoice_number} от {act.invoice_date:%d.%m.%Y} г." \
        if act.invoice_date else f"счет-фактура {act.invoice_number}"
    _put(ws, "AL25", invoice)
    _put(ws, "U33", act.supplier)
    _put(ws, "N37", act.supplier)
    if act.contract_number:
        _put(ws, "AT42", act.contract_number)
    if cd := act.contract_date:
        _put(ws, "BN42", _two(cd.day))
        _put(ws, "BU42", _two(cd.month))
        _put(ws, "CG42", str(cd.year))
    _put(ws, "U43", act.invoice_number)
    if idate := act.invoice_date:
        _put(ws, "AO43", _two(idate.day))
        _put(ws, "AV43", _two(idate.month))
        _put(ws, "BH43", str(idate.year))
    _put(ws, "T48", act.delivery)
    _put(ws, "CP48", act.waybill)
    receipt = f"{act.receipt_date:%d.%m.%Y}"
    for cell in ("AQ60", "BP60", "CT60"):
        _put(ws, cell, receipt)


# Колонки строки товара на втором листе: куда что пишется.
_ITEM_CELLS = {
    "name": "A", "unit": "AL", "barcode_doc": "BS", "qty_doc": "BY", "price_doc": "CD",
    "sum_doc": "CQ", "barcode_fact": "DR", "qty_fact": "DV", "price_fact": "DY",
    "sum_fact": "EF", "qty_defect": "EM", "sum_defect": "ET",
    "qty_short": "FO", "sum_short": "FW",
    "qty_surplus": "GE", "sum_surplus": "GM",
}


def _qty(value: float) -> float | int:
    return int(value) if float(value).is_integer() else value


def _fill_items(ws, act: TorgAct) -> None:
    receipt = act.receipt_date
    _put(ws, "EB20", _two(receipt.day))
    _put(ws, "EI20", _two(receipt.month))
    _put(ws, "FA20", str(receipt.year))

    count = len(act.items)
    _resize_items(ws, count)
    for offset, item in enumerate(act.items):
        row = ITEMS_FIRST_ROW + offset
        values = {
            "name": item.name, "unit": item.unit,
            "barcode_doc": item.barcode, "qty_doc": _qty(item.qty_doc),
            "price_doc": item.price, "sum_doc": item.sum_doc,
            "barcode_fact": item.barcode, "qty_fact": _qty(item.qty_fact),
            "price_fact": item.price, "sum_fact": item.sum_fact,
        }
        # Колонки брака и недостачи заполняются только там, где они есть: нули
        # на бланке читаются как «проверили, и ноль», а не «не применимо».
        if item.qty_defect:
            values["qty_defect"] = _qty(item.qty_defect)
            values["sum_defect"] = item.sum_defect
        if item.qty_short:
            values["qty_short"] = _qty(item.qty_short)
            values["sum_short"] = item.sum_short
        if item.qty_surplus:
            values["qty_surplus"] = _qty(item.qty_surplus)
            values["sum_surplus"] = item.sum_surplus
        for key, value in values.items():
            ws[f"{_ITEM_CELLS[key]}{row}"] = value
        # Штрих-код — текстом: иначе Excel покажет 8.01733E+12.
        for key in ("barcode_doc", "barcode_fact"):
            ws[f"{_ITEM_CELLS[key]}{row}"].number_format = "@"

    total_row = ITEMS_FIRST_ROW + count
    totals = act.totals()
    ws[f"BS{total_row}"] = "Итого"
    ws[f"BY{total_row}"] = _qty(totals["qty_doc"])
    ws[f"CQ{total_row}"] = totals["sum_doc"]
    ws[f"DV{total_row}"] = _qty(totals["qty_fact"])
    ws[f"EF{total_row}"] = totals["sum_fact"]
    if totals["qty_defect"]:
        ws[f"EM{total_row}"] = _qty(totals["qty_defect"])
        ws[f"ET{total_row}"] = totals["sum_defect"]
    if totals["qty_short"]:
        ws[f"FO{total_row}"] = _qty(totals["qty_short"])
        ws[f"FW{total_row}"] = totals["sum_short"]
    if totals["qty_surplus"]:
        ws[f"GE{total_row}"] = _qty(totals["qty_surplus"])
        ws[f"GM{total_row}"] = totals["sum_surplus"]


def _resize_items(ws, count: int) -> None:
    """Делает в таблице товаров ровно `count` строк.

    Бланк — это сетка объединённых ячеек, и обычная вставка строки её ломает:
    openpyxl сдвигает ячейки, но не объединения, высоты строк и область печати.
    Поэтому всё, что ниже строк товаров (итог), сдвигается вручную вместе с
    объединениями и высотами, а новые строки копируют вид последней.
    """
    shift = count - ITEMS_IN_TEMPLATE
    if shift == 0:
        return
    total_row = ITEMS_FIRST_ROW + ITEMS_IN_TEMPLATE
    last_item = total_row - 1
    max_row = ws.max_row
    max_col = ws.max_column

    # Образец новой строки: объединения и стиль последней строки товаров.
    pattern_merges = [(r.min_col, r.max_col) for r in ws.merged_cells.ranges
                      if r.min_row == last_item and r.max_row == last_item]
    pattern_height = ws.row_dimensions[last_item].height
    pattern_styles = {col: copy(ws.cell(last_item, col)._style)
                      for col in range(1, max_col + 1)}

    below = [r for r in ws.merged_cells.ranges if r.min_row >= total_row]
    if shift < 0:
        # Лишние строки товаров: их объединения уходят вместе с ними.
        dropped = [r for r in ws.merged_cells.ranges
                   if total_row + shift <= r.min_row < total_row]
        for r in dropped:
            ws.unmerge_cells(str(r))
    heights = {row: ws.row_dimensions[row].height for row in range(total_row, max_row + 1)}
    spans = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in below]
    for r in below:
        ws.unmerge_cells(str(r))

    ws.move_range(f"A{total_row}:{get_column_letter(max_col)}{max_row}", rows=shift)

    for min_row, min_col, max_r, max_c in spans:
        ws.merge_cells(start_row=min_row + shift, start_column=min_col,
                       end_row=max_r + shift, end_column=max_c)
    for row, height in heights.items():
        ws.row_dimensions[row + shift].height = height

    if shift > 0:
        for offset in range(shift):
            row = total_row + offset
            for col, style in pattern_styles.items():
                ws.cell(row, col)._style = copy(style)
            for min_col, max_c in pattern_merges:
                ws.merge_cells(start_row=row, start_column=min_col,
                               end_row=row, end_column=max_c)
            ws.row_dimensions[row].height = pattern_height

    if ws.print_area:
        area = ws.print_area.split("!")[-1].replace("$", "")
        match = re.fullmatch(r"([A-Z]+\d+):([A-Z]+)(\d+)", area)
        if match:
            ws.print_area = f"{match.group(1)}:{match.group(2)}{int(match.group(3)) + shift}"


def _fill_third_page(ws, act: TorgAct, s: TorgSettings) -> None:
    lines = wrap_lines(act.description, DESCRIPTION_LINES)
    for cell, text in zip(("AW20", "A21", "A22", "A23"), lines + [""] * 4):
        ws[cell] = text
    lines = wrap_lines(act.conclusion, CONCLUSION_LINES)
    for cell, text in zip(("Y24", "A25", "A26"), lines + [""] * 3):
        ws[cell] = text

    ws["AC30"] = s.chair_title
    ws["CA30"] = s.chair_name
    for row, index in ((33, 0), (35, 1), (37, 2)):
        title, name = s.members[index] if index < len(s.members) else ("", "")
        ws[f"AC{row}"] = title
        ws[f"CA{row}"] = name


CODES_SHEET = "стр4"
_CODES_HEADER = ("Номер строки", "Код номенклатуры", "Штрих-код", "Код маркировки",
                 "Брак", "Недостача", "Излишек")
_CODES_WIDTHS = (13.7, 18.3, 19.1, 97.3, 9.1, 12.9, 9.1)
# Управляющие символы в ячейку Excel не пишутся: разделитель групп (GS), который
# сканер ставит внутри кода, openpyxl отвергает целиком.
_CONTROL = re.compile(r"[\x00-\x1f\x7f]")


def sheet_code(code: str) -> str:
    """Код маркировки в виде для ячейки: без управляющих символов и пробелов по краям."""
    return _CONTROL.sub("", code or "").strip()


def _fill_codes_page(book: openpyxl.Workbook, act: TorgAct) -> None:
    """Четвёртая страница — коды маркировки позиций акта.

    Склад прикладывает её к акту: по каждому коду видно, какой экземпляр
    бракованный или лишний. Строка — одна заявка с кодом; в колонке «Номер
    строки» стоит номер позиции в таблице второй страницы, а количество — 1:
    код означает один экземпляр. У акта без кодов страницы нет вовсе — пустой
    лист поставщику ни к чему.
    """
    rows = [(number, item) for number, item in enumerate(act.items, 1)
            if sheet_code(item.mark_code)]
    if not rows:
        return
    ws = book.create_sheet(CODES_SHEET)
    ws.append(list(_CODES_HEADER))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for number, item in rows:
        # Номер и штрих-код — текстом: иначе Excel съест ведущие нули и покажет
        # 8.01733E+12. Количество стоит в колонке того, что случилось.
        ws.append([str(number), "", item.barcode, sheet_code(item.mark_code),
                   1 if item.qty_defect else None, 1 if item.qty_short else None,
                   1 if item.qty_surplus else None])
        row = ws.max_row
        for column in "ABCD":
            ws[f"{column}{row}"].number_format = "@"
        ws[f"A{row}"].alignment = Alignment(horizontal="center")
    for column, width in zip("ABCDEFG", _CODES_WIDTHS):
        ws.column_dimensions[column].width = width


def _clear_yellow(book: openpyxl.Workbook) -> None:
    """Жёлтое в образце — пометка «заполняет склад». В готовом акте её быть не
    должно: бланк уходит поставщику и печатается."""
    from openpyxl.styles import PatternFill

    nothing = PatternFill(fill_type=None)
    for ws in book:
        for row in ws.iter_rows():
            for cell in row:
                fill = cell.fill
                if fill and fill.fill_type and str(fill.fgColor.rgb) == "FFFFFF00":
                    cell.fill = nothing
