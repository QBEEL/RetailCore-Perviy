"""Сверка приехавшего товара с кодами из УПД.

Проверяется главное свойство сверки: один и тот же экземпляр опознаётся и
тогда, когда в документе он записан кодом идентификации, а сканер отдал его
вместе с криптохвостом и без разделителей. Если это свойство сломается, сверка
объявит недостачей всю коробку, а расхождения станут нормой — и им перестанут
верить.

Ни сети, ни сертификата здесь нет намеренно: сверка обходится без них, и тест
это тоже утверждает.
"""
from __future__ import annotations

import os
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.marking import reconcile, upd
from app.core.marking.codes import GS
from app.core.marking.reconcile import Reconciliation, Verdict

# Документ того же строения, что приходит из Диадока: версия 5.03, кодировка
# windows-1251, коды экземпляров — в «НомСредИдентТов» внутри «ДопСведТов».
DOCUMENT = """<?xml version="1.0" encoding="windows-1251"?>
<Файл ИдФайл="ON_NSCHFDOPPR_test" ВерсФорм="5.03">
  <Документ КНД="1115131" Функция="СЧФДОП">
    <СвСчФакт НомерДок="3349" ДатаДок="09.09.2026">
      <СвПрод>
        <ИдСв><СвЮЛУч НаимОрг="НЭЙТИВ ООО" ИННЮЛ="7724542709" КПП="774301001"/></ИдСв>
      </СвПрод>
      <СвПокуп>
        <ИдСв>
          <СвИП ИННФЛ="254009895745">
            <ФИО Фамилия="Саух" Имя="Мария" Отчество="Николаевна"/>
          </СвИП>
        </ИдСв>
      </СвПокуп>
    </СвСчФакт>
    <ТаблСчФакт>
      <СведТов НомСтр="1" НаимТов="Бюстгальтер, какао (M)" ОКЕИ_Тов="796"
               НаимЕдИзм="шт" КолТов="2.00" ЦенаТов="3503.27" СтТовУчНал="7497.00">
        <ДопСведТов ПрТовРаб="1" КодТов="04680962315457">
          <НомСредИдентТов>
            <КИЗ>0104680962315457215hsFUtaRDI3mf</КИЗ>
            <КИЗ>0104680962315457215X+iQRNpbsQ'z</КИЗ>
          </НомСредИдентТов>
        </ДопСведТов>
      </СведТов>
      <СведТов НомСтр="2" НаимТов="Трусы, чёрные (L)" ОКЕИ_Тов="796"
               НаимЕдИзм="шт" КолТов="1.00" ЦенаТов="1628.50" СтТовУчНал="1742.50">
        <ДопСведТов ПрТовРаб="1" КодТов="04680962314627">
          <НомСредИдентТов>
            <КИЗ>0104680962314627215lGhszhqF%9Rx</КИЗ>
          </НомСредИдентТов>
        </ДопСведТов>
      </СведТов>
    </ТаблСчФакт>
  </Документ>
</Файл>
"""

FIRST = "0104680962315457215hsFUtaRDI3mf"
SECOND = "0104680962315457215X+iQRNpbsQ'z"
THIRD = "0104680962314627215lGhszhqF%9Rx"

# Тот же первый экземпляр, каким его отдаёт сканер: криптохвост на месте,
# разделителя нет — поле ввода срезает неотображаемые знаки.
SCANNED = FIRST + "93Zx1q"


@pytest.fixture
def document(tmp_path) -> Path:
    path = tmp_path / "ON_NSCHFDOPPR_2BM-1.xml"
    path.write_bytes(DOCUMENT.encode("cp1251"))
    return path


# --- чтение документа -------------------------------------------------------------

def test_читает_упд_с_кодами(document):
    read = upd.read(str(document))

    assert read.number == "3349"
    assert read.date == "09.09.2026"
    assert read.seller == "НЭЙТИВ ООО"
    assert read.seller_inn == "7724542709"
    assert read.buyer == "Саух Мария Николаевна"
    assert len(read.lines) == 2
    assert len(read.marks) == 3

    first = read.lines[0]
    assert first.name == "Бюстгальтер, какао (M)"
    assert first.gtin == "04680962315457"
    assert first.quantity == 2.0
    assert [mark.value for mark in first.marks] == [FIRST, SECOND]


def test_читает_упд_из_архива_диадока(document, tmp_path):
    """Архив скачивают целиком, и распаковывать его ради одного файла незачем."""
    archive = tmp_path / "Диадок.zip"
    with zipfile.ZipFile(archive, "w") as pack:
        pack.write(document, "ON_NSCHFDOPPR_2BM-1.xml")
        pack.writestr("Протокол.pdf", b"%PDF-1.4")
        pack.writestr("УС/DP_UNISOOBSCH.xml", "<Файл><Документ/></Файл>")

    read = upd.read(str(archive))

    assert read.number == "3349"
    assert len(read.marks) == 3


def test_документ_без_кодов_отвергается_с_объяснением(tmp_path):
    """Молчаливая пустая сверка хуже отказа: непонятно, чего ждать дальше."""
    path = tmp_path / "пустой.xml"
    path.write_bytes(DOCUMENT.replace(
        "<КИЗ>0104680962315457215hsFUtaRDI3mf</КИЗ>", "")
        .replace("<КИЗ>0104680962315457215X+iQRNpbsQ'z</КИЗ>", "")
        .replace("<КИЗ>0104680962314627215lGhszhqF%9Rx</КИЗ>", "")
        .encode("cp1251"))

    with pytest.raises(upd.UpdProblem) as problem:
        upd.read(str(path))

    assert "нет кодов маркировки" in str(problem.value)


def test_нехватка_кодов_в_документе_видна_до_сканирования(tmp_path):
    """Товара две штуки, код один — расхождение самого документа."""
    path = tmp_path / "недобор.xml"
    path.write_bytes(
        DOCUMENT.replace("<КИЗ>0104680962315457215X+iQRNpbsQ'z</КИЗ>", "")
        .encode("cp1251"))

    read = upd.read(str(path))

    assert read.lines[0].shortage == 1
    assert "кодов не хватает" in read.summary


# --- сверка ------------------------------------------------------------------------

def test_код_со_сканера_узнаётся_несмотря_на_криптохвост(document):
    """Главное свойство сверки: этикетка и документ — про одну и ту же вещь."""
    session = Reconciliation(upd.read(str(document)))

    scan = session.scan(SCANNED)

    assert scan.verdict is Verdict.MATCHED
    assert scan.line is not None and scan.line.number == "1"
    assert session.done == 1
    assert session.left == 2


def test_разделители_и_признак_символики_не_мешают(document):
    """Сканеры ставят ]d2 впереди и GS между полями — код от этого не другой."""
    session = Reconciliation(upd.read(str(document)))

    scan = session.scan(f"]d2{FIRST}{GS}93Zx1q")

    assert scan.verdict is Verdict.MATCHED
    assert session.done == 1


def test_повтор_не_засчитывается_вторым_экземпляром(document):
    """Дважды пикнутая вещь — это одна вещь, а не закрытая строка."""
    session = Reconciliation(upd.read(str(document)))

    session.scan(FIRST)
    repeat = session.scan(FIRST)

    assert repeat.verdict is Verdict.REPEAT
    assert session.done == 1
    assert session.repeats == 1


def test_чужой_экземпляр_того_же_товара_отличается_от_чужого_товара(document):
    """Разговор с поставщиком в этих случаях разный, и подсказка тоже."""
    session = Reconciliation(upd.read(str(document)))

    same_goods = session.scan("0104680962315457215ZZZZZZZZZZZZ")
    other_goods = session.scan("0104607091398595215AbcDefGhiJkL")

    assert same_goods.verdict is Verdict.UNKNOWN
    assert "строки 1" in same_goods.note
    assert other_goods.verdict is Verdict.UNKNOWN
    assert other_goods.note == "такого товара в документе нет"
    assert len(session.extra) == 2


def test_нечитаемый_код_не_путается_с_лишним(document):
    session = Reconciliation(upd.read(str(document)))

    scan = session.scan("МУСОР")

    assert scan.verdict is Verdict.BROKEN
    assert scan.note


def test_отмена_возвращает_строку_в_работу(document):
    """Сканер стреляет по соседней коробке, и без отмены выход один — заново."""
    session = Reconciliation(upd.read(str(document)))
    session.scan(FIRST)

    undone = session.undo()

    assert undone is not None and undone.verdict is Verdict.MATCHED
    assert session.done == 0
    assert session.left == 3
    assert session.undo() is None


def test_отмена_снимает_и_лишний_код(document):
    session = Reconciliation(upd.read(str(document)))
    session.scan("0104680962315457215ZZZZZZZZZZZZ")

    session.undo()

    assert session.extra == []


def test_сверка_сходится_когда_всё_найдено(document):
    session = Reconciliation(upd.read(str(document)))

    for code in (SCANNED, SECOND, THIRD):
        session.scan(code)

    assert session.complete
    assert session.left == 0
    assert session.missing == []
    assert "расхождений нет" in session.summary


def test_лишний_код_не_даёт_сверке_сойтись(document):
    """Всё из документа нашлось, но приехало ещё что-то — это не «сошлось»."""
    session = Reconciliation(upd.read(str(document)))
    for code in (FIRST, SECOND, THIRD):
        session.scan(code)

    session.scan("0104680962315457215ZZZZZZZZZZZZ")

    assert not session.complete


def test_ход_по_строкам_считается_отдельно(document):
    session = Reconciliation(upd.read(str(document)))
    session.scan(FIRST)

    first, second = session.progress

    assert (first.expected, first.scanned, first.left) == (2, 1, 1)
    assert first.state == "Не хватает 1"
    assert second.state == "Не начато"


def test_код_записанный_в_документе_дважды_замечается(tmp_path):
    """Сверить такую строку до конца нельзя: второй такой вещи в коробке нет."""
    path = tmp_path / "дубль.xml"
    path.write_bytes(DOCUMENT.replace(
        "<КИЗ>0104680962315457215X+iQRNpbsQ'z</КИЗ>",
        "<КИЗ>0104680962315457215hsFUtaRDI3mf</КИЗ>").encode("cp1251"))

    session = Reconciliation(upd.read(str(path)))

    assert session.repeated_in_document == [FIRST]
    assert session.total == 2


def test_начать_заново_забывает_сверенное_но_не_документ(document):
    session = Reconciliation(upd.read(str(document)))
    session.scan(FIRST)

    session.reset()

    assert session.done == 0
    assert session.total == 3
    assert session.extra == []


# --- акт сверки ---------------------------------------------------------------------

def test_акт_сверки_записывается(document, tmp_path):
    session = Reconciliation(upd.read(str(document)))
    session.scan(SCANNED)
    session.scan("0104680962315457215ZZZZZZZZZZZZ")

    saved = reconcile.save_report(session, str(tmp_path / "акт.xlsx"))

    assert os.path.exists(saved)
    from openpyxl import load_workbook
    book = load_workbook(saved)
    assert book.sheetnames == ["Сверка", "Коды"]
    # В листе кодов каждый код документа плюс каждый лишний скан.
    assert book["Коды"].max_row == 1 + 3 + 1


def test_имя_акта_по_умолчанию_называет_документ(document, tmp_path):
    session = Reconciliation(upd.read(str(document)))

    name = os.path.basename(reconcile.default_name(session, str(tmp_path)))

    assert name == "Сверка УПД 3349 от 09.09.2026.xlsx"


# --- раздел на вкладке ----------------------------------------------------------------

@pytest.fixture(scope="module")
def application():
    pytest.importorskip("PySide6")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


@pytest.fixture
def page(application, tmp_path, monkeypatch, document):
    """Вкладка с подменённым окном выбора файла. Ни сети, ни КриптоПро."""
    from app.core.settings import AppSettings
    from app.ui import marking_page

    class Dialog:
        @staticmethod
        def getOpenFileName(*_args, **_kwargs):
            return str(document), ""

        @staticmethod
        def getSaveFileName(*_args, **_kwargs):
            return str(tmp_path / "акт.xlsx"), ""

    monkeypatch.setattr(marking_page, "QFileDialog", Dialog)
    settings = AppSettings(_path=str(tmp_path / "settings.json"))
    # Без restore(): он читает сертификаты и журнал, а сверке они не нужны —
    # в этом и смысл раздела.
    return marking_page.MarkingPage(settings, lambda *_: None)


def test_загруженный_упд_открывает_сверку_и_ставит_курсор_в_поле(page):
    """Сканер печатает туда, где курсор, — иначе первый код уедет в никуда."""
    from app.ui.marking_page import RECONCILE_TAB

    page.load_upd()

    assert page.tabs.currentIndex() == RECONCILE_TAB
    assert page.scan_edit.isEnabled()
    assert page.tile_upd_codes._value.text() == "3"
    assert "НЭЙТИВ ООО" in page.upd_hint.text()


def test_скан_закрывает_экземпляр_и_отвечает_зелёным(page):
    from app.ui.theme import Palette

    page.load_upd()
    page.scan_edit.setText(SCANNED)
    page.accept_scan()

    assert "Есть в документе" in page.verdict_label.text()
    assert Palette.SUCCESS in page.verdict_label.styleSheet()
    assert page.tile_scanned._value.text() == "1"
    assert page.tile_left._value.text() == "2"
    # Поле очищается всегда: оставленный код сканер допишет своим.
    assert page.scan_edit.text() == ""


def test_лишний_код_показывается_отдельно_и_только_когда_он_есть(page):
    page.load_upd()
    # isHidden, а не isVisible: окно в тесте не показано, и видимым не
    # считается ни один его потомок.
    assert page.extra_card.isHidden()

    page.scan_edit.setText("0104680962315457215ZZZZZZZZZZZZ")
    page.accept_scan()

    assert not page.extra_card.isHidden()
    assert page.extra.rowCount() == 1
    assert page.tile_extra._value.text() == "1"


def test_отмена_скана_возвращает_позицию_в_работу(page):
    page.load_upd()
    page.scan_edit.setText(SCANNED)
    page.accept_scan()

    page.undo_scan()

    assert page.tile_scanned._value.text() == "0"
    assert page.tile_left._value.text() == "3"


def test_только_незакрытые_убирает_сверенные_строки(page):
    page.load_upd()
    assert page.progress.rowCount() == 2

    page.scan_edit.setText(THIRD)
    page.accept_scan()
    page.open_only_box.setChecked(True)

    assert page.progress.rowCount() == 1
    assert page.progress.item(0, 0).text() == "Бюстгальтер, какао (M)"


def test_коды_документа_уходят_на_вкладку_проверки(page):
    """Сверка и проверка отвечают на разные вопросы, и один без другого неполон."""
    from app.ui.marking_page import CHECK_TAB

    page.load_upd()
    page.codes_to_check()

    assert page.tabs.currentIndex() == CHECK_TAB
    assert page.codes_edit.toPlainText().splitlines() == [FIRST, SECOND, THIRD]


def test_f5_на_пустой_сверке_предлагает_выбрать_документ(page):
    """Выбрать УПД — первое действие раздела, им F5 и занят."""
    from app.ui.marking_page import RECONCILE_TAB

    page.tabs.setCurrentIndex(RECONCILE_TAB)
    page.run_current()

    assert page.tile_upd_codes._value.text() == "3"


def test_ctrl_s_на_сверке_сохраняет_акт(page, application, tmp_path):
    from PySide6.QtCore import QThreadPool

    page.load_upd()
    page.scan_edit.setText(FIRST)
    page.accept_scan()

    page.save_current()
    for _ in range(6):
        QThreadPool.globalInstance().waitForDone(3000)
        application.processEvents()

    assert (tmp_path / "акт.xlsx").exists()


def test_без_документа_сканировать_нечем(page):
    assert not page.scan_edit.isEnabled()
    assert not page.upd_report_button.isEnabled()
    assert "Документ не загружен" in page.upd_hint.text()
