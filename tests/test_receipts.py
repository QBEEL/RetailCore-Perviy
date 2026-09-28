"""Аналитика по чекам: разбор выгрузки 1С, выводы, Excel и вкладка.

Строки собираются здесь же, в том виде, в каком их печатает 1С: параметры над
шапкой, шапка с объединёнными ячейками, строки товаров и «Итого» в конце.
"""
from __future__ import annotations

import os
import sys
from datetime import date, datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core import receipts
from app.core.receipts.parse import parse_rows

HEAD = ["Документ", None, None, "Вид документа", "Дата", "Магазин", None, "Гость",
        "Номенклатура с характеристикой", "Категория товара", "Бренд", "Цена", "Акция",
        "% авт.", "Продавец", "Год", "Месяц", "Количество", "Себестоимость",
        "Сумма авт.", "Скидка бонусами", "Сумма", "Начислено бонусов"]


def _line(number: str, when: str, store: str, guest: str, item: str, brand: str,
          price: float, qty: float, cost: float, amount: float, *, kind: str = "Чек",
          seller: str = "Сивкова Маргарита", auto: float | None = None,
          bonus: float | None = None, promo: str | None = None) -> list:
    day = when.split()[0]
    title = ("Чек ККМ на возврат " if "возврат" in kind else "Чек ККМ ") + f"{number} от {when}"
    return [title, None, None, kind, day, store, None, guest, item, "Личная гигиена",
            brand, price, promo, None, seller, 2026, int(day[3:5]), qty, cost, auto, bonus,
            amount, None]


def _rows(lines: list[list], *, period: str = "Период: 01.08.2026 - 31.08.2026",
          total: list | None = None) -> list[list]:
    params = [
        [None] * 23,
        ["Параметры:", None, period] + [None] * 20,
        [None, None, "Подразделение: Универмаг; Уссурийск"] + [None] * 20,
        ["Отбор:", None, "Номенклатура В группе из списка \"Biorepair\""] + [None] * 20,
        [None] * 23,
    ]
    if total is None:
        sums = [sum((line[i] or 0) for line in lines) for i in (17, 18, 19, 20, 21, 22)]
        total = ["Итого"] + [None] * 16 + sums
    return [*params, HEAD, *lines, total]


LINES = [
    _line("IP00-1", "01.08.2026 9:34:26", "Универмаг", "Иванова Анна", "Паста BlanX, 75 мл",
          "BLANX", 750, 1, 445, 750),
    _line("IP00-2", "01.08.2026 13:05:00", "Универмаг", "Иванова Анна", "Паста BlanX, 75 мл",
          "BLANX", 750, 2, 890, 1350, auto=150, promo="Скидка_BLANX_20%"),
    _line("IP00-2", "01.08.2026 13:05:00", "Универмаг", "Иванова Анна", "Свеча Banka Home",
          "BANKA HOME", 900, 1, 430, 600, bonus=300),
    _line("IP00-3", "02.08.2026 18:10:00", "Уссурийск", "Розничный покупатель",
          "Паста BlanX, 75 мл", "BLANX", 750, 1, 445, 750, seller="Самостоятельная покупка"),
    _line("IP00-9", "03.08.2026 19:07:04", "Уссурийск", "Петров Пётр", "Паста BlanX, 75 мл",
          "BLANX", -750, -1, -445, -750, kind="Чек на возврат"),
]


# --- разбор -----------------------------------------------------------------------------

def test_период_и_параметры_из_шапки():
    data = parse_rows(_rows(LINES), source="чеки.xlsx")
    assert (data.start, data.end) == (date(2026, 8, 1), date(2026, 8, 31))
    assert data.days == 31
    assert data.units == "Универмаг; Уссурийск"
    assert "Biorepair" in data.selection
    assert data.missing == [] and data.balanced is True
    assert len(data.lines) == 5


def test_время_чека_из_номера_документа():
    data = parse_rows(_rows(LINES))
    assert data.lines[0].moment == datetime(2026, 8, 1, 9, 34, 26)


def test_период_по_датам_если_в_шапке_нет():
    data = parse_rows(_rows(LINES, period="Отчёт без периода"))
    assert (data.start, data.end) == (date(2026, 8, 1), date(2026, 8, 3))
    assert any("не найден" in text for text in data.warnings)


def test_другой_формат_периода():
    data = parse_rows(_rows(LINES, period="Период: с 01.08.2026 по 15.08.2026"))
    assert data.end == date(2026, 8, 15)


def test_чужой_файл_объясняет_чего_нет():
    rows = [["Номенклатура", "Расход", "Конечный остаток"], ["Паста", 1, 2]]
    with pytest.raises(receipts.ParseError, match="не хватает — Документ, Дата, Магазин"):
        parse_rows(rows, source="ведомость.xlsx")


def test_несошедшийся_итог_виден():
    total = ["Итого"] + [None] * 16 + [5, 1000, 150, 300, 99999, 0]
    data = parse_rows(_rows(LINES, total=total))
    assert data.balanced is False
    assert any("«Сумма» не сходится" in text for text in data.warnings)


def test_колонки_в_другом_порядке_и_без_необязательных():
    head = ["Магазин", "Дата", "Документ", "Номенклатура", "Сумма", "Количество"]
    rows = [head, ["Универмаг", "01.08.2026", "Чек ККМ IP00-1 от 01.08.2026 10:00:00",
                   "Паста", 750, 1]]
    data = parse_rows(rows)
    assert data.lines[0].amount == 750 and data.lines[0].store == "Универмаг"
    assert "cost" in data.missing and not data.has("seller")


# --- выводы -----------------------------------------------------------------------------

def test_выручка_сходится_с_файлом_а_средний_чек_без_возвратов():
    report = receipts.build(parse_rows(_rows(LINES)))
    total = report.total
    assert total.revenue == pytest.approx(750 + 1350 + 600 + 750 - 750)
    assert total.checks == 3, "возврат — не чек продажи"
    assert total.average_check == pytest.approx((750 + 1350 + 600 + 750) / 3)
    assert total.returned == pytest.approx(750)
    assert total.profit == pytest.approx(total.revenue - total.cost)


def test_скидка_доля_от_суммы_до_скидки():
    report = receipts.build(parse_rows(_rows(LINES)))
    assert report.total.discount == pytest.approx(450)
    assert report.total.discount_share == pytest.approx(450 / (2700 + 450))


def test_гости_и_покупки_без_продавца():
    report = receipts.build(parse_rows(_rows(LINES)))
    # «Розничный покупатель» — не гость; вернувшийся только с возвратом — тоже.
    assert report.guests == 1 and report.repeat_guests == 1
    assert report.self_service.revenue == pytest.approx(750)
    assert all(group.name != "Самостоятельная покупка" for group in report.sellers)


def test_срезы_по_магазинам_брендам_часам_и_акциям():
    report = receipts.build(parse_rows(_rows(LINES)))
    assert [group.name for group in report.stores] == ["Универмаг", "Уссурийск"]
    assert report.brands[0].name == "BLANX"
    assert dict((hour, group.checks) for hour, group in report.hours) == {9: 1, 13: 1, 18: 1}
    assert [group.name for group in report.promos] == ["Скидка_BLANX_20%"]
    assert len(report.returns) == 1
    item = next(group for group in report.items if group.name == "Паста BlanX, 75 мл")
    assert item.brand == "BLANX" and len(item.stores) == 2


def test_чек_в_двух_файлах_учитывается_один_раз():
    first = parse_rows(_rows(LINES[:3]), source="август-1.xlsx")
    second = parse_rows(_rows(LINES[1:4], period="Период: 01.08.2026 - 31.08.2026"),
                        source="август-2.xlsx")
    merged = receipts.merge([first, second])
    assert len(merged.lines) == 4
    assert any("перекрываются" in text for text in merged.warnings)


def test_тот_же_файл_дважды_не_удваивает():
    part = parse_rows(_rows(LINES), source="чеки.xlsx")
    merged = receipts.merge([part, parse_rows(_rows(LINES), source="чеки.xlsx")])
    assert len(merged.lines) == len(LINES)


# --- Excel ------------------------------------------------------------------------------

def test_excel_с_листами_и_числами(tmp_path):
    from openpyxl import load_workbook

    report = receipts.build(parse_rows(_rows(LINES)))
    path = receipts.save(report, str(tmp_path / "разбор.xlsx"))
    book = load_workbook(path)
    assert book.sheetnames == ["Главное", "Магазины", "Бренды", "Категории", "Товары",
                               "Продавцы", "По дням", "По часам", "Акции", "Возвраты"]
    main = {row[0]: row[1] for row in book["Главное"].iter_rows(values_only=True) if row[0]}
    assert main["Период"] == "01.08.2026 — 31.08.2026"
    assert main["Выручка"] == pytest.approx(2700)
    assert book["Магазины"]["B2"].number_format == "#,##0.00"
    assert "01.08.2026–31.08.2026" in receipts.default_name(str(tmp_path), report)


# --- вкладка ----------------------------------------------------------------------------

@pytest.fixture(scope="module")
def application():
    pytest.importorskip("PySide6")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_вкладка_показывает_период_и_срезы(application):
    from app.core.settings import AppSettings
    from app.ui.widgets.receipts_tab import ReceiptsTab

    tab = ReceiptsTab(AppSettings(), lambda *_: None)
    tab.show_report(receipts.build(parse_rows(_rows(LINES))))
    assert tab.period_title.text() == "Период: 01.08.2026 — 31.08.2026 · 31 дн."
    stores = tab.tables["stores"]
    assert stores.rowCount() == 2 and stores.item(0, 0).text() == "Универмаг"
    assert tab.save_button.isEnabled()
    assert tab.views.tabText(tab.views.indexOf(tab.tables["returns"])) == "Возвраты (1)"


def test_вкладка_есть_в_отчётности(application):
    from app.core.settings import AppSettings
    from app.ui.reports_page import ReportsPage

    page = ReportsPage(AppSettings(), lambda *_: None)
    titles = [page.tabs.tabText(i) for i in range(page.tabs.count())]
    assert titles[-1] == "Аналитика по чекам"
    page.tabs.setCurrentIndex(2)
    assert "Аналитика по чекам" in page.subtitle.text()
