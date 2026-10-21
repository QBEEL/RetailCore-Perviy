"""Торг-2: акт ТОРГ-2 из программы, справочник брендов, окна и страница."""
from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PySide6.QtWidgets import QAbstractItemDelegate, QApplication, QMessageBox  # noqa: E402

from app.core import appdata  # noqa: E402
from app.core.claims import (  # noqa: E402
    Brand,
    Claim,
    ClaimFile,
    TorgAct,
    TorgItem,
    TorgSettings,
    brand_key,
    has_brand,
    remote,
    torg,
)
from app.core.payments import transport  # noqa: E402
from app.core.settings import AppSettings  # noqa: E402
from app.ui import claims_page  # noqa: E402
from app.ui.claims_page import ClaimsPage  # noqa: E402
from app.ui.widgets import claim_dialogs, torg_dialogs  # noqa: E402
from app.ui.widgets.claim_dialogs import ClaimDialog  # noqa: E402
from app.ui.widgets.torg_dialogs import (  # noqa: E402
    COL_DEFECT,
    COL_DOC,
    COL_FACT,
    COL_PRICE,
    COL_SHORT,
    BrandsDialog,
    TorgDialog,
    TorgSettingsDialog,
    parse_number,
)

SETTINGS = TorgSettings(
    organization="ИП Тестов Тест Тестович, 690000, г. Тест, ул. Тестовая 1",
    director_title="ИП Тестов Т. Т.", director_name="Тестов Т. Т.",
    place="690000, г. Тест, ул. Склад 2", chair_name="Иванова И. И.",
    members=[("товаровед", "Петров П."), ("товаровед", "Сидоров С.")])


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def clean_session():
    transport.session.clear()
    yield
    transport.session.clear()


def item(number: int = 1, **over) -> TorgItem:
    data = dict(claim_id=number, name=f"Товар {number} Brand 50 мл",
                barcode=f"46000000000{number:02d}", price=375.0, qty_doc=36,
                qty_short=0, qty_defect=11)
    data.update(over)
    return TorgItem(**data)


def act(*items: TorgItem, **over) -> TorgAct:
    data = dict(
        number=63, act_date=date(2026, 10, 20), receipt_date=date(2026, 10, 19),
        supplier='ООО "Тест Поставка" 109380, Москва, ул Тестовая, д. 4',
        contract_number="ТП-1/2023", contract_date=date(2023, 1, 24),
        invoice_number="4455", invoice_date=date(2026, 9, 28),
        delivery="Транспортная компания", description="Замятие упаковки.",
        conclusion="выставить УКД на бракованную продукцию",
        items=list(items) or [item()])
    data.update(over)
    return TorgAct(**data)


def built(a: TorgAct, settings: TorgSettings = SETTINGS) -> openpyxl.Workbook:
    return torg.build_workbook(a, settings)


def all_strings(book: openpyxl.Workbook) -> list[str]:
    return [str(c.value) for ws in book for row in ws.iter_rows() for c in row
            if c.value is not None]


# --- расчёт --------------------------------------------------------------------

class TestAmounts:
    def test_брак_входит_в_поступившее(self) -> None:
        # Как в образце склада: 36 по документам, 36 фактически, 11 из них брак.
        row = item(qty_doc=36, qty_defect=11, qty_short=0)
        assert row.qty_fact == 36 and row.sum_defect == 4125.0 and row.sum_fact == 13500.0

    def test_недостача_уменьшает_поступившее(self) -> None:
        row = item(qty_doc=36, qty_defect=0, qty_short=3)
        assert row.qty_fact == 33 and row.sum_short == 1125.0 and row.sum_fact == 12375.0

    def test_итоги_по_строкам(self) -> None:
        total = act(item(1), item(2, qty_defect=0, qty_short=3)).totals()
        assert total["qty_doc"] == 72 and total["sum_doc"] == 27000.0
        assert total["qty_defect"] == 11 and total["qty_short"] == 3
        assert total["sum_defect"] == 4125.0 and total["sum_short"] == 1125.0

    def test_деньги_округляются_до_копеек(self) -> None:
        assert item(price=10.005, qty_doc=3, qty_defect=1).sum_defect == 10.01


class TestProblems:
    def test_готовый_акт_без_замечаний(self) -> None:
        assert act().problems() == []

    @pytest.mark.parametrize("change, text", [
        ({"supplier": " "}, "поставщика"),
        ({"invoice_number": ""}, "номер счёта-фактуры"),
        ({"invoice_date": None}, "дату счёта-фактуры"),
        ({"items": []}, "нет товаров"),
    ])
    def test_шапка(self, change, text) -> None:
        assert any(text in p for p in act(**change).problems())

    @pytest.mark.parametrize("row, text", [
        (item(qty_doc=0), "количество по документам"),
        (item(price=0), "цену"),
        (item(qty_short=40, qty_defect=0), "недостача больше"),
        (item(qty_defect=50), "брака больше"),
        (item(qty_defect=0, qty_short=0), "недостачу, брак или излишек"),
        (item(qty_defect=-1), "отрицательным"),
    ])
    def test_строки(self, row, text) -> None:
        assert any(text in p for p in act(row).problems())

    def test_слишком_длинное_описание_не_пропускается(self) -> None:
        assert any("Описание" in p for p in act(description="слово " * 200).problems())

    def test_пустые_реквизиты_названы(self) -> None:
        assert TorgSettings().problems() == [
            "организация", "руководитель", "место приёмки", "председатель комиссии"]
        assert SETTINGS.problems() == []

    def test_подсказка_текста_по_содержимому(self) -> None:
        assert "бракованную" in act(item(qty_short=0, qty_defect=2)).suggested_text()[1]
        assert "недопоставленную продукцию" in act(
            item(qty_short=2, qty_defect=0)).suggested_text()[1]
        assert "и бракованную" in act(item(1), item(2, qty_defect=0, qty_short=1)
                                      ).suggested_text()[1]


class TestText:
    def test_перенос_по_словам(self) -> None:
        assert torg.wrap_lines("aaa bbb ccc ddd", (7, 7)) == ["aaa bbb", "ccc ddd"]

    def test_явный_перенос_строки_уважается(self) -> None:
        assert torg.wrap_lines("первая\nвторая", (50, 50)) == ["первая", "вторая"]

    def test_что_не_помещается_определяется(self) -> None:
        assert torg.overflow("a " * 100, (10, 10)) is True
        assert torg.overflow("короткий текст", torg.DESCRIPTION_LINES) is False
        # Слово длиннее строки — тоже не помещается, а не разрезается молча.
        assert torg.overflow("х" * 100, (10, 10)) is True

    def test_пустой_текст_помещается(self) -> None:
        assert torg.overflow("", (5,)) is False and torg.wrap_lines("", (5,)) == []


class TestNames:
    def test_имя_поставщика_без_адреса(self) -> None:
        assert torg.short_supplier('ООО "Ромашка" 109380, Москва') == "ООО Ромашка"
        assert torg.short_supplier("ИП Иванов, г. Москва") == "ИП Иванов"

    def test_имя_файла(self) -> None:
        assert torg.file_name(act()) == "ТОРГ-2 №63 ООО Тест Поставка 20-10-2026.xlsx"

    def test_имя_файла_без_запрещённых_знаков(self) -> None:
        name = torg.file_name(act(supplier='ООО "А/Б:В" 123456, адрес'))
        assert not set('/:"<>|?*') & set(name)

    def test_оборот_данных_акта_через_сервер(self) -> None:
        a = act(item(1), item(2, qty_defect=0, qty_short=2))
        back = TorgAct.from_parts(id=5, number=a.number, act_date=a.act_date,
                                  supplier=a.supplier, payload=a.payload())
        assert back.payload() == a.payload() and back.id == 5
        assert back.items[1].qty_short == 2

    def test_акт_по_заявкам(self) -> None:
        claims = [Claim(id=1, kind="shortage", product="Крем", spec="50 мл",
                        barcode="111", quantity=3.0),
                  Claim(id=2, kind="defect", brand="Rink", barcode="222", quantity=2.0)]

        result = torg.act_from_claims(claims, 7, date(2026, 10, 20))

        assert result.number == 7 and result.claim_ids == [1, 2]
        first, second = result.items
        assert first.name == "Крем 50 мл" and first.qty_short == 3 and first.qty_defect == 0
        assert second.name == "Rink" and second.qty_defect == 2 and second.qty_short == 0


# --- бланк ---------------------------------------------------------------------

class TestWorkbook:
    def test_в_шаблоне_пусты_все_поля_которые_заполняет_программа(self) -> None:
        """Шаблон сделан из заполненного акта склада. Всё, что склад вписывал, из
        него убрано, иначе чужие данные попали бы в каждый новый акт."""
        from tools import make_torg_template as tool

        book = openpyxl.load_workbook(torg.template_path())
        addresses = [("стр1", a) for a in tool.FIRST] + [("стр2", a) for a in tool.SECOND]
        addresses += [("стр2", f"{c}{r}") for r in tool.SECOND_ROWS for c in tool.SECOND_COLUMNS]
        addresses += [("стр3", a) for a in tool.THIRD]
        for sheet, address in addresses:
            assert book[sheet][address].value in (None, ""), f"{sheet}!{address}"

    def test_в_шаблоне_нет_реквизитов_организаций(self) -> None:
        """Репозиторий публичный: в шаблоне только текст формы."""
        import re

        book = openpyxl.load_workbook(torg.template_path())
        for text in all_strings(book):
            assert not re.search(r"\b(ИП|ООО|АО|ОАО)\b", text), text
            # Длинные числа — ИНН, телефоны, счета; единственное допустимое — код формы.
            for number in re.findall(r"[0-9]{6,}", text):
                assert number == "0330202", text

    def test_в_шаблоне_нет_имён_авторов(self) -> None:
        book = openpyxl.load_workbook(torg.template_path())
        assert book.properties.creator == "RetailCore"
        assert not book.properties.lastModifiedBy

    def test_в_акте_свойства_файла_чистые(self) -> None:
        book = built(act())
        assert book.properties.creator == "RetailCore" and not book.properties.lastModifiedBy

    def test_нет_жёлтой_заливки(self) -> None:
        book = built(act())
        for ws in book:
            for row in ws.iter_rows():
                for cell in row:
                    fill = cell.fill
                    assert not (fill.fill_type and str(fill.fgColor.rgb) == "FFFFFF00"), cell.coordinate

    def test_первая_страница(self) -> None:
        ws = built(act())["стр1"]
        assert ws["AY17"].value == "63" and ws["CI11"].value == "63"
        assert (ws["CI12"].value, ws["CP12"].value, ws["CW12"].value) == ("20", "10", "2026")
        assert ws["CI20"].value == "Октября" and ws["BS24"].value == "октября"
        assert ws["A6"].value == SETTINGS.organization
        assert ws["CK18"].value == "Тестов Т. Т." and ws["Z23"].value == SETTINGS.place
        assert ws["U33"].value == ws["N37"].value == act().supplier
        assert ws["AL25"].value == "счет-фактура 4455 от 28.09.2026 г."
        assert ws["AT42"].value == "ТП-1/2023"
        assert (ws["BN42"].value, ws["BU42"].value, ws["CG42"].value) == ("24", "01", "2023")
        assert (ws["AO43"].value, ws["AV43"].value, ws["BH43"].value) == ("28", "09", "2026")
        assert ws["T48"].value == "Транспортная компания"
        assert ws["AQ60"].value == ws["BP60"].value == ws["CT60"].value == "19.10.2026"

    def test_строка_товара(self) -> None:
        ws = built(act())["стр2"]
        assert ws["A40"].value == "Товар 1 Brand 50 мл" and ws["AL40"].value == "шт"
        assert ws["BS40"].value == ws["DR40"].value == "4600000000001"
        assert (ws["BY40"].value, ws["CD40"].value, ws["CQ40"].value) == (36, 375.0, 13500.0)
        assert (ws["DV40"].value, ws["EF40"].value) == (36, 13500.0)
        assert (ws["EM40"].value, ws["ET40"].value) == (11, 4125.0)
        # Недостачи нет — колонки пустые, а не нули.
        assert ws["FO40"].value is None and ws["FW40"].value is None

    def test_штрих_код_хранится_текстом(self) -> None:
        ws = built(act())["стр2"]
        assert ws["BS40"].number_format == "@" and isinstance(ws["BS40"].value, str)

    def test_недостача(self) -> None:
        ws = built(act(item(qty_doc=36, qty_defect=0, qty_short=3)))["стр2"]
        assert ws["DV40"].value == 33 and ws["EF40"].value == 12375.0
        assert (ws["FO40"].value, ws["FW40"].value) == (3, 1125.0)
        assert ws["EM40"].value is None

    @pytest.mark.parametrize("count", [1, 2, 3, 6])
    def test_итоговая_строка_уезжает_за_товарами(self, count) -> None:
        a = act(*[item(i) for i in range(1, count + 1)])
        ws = built(a)["стр2"]
        total_row = 40 + count
        assert ws[f"BS{total_row}"].value == "Итого"
        assert ws[f"BY{total_row}"].value == 36 * count
        assert ws[f"EM{total_row}"].value == 11 * count
        assert ws[f"CQ{total_row}"].value == 13500.0 * count
        # После товаров — ничего лишнего, кроме итога.
        assert ws[f"A{40 + count - 1}"].value == f"Товар {count} Brand 50 мл"
        if count < 3:
            assert ws["A42"].value is None

    @pytest.mark.parametrize("count", [1, 4])
    def test_объединения_строк_товаров_целы(self, count) -> None:
        ws = built(act(*[item(i) for i in range(1, count + 1)]))["стр2"]
        anchors = {(r.min_row, r.min_col) for r in ws.merged_cells.ranges}
        for row in range(40, 40 + count):
            # Каждая строка товара — те же объединения, что в образце: название
            # на 27 колонок, штрих-код и т.д.
            assert (row, 1) in anchors and (row, 71) in anchors
        # Итоговая строка сохранила свои объединения под товарами.
        assert (40 + count, 71) in anchors
        assert (40 + count + count, 1) not in anchors or count == 0

    def test_область_печати_растёт_с_товарами(self) -> None:
        small = built(act(item(1), item(2)))["стр2"].print_area
        big = built(act(*[item(i) for i in range(1, 6)]))["стр2"].print_area
        one = built(act(item(1)))["стр2"].print_area
        assert one.endswith("41") and small.endswith("42") and big.endswith("45")

    def test_высоты_строк_товаров_одинаковы(self) -> None:
        ws = built(act(*[item(i) for i in range(1, 6)]))["стр2"]
        heights = {ws.row_dimensions[row].height for row in range(40, 45)}
        assert len(heights) <= 2  # 45.0 и 43.5 из образца

    def test_дата_вскрытия_тары(self) -> None:
        ws = built(act())["стр2"]
        assert (ws["EB20"].value, ws["EI20"].value, ws["FA20"].value) == ("19", "10", "2026")

    def test_третья_страница(self) -> None:
        ws = built(act())["стр3"]
        assert ws["AW20"].value == "Замятие упаковки."
        assert ws["Y24"].value == "выставить УКД на бракованную продукцию"
        assert ws["CA30"].value == "Иванова И. И." and ws["AC30"].value == "Руководитель склада"
        assert (ws["AC33"].value, ws["CA33"].value) == ("товаровед", "Петров П.")
        assert ws["CA35"].value == "Сидоров С."
        # Третьего члена комиссии нет — строка пустая, а не чужая.
        assert ws["CA37"].value in (None, "") and ws["AC37"].value in (None, "")

    def test_длинное_описание_идёт_на_следующие_строки(self) -> None:
        text = "Товар ненадлежащего качества, обнаружено при приёмке товара на складе. " * 2
        ws = built(act(description=text))["стр3"]
        assert ws["AW20"].value and ws["A21"].value
        assert " ".join(filter(None, [ws["AW20"].value, ws["A21"].value, ws["A22"].value])
                        ).split() == text.split()

    def test_шаблон_не_меняется_при_сборке(self) -> None:
        before = Path(torg.template_path()).read_bytes()
        built(act(*[item(i) for i in range(1, 5)]))
        assert Path(torg.template_path()).read_bytes() == before


class TestSave:
    def test_файл_записывается_и_читается(self, tmp_path) -> None:
        path = torg.save(act(), SETTINGS, tmp_path / "акт.xlsx")

        assert path.exists() and not list(tmp_path.glob("*.part"))
        assert openpyxl.load_workbook(path)["стр1"]["AY17"].value == "63"

    def test_с_замечаниями_файл_не_создаётся(self, tmp_path) -> None:
        with pytest.raises(ValueError, match="поставщика"):
            torg.save(act(supplier=""), SETTINGS, tmp_path / "акт.xlsx")
        with pytest.raises(ValueError, match="организация"):
            torg.save(act(), TorgSettings(), tmp_path / "акт.xlsx")
        assert not list(tmp_path.iterdir())


# --- бренды --------------------------------------------------------------------

class TestBrands:
    def test_сравнение_без_регистра_пробелов_и_ё(self) -> None:
        assert brand_key("  Biorepair  ") == "biorepair"
        assert brand_key("ЁЖИК   Жук") == "ежик жук"

    def test_наличие_в_списке(self) -> None:
        brands = [Brand(1, "BIOREPAIR"), Brand(2, "Рэд Булл")]
        assert has_brand(brands, "biorepair") and has_brand(brands, " рэд  булл ")
        assert not has_brand(brands, "Rink") and not has_brand(brands, "  ")

    def test_справочник_с_сервера(self, monkeypatch) -> None:
        monkeypatch.setattr(transport, "get", lambda path, *a: [{"id": 1, "name": "Rink"}])
        assert remote.brands() == [Brand(1, "Rink")]

    def test_старый_сервер_без_справочника_не_мешает(self, monkeypatch) -> None:
        def missing(path, *a):
            raise transport.ServerError("Not Found", 404)

        monkeypatch.setattr(transport, "get", missing)
        assert remote.brands() == []

    def test_добавление(self, monkeypatch) -> None:
        sent = []
        monkeypatch.setattr(transport, "post", lambda path, body=None, *a: (
            sent.append((path, body)) or {"id": 9, "name": "Новый"}))
        assert remote.add_brand("Новый") == Brand(9, "Новый")
        assert sent == [("/api/claim-brands", {"name": "Новый"})]


class TestBrandsDialog:
    def test_список_и_поиск(self, app) -> None:
        dialog = BrandsDialog([Brand(1, "Rink"), Brand(2, "Biorepair")], False)
        assert dialog.list.count() == 2
        dialog.search.setText("repa")
        assert dialog.list.count() == 1 and dialog.list.item(0).text() == "Biorepair"

    def test_удалять_может_только_админ(self, app) -> None:
        assert BrandsDialog([], False).delete_button.isHidden()
        assert not BrandsDialog([], True).delete_button.isHidden()

    def test_добавление_попадает_в_общий_список(self, app, monkeypatch) -> None:
        brands = [Brand(1, "Rink")]
        monkeypatch.setattr(torg_dialogs.remote, "add_brand", lambda name: Brand(5, name))
        dialog = BrandsDialog(brands, False)
        dialog.new_name.setText("  Новый   бренд ")

        dialog._add()

        assert brands[-1] == Brand(5, "Новый бренд")
        assert dialog.list.count() == 1 and dialog.new_name.text() == ""

    def test_отказ_сервера_виден(self, app, monkeypatch) -> None:
        def refuse(name):
            raise transport.ServerError("нет связи")

        monkeypatch.setattr(torg_dialogs.remote, "add_brand", refuse)
        dialog = BrandsDialog([], False)
        dialog.new_name.setText("X")

        dialog._add()

        assert "нет связи" in dialog.error.text()

    def test_удаление(self, app, monkeypatch) -> None:
        brands = [Brand(1, "Rink"), Brand(2, "Biorepair")]
        gone = []
        monkeypatch.setattr(torg_dialogs.remote, "delete_brand", gone.append)
        dialog = BrandsDialog(brands, True)
        dialog.list.setCurrentRow(0)  # по алфавиту первый — Biorepair

        dialog._delete()

        assert gone == [2] and brands == [Brand(1, "Rink")]


BRANDS = [Brand(1, "BIOREPAIR"), Brand(2, "Rink"), Brand(3, "Рэд Булл")]


class TestBrandInClaimDialog:
    def make(self, app, claim=None) -> ClaimDialog:
        return ClaimDialog(claim, [], list(BRANDS))

    def test_подсказка_берёт_бренды_из_справочника(self, app) -> None:
        dialog = self.make(app)
        assert dialog.brand.completer().model().stringList() == ["BIOREPAIR", "Rink", "Рэд Булл"]

    def test_подсказка_ищет_по_любому_месту_и_без_регистра(self, app) -> None:
        dialog = self.make(app)
        completer = dialog.brand.completer()
        completer.setCompletionPrefix("пар")
        assert completer.completionCount() == 0
        completer.setCompletionPrefix("epair")
        assert completer.completionCount() == 1

    def test_неизвестный_бренд_предупреждает(self, app) -> None:
        dialog = self.make(app)
        dialog.show()
        dialog.brand.setText("Незнакомый")
        assert not dialog.brand_note.isHidden()
        dialog.brand.setText("rink")
        assert dialog.brand_note.isHidden()

    def test_без_справочника_предупреждений_нет(self, app) -> None:
        dialog = ClaimDialog(None, [], [])
        dialog.show()
        dialog.brand.setText("Что угодно")
        assert dialog.brand_note.isHidden()

    def test_написание_приводится_к_списку(self, app) -> None:
        dialog = self.make(app)
        dialog.brand.setText("biorepair")
        dialog._brand_done()
        assert dialog.brand.text() == "BIOREPAIR"

    def ask(self, monkeypatch, answer) -> list:
        asked = []
        monkeypatch.setattr(claim_dialogs.QMessageBox, "question",
                            lambda *a, **k: asked.append(a[2]) or answer)
        return asked

    def test_новый_бренд_спрашивает_и_запоминает_ответ(self, app, monkeypatch) -> None:
        asked = self.ask(monkeypatch, QMessageBox.StandardButton.Yes)
        dialog = self.make(app)
        dialog.brand.setText("Новый бренд")

        dialog._accept()

        assert dialog.result() == 1 and dialog.add_brand is True
        assert "Новый бренд" in asked[0]

    def test_можно_отказаться_добавлять(self, app, monkeypatch) -> None:
        self.ask(monkeypatch, QMessageBox.StandardButton.No)
        dialog = self.make(app)
        dialog.brand.setText("Новый бренд")

        dialog._accept()

        assert dialog.result() == 1 and dialog.add_brand is False

    def test_отмена_оставляет_окно_открытым(self, app, monkeypatch) -> None:
        self.ask(monkeypatch, QMessageBox.StandardButton.Cancel)
        dialog = self.make(app)
        dialog.brand.setText("Новый бренд")

        dialog._accept()

        assert dialog.result() == 0

    def test_известный_бренд_не_спрашивает(self, app, monkeypatch) -> None:
        asked = self.ask(monkeypatch, QMessageBox.StandardButton.Yes)
        dialog = self.make(app)
        dialog.brand.setText("Rink")

        dialog._accept()

        assert asked == [] and dialog.result() == 1 and dialog.add_brand is False


# --- окно акта -----------------------------------------------------------------

def make_dialog(app, a: TorgAct | None = None, settings: TorgSettings = SETTINGS,
                suppliers: list[dict] | None = None) -> TorgDialog:
    base = a or torg.act_from_claims(
        [Claim(id=1, kind="defect", product="Крем", spec="50 мл", barcode="4600",
               quantity=2.0),
         Claim(id=2, kind="shortage", product="Маска", barcode="4601", quantity=3.0)],
        63, date(2026, 10, 20))
    return TorgDialog(base, settings, suppliers or [])


def set_cell(dialog: TorgDialog, row: int, col: int, text: str) -> None:
    dialog.table.item(row, col).setText(text)


class TestParseNumber:
    @pytest.mark.parametrize("text, value", [
        ("375", 375.0), ("375,5", 375.5), ("1 200", 1200.0), ("12.50", 12.5),
        ("", 0.0), ("  ", 0.0), ("\u00a01\u00a0000", 1000.0)])
    def test_числа(self, text, value) -> None:
        assert parse_number(text) == value

    @pytest.mark.parametrize("text", ["abc", "1,2,3", "nan", "inf", "1e400"])
    def test_мусор(self, text) -> None:
        assert parse_number(text) is None


class TestTorgDialog:
    def test_строки_берутся_из_заявок(self, app) -> None:
        dialog = make_dialog(app)
        assert dialog.table.rowCount() == 2
        assert dialog.table.item(0, 0).text() == "Крем 50 мл"
        assert dialog.table.item(0, COL_DEFECT).text() == "2"
        assert dialog.table.item(1, COL_SHORT).text() == "3"
        assert dialog.table.item(0, COL_SHORT).text() == ""

    def test_номер_и_даты_предзаполнены(self, app) -> None:
        dialog = make_dialog(app)
        assert dialog.number.value() == 63
        assert dialog.act_date.text() == "20.10.2026"

    def test_текст_предложен_по_содержимому(self, app) -> None:
        dialog = make_dialog(app)
        assert "недопоставленную и бракованную" in dialog.conclusion.toPlainText()
        assert dialog.description.toPlainText()

    def test_у_готового_акта_текст_не_перезаписывается(self, app) -> None:
        dialog = make_dialog(app, act(description="Моё описание", conclusion="Моё заключение"))
        assert dialog.description.toPlainText() == "Моё описание"
        assert dialog.conclusion.toPlainText() == "Моё заключение"

    def test_поступившее_считается_само(self, app) -> None:
        dialog = make_dialog(app)
        set_cell(dialog, 1, COL_DOC, "36")
        assert dialog.table.item(1, COL_FACT).text() == "33"
        set_cell(dialog, 1, COL_SHORT, "6")
        assert dialog.table.item(1, COL_FACT).text() == "30"

    def test_итоги_обновляются(self, app) -> None:
        dialog = make_dialog(app)
        set_cell(dialog, 0, COL_PRICE, "375,5")
        set_cell(dialog, 0, COL_DOC, "36")
        assert "брак: 2 шт на 751,00" in dialog.totals.text()

    def test_мусор_в_ячейке_краснеет_и_не_пропускает(self, app) -> None:
        dialog = make_dialog(app)
        set_cell(dialog, 0, COL_PRICE, "дорого")

        dialog._accept()

        assert dialog.result() == 0 and "не число" in dialog.error.text()

    def test_неполный_акт_называет_замечания(self, app) -> None:
        dialog = make_dialog(app)
        dialog._accept()
        text = dialog.error.text()
        assert dialog.result() == 0
        assert "поставщика" in text and "счёта-фактуры" in text

    def test_полный_акт_принимается(self, app) -> None:
        dialog = make_dialog(app)
        dialog.supplier.setCurrentText("ООО Тест, Москва")
        dialog.invoice_number.setText("77")
        dialog.invoice_date.setText("28.09.2026")
        for row in (0, 1):
            set_cell(dialog, row, COL_PRICE, "100")
            set_cell(dialog, row, COL_DOC, "10")

        dialog._accept()

        assert dialog.result() == 1, dialog.error.text()
        result = dialog.act_snapshot()
        assert result.supplier == "ООО Тест, Москва" and result.claim_ids == [1, 2]
        assert result.items[1].qty_fact == 7 and result.problems() == []

    def test_без_реквизитов_акт_не_составить(self, app) -> None:
        dialog = make_dialog(app, settings=TorgSettings())
        dialog._accept()
        assert "Реквизиты" in dialog.error.text()

    def test_неверная_дата_называется(self, app) -> None:
        dialog = make_dialog(app)
        dialog.invoice_date.setText("31.13.2026")
        dialog._accept()
        assert "дата счёта-фактуры" in dialog.error.text()

    def test_выбор_поставщика_подставляет_договор_в_пустое(self, app) -> None:
        hints = [{"supplier": "ООО Старый", "contract_number": "СБ-1",
                  "contract_date": "2023-01-24", "delivery": "ДЛ"}]
        dialog = make_dialog(app, suppliers=hints)

        dialog._pick_supplier(0)

        assert dialog.contract_number.text() == "СБ-1"
        assert dialog.contract_date.text() == "24.01.2023"
        assert dialog.delivery.text() == "ДЛ"

    def test_выбор_поставщика_не_затирает_набранное(self, app) -> None:
        hints = [{"supplier": "ООО Старый", "contract_number": "СБ-1",
                  "contract_date": "2023-01-24", "delivery": "ДЛ"}]
        dialog = make_dialog(app, suppliers=hints)
        dialog.contract_number.setText("МОЙ-5")

        dialog._pick_supplier(0)

        assert dialog.contract_number.text() == "МОЙ-5"

    def test_длинный_текст_предупреждает(self, app) -> None:
        dialog = make_dialog(app)
        dialog.description.setPlainText("слово " * 200)
        assert "Не помещается" in dialog.text_hint.text()

    def test_реквизиты_правятся_из_окна(self, app, monkeypatch) -> None:
        changed = TorgSettings(organization="Другая", director_name="Д. Д.",
                               place="склад", chair_name="Председатель")
        monkeypatch.setattr(torg_dialogs.TorgSettingsDialog, "exec", lambda self: 1)
        monkeypatch.setattr(torg_dialogs.TorgSettingsDialog, "value", lambda self: changed)
        dialog = make_dialog(app)

        dialog._edit_settings()

        assert dialog.settings is changed and dialog.settings_changed is True

    def test_существующий_акт_открывается_со_своими_данными(self, app) -> None:
        saved = act(item(1, qty_doc=40), item(2, qty_defect=0, qty_short=2))
        saved.id = 9
        dialog = make_dialog(app, saved)
        assert "№63" in dialog.windowTitle()
        assert dialog.table.item(0, COL_DOC).text() == "40"
        assert dialog.act_snapshot().id == 9


class TestSettingsDialog:
    def test_пустые_реквизиты_не_сохраняются(self, app) -> None:
        dialog = TorgSettingsDialog(TorgSettings())
        dialog._accept()
        assert dialog.result() == 0 and "Заполните" in dialog.error.text()

    def test_значения_и_комиссия(self, app) -> None:
        dialog = TorgSettingsDialog(SETTINGS)
        dialog._accept()
        assert dialog.result() == 1
        value = dialog.value()
        assert value.organization == SETTINGS.organization
        assert value.members == [("товаровед", "Петров П."), ("товаровед", "Сидоров С.")]

    def test_не_больше_трёх_членов(self, app) -> None:
        assert len(TorgSettingsDialog(SETTINGS)._members) == 3
        many = TorgSettings.from_dict({"members": [["т", str(i)] for i in range(7)]})
        assert len(many.members) == 3

    def test_оборот_через_словарь(self) -> None:
        assert TorgSettings.from_dict(SETTINGS.to_dict()) == SETTINGS

    def test_пустой_ответ_сервера(self) -> None:
        empty = TorgSettings.from_dict(None)
        assert empty.department == "Отдел Склад" and empty.members == []


# --- сервер и страница ---------------------------------------------------------

class FakeTransport:
    def __init__(self, monkeypatch) -> None:
        self.calls: list[tuple] = []
        self.answers: dict[str, object] = {}
        for verb in ("get", "post", "put", "patch", "delete"):
            monkeypatch.setattr(transport, verb, self._make(verb))

    def _make(self, verb):
        def call(path, *args):
            self.calls.append((verb, path, *args))
            return self.answers.get(f"{verb} {path}")
        return call


class TestRemoteActs:
    def test_новый_акт_создаётся_а_старый_правится(self, monkeypatch) -> None:
        fake = FakeTransport(monkeypatch)
        answer = {"id": 4, "number": 63, "act_date": "2026-10-20",
                  "supplier": "ООО Тест", "payload": act().payload()}
        fake.answers["post /api/torg"] = answer
        fake.answers["put /api/torg/4"] = answer

        created = remote.save_act(act())
        a = act()
        a.id = 4
        remote.save_act(a)

        assert fake.calls[0][:2] == ("post", "/api/torg")
        assert fake.calls[0][2]["claim_ids"] == [1]
        assert fake.calls[1][:2] == ("put", "/api/torg/4")
        assert created.id == 4 and created.items[0].qty_defect == 11

    def test_реквизиты(self, monkeypatch) -> None:
        fake = FakeTransport(monkeypatch)
        fake.answers["get /api/torg/settings"] = SETTINGS.to_dict()
        fake.answers["put /api/torg/settings"] = SETTINGS.to_dict()

        assert remote.torg_settings() == SETTINGS
        remote.save_torg_settings(SETTINGS)

        assert fake.calls[1][:2] == ("put", "/api/torg/settings")

    def test_следующий_номер(self, monkeypatch) -> None:
        fake = FakeTransport(monkeypatch)
        fake.answers["get /api/torg/next"] = {"number": 64}
        assert remote.next_act_number(2026) == 64
        assert fake.calls == [("get", "/api/torg/next", {"year": 2026})]

    def test_сервер_без_актов_объясняет_что_делать(self, monkeypatch) -> None:
        def missing(path, *a):
            raise transport.ServerError("Not Found", 404)

        monkeypatch.setattr(transport, "get", missing)
        with pytest.raises(transport.ServerError, match="миграцию"):
            remote.torg_settings()

    def test_файл_акта_заменяет_прежний_и_не_трогает_чужие(self, monkeypatch, tmp_path) -> None:
        a = act(item(1), item(2))
        name = torg.file_name(a)
        files = {1: [ClaimFile(10, 1, name), ClaimFile(11, 1, "фото.jpg")],
                 2: []}
        detached, attached = [], []
        monkeypatch.setattr(remote, "list_files", lambda claim_id: files[claim_id])
        monkeypatch.setattr(remote, "detach_file", detached.append)
        monkeypatch.setattr(remote, "attach_file",
                            lambda claim_id, path: attached.append((claim_id, path)))

        failures = remote.publish_act_file(a, "C:/акт.xlsx")

        assert detached == [10]  # старый акт убран, фото осталось
        assert attached == [(1, "C:/акт.xlsx"), (2, "C:/акт.xlsx")]
        assert failures == []

    def test_закрытая_заявка_не_роняет_остальные(self, monkeypatch) -> None:
        a = act(item(1), item(2))
        attached = []

        def attach(claim_id, path):
            if claim_id == 1:
                raise transport.ServerError("Заявка закрыта")
            attached.append(claim_id)

        monkeypatch.setattr(remote, "list_files", lambda claim_id: [])
        monkeypatch.setattr(remote, "attach_file", attach)

        failures = remote.publish_act_file(a, "x.xlsx")

        assert attached == [2]
        assert len(failures) == 1 and "Заявка №1" in failures[0]


def claim_row(**over) -> Claim:
    data = dict(id=1, kind="defect", brand="Rink", product="Крем", barcode="4600",
                quantity=2.0, status="new", applicant_id=3, manager_id=7,
                act_id=None, act_number=None)
    data.update(over)
    return Claim(**data)


class Notices(list):
    def __call__(self, text, kind) -> None:
        self.append((text, kind))


def make_page(app) -> tuple[ClaimsPage, Notices]:
    notices = Notices()
    return ClaimsPage(AppSettings(), notices), notices


class TestPageTorg:
    def load(self, page: ClaimsPage, *claims_: Claim) -> None:
        page._apply((claims_page.TAB_OPEN, list(claims_), [], []))
        page.table.selectAll()

    def test_без_выделения_подсказывает(self, app) -> None:
        page, notices = make_page(app)
        page.compose_torg()
        assert "Выделите заявки" in notices[0][0]

    def test_кнопка_акта_включается_с_выделением(self, app) -> None:
        page, _ = make_page(app)
        page._apply((claims_page.TAB_OPEN, [claim_row()], [], []))
        assert not page.torg_button.isEnabled()
        page.table.selectRow(0)
        assert page.torg_button.isEnabled()

    def test_заявки_из_разных_актов_не_смешиваются(self, app, monkeypatch) -> None:
        page, notices = make_page(app)
        started = []
        monkeypatch.setattr(claims_page, "run_task", lambda *a, **k: started.append(a))
        self.load(page, claim_row(id=1, act_id=5, act_number=5), claim_row(id=2))

        page.compose_torg()

        assert started == [] and "разных актов" in notices[0][0]

    def test_закрытые_заявки_в_новый_акт_не_идут(self, app, monkeypatch) -> None:
        page, notices = make_page(app)
        monkeypatch.setattr(claims_page, "run_task", lambda *a, **k: pytest.fail("не должен"))
        self.load(page, claim_row(status="closed"))

        page.compose_torg()

        assert "Закрытые" in notices[0][0]

    def test_акт_уже_составленной_заявки_открывается(self, app, monkeypatch) -> None:
        page, _ = make_page(app)
        started = []
        monkeypatch.setattr(claims_page, "run_task", lambda fn, *a, **k: started.append(a))
        self.load(page, claim_row(status="closed", act_id=5, act_number=5))

        page.compose_torg()

        assert started and started[0][1] == 5

    def test_колонка_торг(self) -> None:
        assert ClaimsPage._torg_text(claim_row(act_number=62, files=1)) == "акт №62 · файлов: 1"
        assert ClaimsPage._torg_text(claim_row(files=2)) == "файлов: 2"
        assert ClaimsPage._torg_text(claim_row()) == ""

    def test_подготовка_нового_акта(self, monkeypatch) -> None:
        monkeypatch.setattr(remote, "torg_settings", lambda: SETTINGS)
        monkeypatch.setattr(remote, "torg_suppliers", lambda: [{"supplier": "ООО А"}])
        monkeypatch.setattr(remote, "next_act_number", lambda year: 64)

        result, settings, suppliers = claims_page._prepare_act([claim_row()], None)

        assert result.number == 64 and result.claim_ids == [1]
        assert settings is SETTINGS and suppliers == [{"supplier": "ООО А"}]

    def test_без_справочника_поставщиков_акт_всё_равно_готовится(self, monkeypatch) -> None:
        def refuse():
            raise transport.ServerError("нет")

        monkeypatch.setattr(remote, "torg_settings", lambda: SETTINGS)
        monkeypatch.setattr(remote, "torg_suppliers", refuse)
        monkeypatch.setattr(remote, "next_act_number", lambda year: 1)

        assert claims_page._prepare_act([claim_row()], None)[2] == []

    def test_подготовка_существующего_акта(self, monkeypatch) -> None:
        stored = act()
        stored.id = 5
        monkeypatch.setattr(remote, "torg_settings", lambda: SETTINGS)
        monkeypatch.setattr(remote, "torg_suppliers", lambda: [])
        monkeypatch.setattr(remote, "get_act", lambda act_id: stored)

        assert claims_page._prepare_act([claim_row()], 5)[0] is stored

    def test_порядок_сначала_сервер_потом_файл(self, monkeypatch, tmp_path) -> None:
        order = []
        saved = act()
        saved.id = 5
        monkeypatch.setattr(remote, "save_torg_settings", lambda s: order.append("settings"))
        monkeypatch.setattr(remote, "save_act", lambda a: order.append("act") or saved)
        monkeypatch.setattr(torg, "save", lambda a, s, p: order.append("file"))
        monkeypatch.setattr(remote, "publish_act_file",
                            lambda a, p: order.append("publish") or [])

        result, path, failures = claims_page._make_act(act(), SETTINGS, True, "x.xlsx")

        assert order == ["settings", "act", "file", "publish"]
        assert result is saved and path == "x.xlsx" and failures == []

    def test_реквизиты_не_пишутся_если_не_менялись(self, monkeypatch) -> None:
        order = []
        saved = act()
        monkeypatch.setattr(remote, "save_torg_settings", lambda s: order.append("settings"))
        monkeypatch.setattr(remote, "save_act", lambda a: saved)
        monkeypatch.setattr(torg, "save", lambda a, s, p: None)
        monkeypatch.setattr(remote, "publish_act_file", lambda a, p: [])

        claims_page._make_act(act(), SETTINGS, False, "x.xlsx")

        assert order == []

    def test_не_собравшийся_бланк_не_прикладывается(self, monkeypatch) -> None:
        published = []
        monkeypatch.setattr(remote, "save_act", lambda a: act())

        def broken(a, s, p):
            raise ValueError("не заполнено")

        monkeypatch.setattr(torg, "save", broken)
        monkeypatch.setattr(remote, "publish_act_file", lambda a, p: published.append(1))

        with pytest.raises(ValueError):
            claims_page._make_act(act(), SETTINGS, False, "x.xlsx")
        assert published == []

    def test_новый_бренд_добавляется_до_заявки(self, monkeypatch) -> None:
        order = []
        monkeypatch.setattr(remote, "add_brand", lambda name: order.append(("brand", name)))
        monkeypatch.setattr(remote, "create_claim",
                            lambda fields: order.append(("claim", 0)) or claim_row())

        claims_page._create_many([{"brand": "Новый"}], [], ["Новый"])

        assert order == [("brand", "Новый"), ("claim", 0)]

    def test_отказ_справочника_не_мешает_заявке(self, monkeypatch) -> None:
        def refuse(name):
            raise transport.ServerError("нет")

        monkeypatch.setattr(remote, "add_brand", refuse)
        monkeypatch.setattr(remote, "create_claim", lambda fields: claim_row())

        assert claims_page._create_many([{"brand": "X"}], [], ["X"])[0][0].id == 1

    def test_без_нового_бренда_справочник_не_трогается(self, monkeypatch) -> None:
        monkeypatch.setattr(remote, "add_brand", lambda name: pytest.fail("не должен"))
        monkeypatch.setattr(remote, "create_claim", lambda fields: claim_row())

        claims_page._create_many([{"brand": "Rink"}], [], [])

    def test_страница_читает_бренды_вместе_со_списком(self, monkeypatch) -> None:
        monkeypatch.setattr(remote, "list_claims", lambda **k: [claim_row()])
        monkeypatch.setattr(remote, "managers", lambda: [])
        monkeypatch.setattr(remote, "brands", lambda: [Brand(1, "Rink")])

        tab, items, managers, brands = claims_page._fetch(claims_page.TAB_OPEN, "")

        assert len(items) == 1 and brands == [Brand(1, "Rink")]

    def test_бренды_обновляются_на_месте(self, app) -> None:
        page, _ = make_page(app)
        shared = page._brands
        page._apply((claims_page.TAB_OPEN, [], [], [Brand(1, "Rink")]))
        assert page._brands is shared and shared == [Brand(1, "Rink")]


# --- несколько позиций в одной поставке ----------------------------------------

def fill_position(dialog: ClaimDialog, brand: str, product: str, quantity: float = 1,
                  barcode: str = "") -> None:
    dialog.brand.setText(brand)
    dialog.product.setText(product)
    dialog.quantity.setValue(quantity)
    dialog.barcode.setText(barcode)


class TestSeveralPositions:
    def make(self, app) -> ClaimDialog:
        return ClaimDialog(None, [], list(BRANDS))

    def test_позиция_встаёт_в_список_и_поля_очищаются(self, app) -> None:
        dialog = self.make(app)
        dialog.kind.setCurrentIndex(dialog.kind.findData("defect"))
        fill_position(dialog, "Rink", "Крем", 2, "4600")

        dialog._add_position()

        assert dialog.positions.count() == 1
        assert "Брак: Rink · Крем — 2 шт" in dialog.positions.item(0).text()
        assert dialog.product.text() == dialog.barcode.text() == dialog.spec.text() == ""
        assert dialog.quantity.value() == 1
        # Вид остаётся: в поставке чаще всего всё одного рода.
        assert dialog.kind.currentData() == "defect"

    def test_бренд_остаётся_после_добавления_позиции(self, app) -> None:
        dialog = self.make(app)
        fill_position(dialog, "Rink", "Крем")

        dialog._add_position()

        assert dialog.brand.text() == "Rink"

    def test_второй_бренд_печатается_поверх_прежнего(self, app) -> None:
        dialog = self.make(app)
        fill_position(dialog, "Rink", "Крем")
        dialog._add_position()
        fill_position(dialog, "BIOREPAIR", "Паста")
        dialog._add_position()

        assert [c["brand"] for c in dialog.claims_fields()] == ["Rink", "BIOREPAIR"]

    def test_оставшийся_бренд_не_становится_лишней_позицией(self, app) -> None:
        dialog = self.make(app)
        fill_position(dialog, "Rink", "Крем")
        dialog._add_position()
        fill_position(dialog, "Rink", "Маска")
        dialog._add_position()
        # В поле остался бренд прошлой позиции, больше ничего не набрано.

        dialog._accept()

        assert dialog.result() == 1
        assert [c["product"] for c in dialog.claims_fields()] == ["Крем", "Маска"]

    def test_бренд_без_товара_как_единственная_позиция_допустим(self, app) -> None:
        dialog = self.make(app)
        dialog.brand.setText("Rink")

        dialog._accept()

        assert dialog.result() == 1 and len(dialog.claims_fields()) == 1

    def test_несколько_позиций_дают_несколько_заявок_с_общими_данными(self, app) -> None:
        dialog = self.make(app)
        dialog.comment.setPlainText("Замятие упаковки")
        fill_position(dialog, "Rink", "Крем", 2)
        dialog._add_position()
        fill_position(dialog, "BIOREPAIR", "Паста", 11)
        dialog._add_position()

        result = dialog.claims_fields()

        assert [c["product"] for c in result] == ["Крем", "Паста"]
        assert [c["quantity"] for c in result] == [2.0, 11.0]
        assert all(c["comment"] == "Замятие упаковки" for c in result)

    def test_набранная_но_не_добавленная_позиция_не_теряется(self, app) -> None:
        dialog = self.make(app)
        fill_position(dialog, "Rink", "Крем")
        dialog._add_position()
        fill_position(dialog, "BIOREPAIR", "Паста")

        dialog._accept()

        assert dialog.result() == 1
        assert [c["product"] for c in dialog.claims_fields()] == ["Крем", "Паста"]

    def test_одна_позиция_как_раньше(self, app) -> None:
        dialog = self.make(app)
        fill_position(dialog, "Rink", "Крем")

        dialog._accept()

        assert dialog.result() == 1 and len(dialog.claims_fields()) == 1

    def test_без_позиций_не_принимает(self, app) -> None:
        dialog = self.make(app)
        dialog._accept()
        assert dialog.result() == 0 and "хотя бы бренд" in dialog.error.text()

    def test_пустую_позицию_добавить_нельзя(self, app) -> None:
        dialog = self.make(app)
        dialog._add_position()
        assert dialog.positions.count() == 0 and "хотя бы бренд" in dialog.error.text()

    def test_позицию_можно_убрать(self, app) -> None:
        dialog = self.make(app)
        for name in ("Крем", "Паста"):
            fill_position(dialog, "Rink", name)
            dialog._add_position()
        dialog.positions.setCurrentRow(0)

        dialog._remove_position()

        assert [p["product"] for p in dialog._positions] == ["Паста"]
        assert dialog.positions.count() == 1

    def test_неверный_срок_годности_не_даёт_добавить(self, app) -> None:
        dialog = self.make(app)
        fill_position(dialog, "Rink", "Крем")
        dialog.expiry.setText("31.13.2027")

        dialog._add_position()

        assert dialog.positions.count() == 0 and "дд.мм.гггг" in dialog.error.text()

    def test_новые_бренды_спрашиваются_один_раз_за_все_позиции(self, app, monkeypatch) -> None:
        asked = []
        monkeypatch.setattr(claim_dialogs.QMessageBox, "question",
                            lambda *a, **k: asked.append(a[2]) or QMessageBox.StandardButton.Yes)
        dialog = self.make(app)
        for brand in ("Новый А", "новый а", "Rink", "Новый Б"):
            fill_position(dialog, brand, "Товар")
            dialog._add_position()

        dialog._accept()

        assert len(asked) == 1 and "Новый А" in asked[0] and "Новый Б" in asked[0]
        assert dialog.new_brands == ["Новый А", "Новый Б"]

    def test_заведённая_заявка_без_списка_позиций(self, app, monkeypatch) -> None:
        monkeypatch.setattr(ClaimDialog, "_reload_files", lambda self: None)
        dialog = ClaimDialog(claim_row(), [], list(BRANDS))
        assert not hasattr(dialog, "positions")


class TestCreateMany:
    def fake(self, monkeypatch, fail_on=()):
        created, attached = [], []

        def create(fields):
            if fields["product"] in fail_on:
                raise transport.ServerError("отказ")
            claim = claim_row(id=len(created) + 1, product=fields["product"])
            created.append(claim)
            return claim

        monkeypatch.setattr(remote, "create_claim", create)
        monkeypatch.setattr(remote, "attach_file", lambda cid, path: attached.append((cid, path)))
        return created, attached

    def test_каждая_позиция_становится_заявкой_и_получает_файлы(self, monkeypatch) -> None:
        created, attached = self.fake(monkeypatch)

        claims, failures = claims_page._create_many(
            [{"product": "А"}, {"product": "Б"}], ["акт.pdf"], [])

        assert [c.id for c in claims] == [1, 2] and failures == []
        assert attached == [(1, "акт.pdf"), (2, "акт.pdf")]

    def test_отказ_по_одной_позиции_не_отменяет_остальные(self, monkeypatch) -> None:
        self.fake(monkeypatch, fail_on=("Б",))

        claims, failures = claims_page._create_many(
            [{"product": "А"}, {"product": "Б"}, {"product": "В"}], [], [])

        assert [c.product for c in claims] == ["А", "В"]
        assert len(failures) == 1 and "Позиция 2" in failures[0]

    def test_если_не_заведена_ни_одна_ошибка_поднимается(self, monkeypatch) -> None:
        self.fake(monkeypatch, fail_on=("А", "Б"))

        with pytest.raises(transport.ServerError, match="Позиция 1 не заведена|отказ"):
            claims_page._create_many([{"product": "А"}, {"product": "Б"}], [], [])

    def test_новые_бренды_добавляются_один_раз_до_заявок(self, monkeypatch) -> None:
        order = []
        monkeypatch.setattr(remote, "add_brand", lambda name: order.append(("brand", name)))
        monkeypatch.setattr(remote, "create_claim",
                            lambda fields: order.append(("claim", fields["product"])) or claim_row())

        claims_page._create_many([{"product": "А"}, {"product": "Б"}], [], ["Новый"])

        assert order == [("brand", "Новый"), ("claim", "А"), ("claim", "Б")]

    def test_после_нескольких_заявок_предлагается_акт(self, app, monkeypatch) -> None:
        page, notices = make_page(app)
        started = []
        monkeypatch.setattr(claims_page, "run_task", lambda fn, *a, **k: started.append((fn, a)))
        monkeypatch.setattr(claims_page.QMessageBox, "question",
                            lambda *a, **k: QMessageBox.StandardButton.Yes)

        page._created(([claim_row(id=1), claim_row(id=2)], []))

        assert "Заведено заявок: 2" in notices[0][0]
        prepare = [item for item in started if item[0] is claims_page._prepare_act]
        assert prepare and [c.id for c in prepare[0][1][0]] == [1, 2]

    def test_можно_отказаться_от_акта(self, app, monkeypatch) -> None:
        page, _ = make_page(app)
        started = []
        monkeypatch.setattr(claims_page, "run_task", lambda fn, *a, **k: started.append(fn))
        monkeypatch.setattr(claims_page.QMessageBox, "question",
                            lambda *a, **k: QMessageBox.StandardButton.No)

        page._created(([claim_row(id=1), claim_row(id=2)], []))

        assert claims_page._prepare_act not in started

    def test_одна_заявка_акт_не_предлагает(self, app, monkeypatch) -> None:
        page, notices = make_page(app)
        monkeypatch.setattr(claims_page, "run_task", lambda *a, **k: None)
        monkeypatch.setattr(claims_page.QMessageBox, "question",
                            lambda *a, **k: pytest.fail("не должен спрашивать"))

        page._created(([claim_row(id=1)], []))

        assert "Заявка заведена" in notices[0][0]


# --- ввод цифр в таблице акта --------------------------------------------------

class TestCellEditor:
    def test_редактор_без_тесных_отступов_темы(self, app) -> None:
        dialog = make_dialog(app)
        delegate = dialog.table.itemDelegateForColumn(COL_PRICE)

        editor = delegate.createEditor(dialog.table, None, None)

        # Отступы темы (7 на 10 пикселей) в строке таблицы съедали цифры.
        assert "padding: 0 8px" in editor.styleSheet()

    def test_редактор_только_у_числовых_колонок(self, app) -> None:
        dialog = make_dialog(app)
        for col in (COL_PRICE, COL_DOC, COL_SHORT, COL_DEFECT):
            assert isinstance(dialog.table.itemDelegateForColumn(col),
                              torg_dialogs._NumberCellDelegate)
        assert not isinstance(dialog.table.itemDelegateForColumn(0),
                              torg_dialogs._NumberCellDelegate)

    def test_редактор_пропускает_цифры_и_не_пропускает_буквы(self, app) -> None:
        from PySide6.QtGui import QValidator

        dialog = make_dialog(app)
        editor = dialog.table.itemDelegateForColumn(COL_PRICE).createEditor(
            dialog.table, None, None)
        validator = editor.validator()

        assert validator.validate("375,50", 0)[0] == QValidator.State.Acceptable
        assert validator.validate("1 200.5", 0)[0] == QValidator.State.Acceptable
        assert validator.validate("abc", 0)[0] == QValidator.State.Invalid

    def test_строки_выше_редактора(self, app) -> None:
        dialog = make_dialog(app)
        dialog.show()
        assert dialog.table.verticalHeader().defaultSectionSize() == torg_dialogs.ROW_HEIGHT
        assert torg_dialogs.ROW_HEIGHT >= 34

    def test_набранное_число_доходит_до_акта(self, app) -> None:
        dialog = make_dialog(app)
        dialog.show()
        index = dialog.table.model().index(0, COL_PRICE)
        dialog.table.setCurrentIndex(index)
        dialog.table.edit(index)
        editor = dialog.table.findChild(torg_dialogs.QLineEdit)
        assert editor is not None and editor.height() >= 30
        editor.setText("375,5")
        dialog.table.commitData(editor)
        dialog.table.closeEditor(editor, QAbstractItemDelegate.EndEditHint.NoHint)

        assert dialog.act_snapshot().items[0].price == 375.5
