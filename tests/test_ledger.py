"""Разбор ведомости по товарам на складах и выводы по ней.

Файл собирается здесь же, а не берётся с чьего-то рабочего стола: проверять
разбор выгрузки по файлу, которого нет в репозитории, нельзя — тест начнёт
падать на другой машине, и причина будет не в коде.

Проверяется то, на чём такой разбор ломается молча: шапка с одноимёнными
складами, строки групп с итогами вложенных товаров, служебные склады «в пути» и
колонка «Итого» рядом с данными. Каждая из этих ловушек даёт не ошибку, а
неверные числа в отчёте, который выглядит целым.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core import ledger
from app.core.ledger import parse

FIELDS = ("Начальный остаток", "Приход", "Расход", "Конечный остаток")


def _book(path: Path, stores: list[str], rows: list[tuple],
          *, unit_label: str = "Ед. изм.", total: list[float] | None = None,
          fields: tuple[str, ...] = FIELDS,
          measures: tuple[str, ...] = ("Количество",),
          labels: list[str] | None = None, total_row: bool = False,
          filename: str = "ведомость.xlsx") -> str:
    """Собирает выгрузку в том виде, в каком её отдаёт 1С: склады в колонках.

    `rows` — кортежи (артикул, номенклатура, ед.изм, [числа по складам]) и
    необязательный пятый элемент — код. Пустая единица измерения означает
    строку группы. Чисел на склад столько, сколько полей в `fields`; для
    показателя «Сумма» в файл пишется число ×100 — чтобы разбор, взявший
    сумму вместо количества, был пойман.
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист_1"
    sheet["A2"] = "Ведомость по товарам на складах"
    sheet["A4"] = "Параметры:"
    sheet["A5"] = "Отбор:"

    labels = labels or ["Артикул", "Номенклатура", unit_label]
    for index, label in enumerate(labels):
        sheet.cell(7, 1 + index, label)
    first = 1 + len(labels)
    step = len(fields)
    width = step * len(measures)
    names = list(stores) + (["Итого"] if total is not None else [])
    for index, name in enumerate(names):
        start = first + index * width
        sheet.cell(7, start, name)
        sheet.merge_cells(start_row=7, start_column=start,
                          end_row=7, end_column=start + width - 1)
        for number, measure in enumerate(measures):
            sheet.cell(8, start + number * step, measure)
            for offset, field in enumerate(fields):
                sheet.cell(9, start + number * step + offset, field)

    sums: dict[int, float] = {}

    def put(row: int, column: int, value: float) -> None:
        sheet.cell(row, column, value)
        sums[column] = sums.get(column, 0) + value

    row = 10
    for row, (article, name, unit, values, *code) in enumerate(rows, 10):
        cells = [article, name, unit] + list(code)
        for index, value in enumerate(cells[:len(labels)]):
            sheet.cell(row, 1 + index, value)
        for index, value in enumerate(values):
            if value is None:
                continue
            store, offset = divmod(index, step)
            for number in range(len(measures)):
                column = first + store * width + number * step + offset
                scale = 100 if number else 1
                if unit:
                    put(row, column, value * scale)
                else:
                    sheet.cell(row, column, value * scale)
        if total is not None and unit:
            # Итог по строке: сумма складов.
            start = first + len(stores) * width
            for offset in range(step):
                pieces = [values[index * step + offset]
                          for index in range(len(stores))
                          if index * step + offset < len(values)
                          and values[index * step + offset] is not None]
                if pieces:
                    put(row, start + offset, sum(pieces))

    if total_row:
        row += 1
        sheet.cell(row, 1, "Итого")
        for column, value in sums.items():
            sheet.cell(row, column, value)

    target = str(path / filename)
    book.save(target)
    return target


def _line(*values) -> list:
    return list(values)


# --- шапка ---------------------------------------------------------------------

def test_склады_читаются_из_шапки(tmp_path):
    path = _book(tmp_path, ["Реми", "ВС - В пути Сэм", "Сити Молл"],
                 [("1", "Мыло", "шт", _line(10, 0, 3, 7, 0, 0, 0, 0, 5, 0, 1, 4))])

    led = ledger.read(path)

    assert [store.title for store in led.stores] == [
        "Реми", "ВС - В пути Сэм", "Сити Молл"]
    assert [store.title for store in led.shops] == ["Реми", "Сити Молл"]
    assert [store.title for store in led.transit] == ["ВС - В пути Сэм"]


def test_одноимённые_склады_разведены_номером(tmp_path):
    """Два склада с одним названием — данные одного не должны затирать другой."""
    path = _book(tmp_path, ["Пирожок", "Пирожок"],
                 [("1", "Мыло", "шт", _line(10, 0, 4, 6, 20, 0, 5, 15))])

    led = ledger.read(path)

    assert [store.title for store in led.stores] == ["Пирожок", "Пирожок (2)"]
    assert [store.doubled for store in led.stores] == [False, True]
    item = led.items[0]
    assert item.move("Пирожок").outgoing == 4
    assert item.move("Пирожок (2)").outgoing == 5
    assert led.total().outgoing == 9, "оба склада должны попасть в сумму"


def test_итоговая_колонка_не_склад(tmp_path):
    path = _book(tmp_path, ["Реми"],
                 [("1", "Мыло", "шт", _line(10, 0, 3, 7))], total=[10, 0, 3, 7])

    led = ledger.read(path)

    assert [store.title for store in led.stores] == ["Реми"]
    assert led.reported is not None and led.reported.outgoing == 3


@pytest.mark.parametrize("label", ["Ед. изм.", "Ед.изм.", "Ед.изм", "ЕД. ИЗМ."])
def test_единица_измерения_узнаётся_в_разных_написаниях(tmp_path, label):
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 3, 7))],
                 unit_label=label)

    assert len(ledger.read(path).items) == 1


def test_чужой_файл_отклоняется_с_объяснением(tmp_path):
    book = Workbook()
    book.active["A1"] = "просто таблица"
    target = str(tmp_path / "чужое.xlsx")
    book.save(target)

    with pytest.raises(parse.ParseError) as error:
        ledger.read(target)

    assert "ведомость" in str(error.value).lower()


# --- строки групп --------------------------------------------------------------

def test_строки_групп_не_складываются_с_товарами(tmp_path):
    """У групп в этой выгрузке стоят итоги вложенных товаров.

    Сложить их вместе с товарами — посчитать всё дважды. Отличает группу
    отсутствие единицы измерения.
    """
    path = _book(tmp_path, ["Реми"], [
        ("", "Красота и уход", "", _line(30, 0, 9, 21)),
        ("1", "Мыло", "шт", _line(10, 0, 4, 6)),
        ("2", "Крем", "шт", _line(20, 0, 5, 15)),
    ])

    led = ledger.read(path)

    assert [item.name for item in led.items] == ["Мыло", "Крем"]
    assert led.total().outgoing == 9


# --- сверка с файлом -----------------------------------------------------------

def test_разбор_сходится_с_итогом_файла(tmp_path):
    path = _book(tmp_path, ["Реми", "Сити Молл"],
                 [("1", "Мыло", "шт", _line(10, 0, 4, 6, 20, 0, 5, 15))],
                 total=[30, 0, 9, 21])

    assert ledger.read(path).balanced


def test_несошедшийся_итог_не_замалчивается(tmp_path):
    """Итог в файле считан по колонке, которую разбор потерял.

    Это главная проверка разбора: если шапку прочитали неверно, числа встанут
    не под тем складом, а отчёт будет выглядеть целым.
    """
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 4, 6))],
                 total=[10, 0, 4, 6])
    # Портим итог так, будто в отчёте был склад, которого разбор не увидел.
    from openpyxl import load_workbook
    book = load_workbook(path)
    book.active.cell(10, 8 + 2, 99)
    book.save(path)

    assert not ledger.read(path).balanced


# --- выводы --------------------------------------------------------------------

@pytest.fixture
def report(tmp_path):
    """Ведомость, в которой есть каждый разбираемый случай."""
    path = _book(tmp_path, ["Реми", "Сити Молл", "ВС - В пути Сэм"], [
        # ходовой везде, остатки есть
        ("1", "Салфетки", "шт", _line(100, 0, 60, 40, 80, 0, 30, 50, 0, 0, 0, 0)),
        # кончился в «Сити Молл», в «Реми» остаток есть
        ("2", "Паста", "шт", _line(50, 0, 10, 40, 12, 0, 12, 0, 0, 0, 0, 0)),
        # остатка меньше, чем ушло, — кончится следующим
        ("3", "Патчи", "шт", _line(30, 0, 20, 10, 0, 0, 0, 0, 0, 0, 0, 0)),
        # отрицательный остаток: расход больше, чем было
        ("4", "Пакет", "шт", _line(5, 0, 23, -18, 0, 0, 0, 0, 0, 0, 0, 0)),
        # лежит в пути, движения нет
        ("5", "Маска", "шт", _line(0, 0, 0, 0, 0, 0, 0, 0, 400, 0, 0, 400)),
    ])
    return ledger.build(ledger.read(path))


def test_топ_упорядочен_по_расходу(report):
    # Салфетки 90, Пакет 23, Паста 22, Патчи 20.
    assert [item.name for item in report.top] == [
        "Салфетки", "Пакет", "Паста", "Патчи", "Маска"]
    assert report.top[0].move.outgoing == 90


def test_доля_расхода_считается_от_всех_магазинов(report):
    # 90 + 22 + 20 + 23 = 155
    assert report.total.outgoing == 155
    assert report.top[0].share == pytest.approx(90 / 155)


def test_товар_ходовой_везде_отличим_от_одного_магазина(report):
    """«В скольких магазинах был расход» — разные решения при равном расходе."""
    by_name = {item.name: item for item in report.items}

    assert by_name["Салфетки"].active == ["Реми", "Сити Молл"]
    assert by_name["Патчи"].active == ["Реми"]


def test_кончилось_а_расход_был(report):
    """Остатка нет, а расход был — то, что нужно смотреть первым."""
    found = {(line.name, line.store) for line in report.shortages}

    assert ("Паста", "Сити Молл") in found
    assert ("Паста", "Реми") not in found, "в «Реми» остаток есть"
    assert ("Салфетки", "Реми") not in found


def test_остатка_меньше_чем_ушло(report):
    found = {(line.name, line.store) for line in report.tight}

    # Патчи: ушло 20, осталось 10. Салфетки: ушло 60, осталось 40 — тоже меньше
    # ушедшего, и попасть в список должны обе.
    assert ("Патчи", "Реми") in found
    assert ("Салфетки", "Реми") in found
    assert ("Паста", "Реми") not in found, "осталось 40 при расходе 10"


def test_отрицательный_остаток_выделен_отдельно(report):
    """Минус — ошибка учёта: завозом не лечится, и причина другая.

    В список кончившихся он попадает тоже: товара на полке действительно нет.
    Но искать по нему нужно непроведённый приход, а не поставщика.
    """
    assert [line.name for line in report.negative] == ["Пакет"]
    assert ("Пакет", "Реми") in {(l.name, l.store) for l in report.shortages}


def test_склады_в_пути_не_попадают_в_расход(report):
    """Продаж в пути не бывает, а остаток видеть нужно."""
    assert all(line.store != "ВС - В пути Сэм" for line in report.shortages)
    assert [(line.name, line.move.closing) for line in report.transit] == [
        ("Маска", 400)]
    # Остаток магазинов: Реми 40+40+10-18=72, Сити Молл 50. Четыреста в пути
    # сюда не входят.
    assert report.total.closing == 122, "в пути в остаток магазинов не входит"


def test_магазины_упорядочены_по_расходу(report):
    assert [total.store.title for total in report.stores] == ["Реми", "Сити Молл"]
    assert report.stores[0].move.outgoing == 113, "60+10+20+23"
    assert report.stores[0].ran_out == 1, "«Пакет» ушёл в минус"


# --- выгрузка ------------------------------------------------------------------

def test_выгрузка_раскладывает_по_листам(report, tmp_path):
    from openpyxl import load_workbook

    target = ledger.save(report, str(tmp_path / "разбор.xlsx"))
    book = load_workbook(target)

    assert book.sheetnames == ["Главное", "Не хватило", "Топ товаров",
                               "По магазинам", "Свод"]
    # Свод: на каждый магазин две колонки, склады «в пути» в него не входят.
    head = [book["Свод"].cell(4, column).value
            for column in range(1, book["Свод"].max_column + 1)]
    assert head[:3] == ["Артикул", "Товар", "Ед."]
    assert "ВС - В пути Сэм" not in head


def test_имя_файла_по_умолчанию_с_датой(tmp_path):
    name = ledger.default_name(str(tmp_path))

    assert name.endswith(".xlsx")
    assert "Ведомость" in name


# --- разные виды отчёта --------------------------------------------------------

def test_отчёт_второго_подразделения(tmp_path):
    """Шапка как в настройке второго подразделения: «Номенклатура.Код» и
    «Номенклатура, Характеристика» вместо «Номенклатуры», итог и колонкой, и
    строкой внизу. Раньше такой файл отклонялся целиком."""
    path = _book(tmp_path, ["Магазин Уссурийск", "Магазин Первый Парфюмерный"], [
        ("", "Biorepair", "", _line(10, 2, 5, 7, 20, 0, 5, 15)),
        ("GA1297700/GA1511900", "Паста Junior", "шт",
         _line(4, 2, 1, 5, 20, 0, 5, 15), "00-00057333"),
        ("GA2019100", "Паста Banana", "шт", _line(6, 0, 4, 2), "00-00095038"),
    ], labels=["Артикул", "Номенклатура, Характеристика", "Ед. изм.",
               "Номенклатура.Код"], total=[], total_row=True)

    led = ledger.read(path)

    assert led.layout == parse.LAYOUT_WIDE
    assert [store.title for store in led.stores] == [
        "Магазин Уссурийск", "Магазин Первый Парфюмерный"]
    assert [item.name for item in led.items] == ["Паста Junior", "Паста Banana"]
    assert [item.code for item in led.items] == ["00-00057333", "00-00095038"]
    assert led.total().outgoing == 10, "группа и строка «Итого» не в счёт"
    assert led.balanced


def test_настоящий_файл_второго_подразделения():
    """Живой образец, если он лежит рядом с проектом. В репозиторий он не
    входит — на другой машине тест пропускается."""
    sample = (Path(__file__).resolve().parents[1] / "Для отчётности"
              / "Пример отчёта второго подразделения.xlsx")
    if not sample.exists():
        pytest.skip("образца нет на этой машине")

    led = ledger.read(str(sample))

    assert len(led.shops) == 3
    assert led.items and led.balanced


@pytest.mark.parametrize("fields", [
    ("Остаток на начало", "Поступление", "Расход", "Остаток на конец"),
    ("Нач. остаток", "Приход", "Расход", "Кон. остаток"),
])
def test_подписи_полей_в_других_написаниях(tmp_path, fields):
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 2, 3, 9))],
                 fields=fields)

    move = ledger.read(path).items[0].move("Реми")

    assert (move.opening, move.incoming, move.outgoing, move.closing) == (10, 2, 3, 9)


@pytest.mark.parametrize("label", ["Товар", "Номенклатура с характеристикой",
                                   "Наименование"])
def test_колонка_товара_в_других_написаниях(tmp_path, label):
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 3, 7))],
                 labels=["Артикул", label, "Ед. изм."])

    assert [item.name for item in ledger.read(path).items] == ["Мыло"]


def test_суммы_не_попадают_в_количество(tmp_path):
    """На склад два блока — «Количество» и «Сумма». Считаются только штуки."""
    path = _book(tmp_path, ["Реми", "Сити Молл"],
                 [("1", "Мыло", "шт", _line(10, 0, 3, 7, 5, 1, 2, 4))],
                 measures=("Количество", "Сумма"), total=[])

    led = ledger.read(path)

    assert [store.title for store in led.stores] == ["Реми", "Сити Молл"]
    assert led.items[0].move("Реми").outgoing == 3
    assert led.items[0].move("Сити Молл").outgoing == 2
    assert led.balanced


def test_только_суммы_отклоняются_с_объяснением(tmp_path):
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 3, 7))],
                 measures=("Сумма",))

    with pytest.raises(parse.ParseError, match="Количество"):
        ledger.read(path)


def test_без_начального_остатка_доли_не_выдумываются(tmp_path):
    """Нет «Начального остатка» и «Прихода» — запас неизвестен."""
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(3, 0))],
                 fields=("Расход", "Конечный остаток"))

    led = ledger.read(path)
    report = ledger.build(led)

    assert led.items[0].move("Реми").outgoing == 3
    assert not led.complete and not report.shares
    assert report.shortages, "расход был, остатка нет — это видно и без запаса"
    assert any("Начальный остаток" in note for note in led.warnings)


def test_без_единицы_группа_узнаётся_по_группировке_строк(tmp_path):
    path = _book(tmp_path, ["Реми"], [
        ("", "Уход", "", _line(30, 0, 9, 21)),
        ("1", "Мыло", "", _line(10, 0, 4, 6)),
        ("2", "Крем", "", _line(20, 0, 5, 15)),
    ], labels=["Артикул", "Номенклатура"])
    from openpyxl import load_workbook
    book = load_workbook(path)
    for row in (11, 12):
        book.active.row_dimensions[row].outline_level = 1
    book.save(path)

    led = ledger.read(path)

    assert [item.name for item in led.items] == ["Мыло", "Крем"]


def _flat_book(path: Path, rows: list[tuple], *, total_row: bool = True) -> str:
    """Плоская таблица: (склад, артикул, номенклатура, ед., 4 числа)."""
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Ведомость по товарам на складах"
    for column, label in enumerate(["Склад", "Артикул", "Номенклатура",
                                    "Ед. изм.", *FIELDS], 1):
        sheet.cell(3, column, label)
    sums = [0.0] * 4
    row = 3
    for row, (store, article, name, unit, *values) in enumerate(rows, 4):
        for column, value in enumerate([store, article, name, unit, *values], 1):
            sheet.cell(row, column, value)
        sums = [total + value for total, value in zip(sums, values)]
    if total_row:
        sheet.cell(row + 1, 1, "Итого")
        for offset, value in enumerate(sums):
            sheet.cell(row + 1, 5 + offset, value)
    target = str(path / "плоская.xlsx")
    book.save(target)
    return target


def test_склад_колонкой(tmp_path):
    """Один товар на двух складах — две строки, в разборе один товар."""
    path = _flat_book(tmp_path, [
        ("Реми", "1", "Мыло", "шт", 10, 0, 4, 6),
        ("Сити Молл", "1", "Мыло", "шт", 20, 0, 5, 15),
        ("Реми", "2", "Крем", "шт", 5, 0, 1, 4),
    ])

    led = ledger.read(path)

    assert led.layout == parse.LAYOUT_FLAT
    assert [store.title for store in led.stores] == ["Реми", "Сити Молл"]
    assert [item.name for item in led.items] == ["Мыло", "Крем"]
    assert led.items[0].move("Сити Молл").outgoing == 5
    assert led.total().outgoing == 10
    assert led.balanced


def _grouped_book(path: Path, blocks: dict[str, list[tuple]], *,
                  levels: bool = True) -> str:
    """Склады группами строк: склад → группа товаров → товары."""
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Ведомость по товарам на складах"
    sheet["B3"] = "Склад"
    sheet["A4"], sheet["B4"], sheet["C4"] = "Артикул", "Номенклатура", "Ед. изм."
    for offset, label in enumerate(FIELDS):
        sheet.cell(4, 4 + offset, label)

    row = 5
    for store, rows in blocks.items():
        # Итоги склада и группы стоят в самих строках — их считать нельзя.
        totals = [sum(values[i] for *_, values in rows) for i in range(4)]
        for level, name in ((0, store), (1, "Уход")):
            sheet.cell(row, 2, name)
            for offset, value in enumerate(totals):
                sheet.cell(row, 4 + offset, value)
            if levels:
                sheet.row_dimensions[row].outline_level = level
            row += 1
        for article, name, values in rows:
            sheet.cell(row, 1, article)
            sheet.cell(row, 2, name)
            sheet.cell(row, 3, "шт")
            for offset, value in enumerate(values):
                sheet.cell(row, 4 + offset, value)
            if levels:
                sheet.row_dimensions[row].outline_level = 2
            row += 1
    target = str(path / "группами.xlsx")
    book.save(target)
    return target


def test_склады_группами_строк(tmp_path):
    path = _grouped_book(tmp_path, {
        "Реми": [("1", "Мыло", [10, 0, 4, 6]), ("2", "Крем", [5, 0, 1, 4])],
        "Сити Молл": [("1", "Мыло", [20, 0, 5, 15])],
    })

    led = ledger.read(path)

    assert led.layout == parse.LAYOUT_GROUPED
    assert [store.title for store in led.stores] == ["Реми", "Сити Молл"]
    assert [item.name for item in led.items] == ["Мыло", "Крем"]
    assert led.items[0].move("Сити Молл").outgoing == 5
    assert led.total().outgoing == 10, "итоги склада и группы не в счёт"


def test_склады_группами_без_группировки_не_угадываются(tmp_path):
    """Без уровней группировки склад не отличить от группы товаров."""
    path = _grouped_book(tmp_path, {"Реми": [("1", "Мыло", [10, 0, 4, 6])]},
                         levels=False)

    with pytest.raises(parse.ParseError, match="группировка"):
        ledger.read(path)


def _single_book(path: Path, filename: str, rows: list[tuple],
                 selection: str = "") -> str:
    """Выгрузка по одному складу: над полями складов нет."""
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Ведомость по товарам на складах"
    sheet["A3"], sheet["C3"] = "Отбор:", selection
    sheet["A5"], sheet["B5"], sheet["C5"] = "Артикул", "Номенклатура", "Ед. изм."
    for offset, label in enumerate(FIELDS):
        sheet.cell(5, 4 + offset, label)
    for row, (article, name, *values) in enumerate(rows, 6):
        for column, value in enumerate([article, name, "шт", *values], 1):
            sheet.cell(row, column, value)
    target = str(path / filename)
    book.save(target)
    return target


def test_один_склад_назван_по_отбору(tmp_path):
    path = _single_book(tmp_path, "выгрузка.xlsx", [("1", "Мыло", 10, 0, 4, 6)],
                        selection='Склад Равно "Магазин Уссурийск"')

    led = ledger.read(path)

    assert led.layout == parse.LAYOUT_SINGLE
    assert [store.title for store in led.stores] == ["Магазин Уссурийск"]
    assert not led.warnings


def test_один_склад_без_отбора_назван_по_файлу(tmp_path):
    path = _single_book(tmp_path, "Реми.xlsx", [("1", "Мыло", 10, 0, 4, 6)])

    led = ledger.read(path)

    assert [store.title for store in led.stores] == ["Реми"]
    assert led.warnings, "имя из файла — догадка, о ней нужно сказать"


def test_нет_конечного_остатка_объясняется(tmp_path):
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 3))],
                 fields=("Начальный остаток", "Приход", "Расход"))

    with pytest.raises(parse.ParseError, match="Конечный остаток"):
        ledger.read(path)


def test_нет_колонки_товара_объясняется_найденным(tmp_path):
    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 3, 7))],
                 labels=["Код товара", "Описание", "Ед. изм."])

    with pytest.raises(parse.ParseError) as error:
        ledger.read(path)

    assert "«Описание»" in str(error.value), "сообщение показывает, что нашлось"


# --- несколько файлов ----------------------------------------------------------

def test_файлы_магазинов_сводятся_в_одну_ведомость(tmp_path):
    first = _single_book(tmp_path, "Реми.xlsx", [("1", "Мыло", 10, 0, 4, 6),
                                                 ("2", "Крем", 5, 0, 1, 4)])
    second = _single_book(tmp_path, "Сити Молл.xlsx",
                          [("1", "Мыло", 20, 0, 5, 15)])

    led = ledger.read_many([first, second])

    assert [store.title for store in led.stores] == ["Реми", "Сити Молл"]
    assert [item.name for item in led.items] == ["Мыло", "Крем"]
    assert led.items[0].move("Сити Молл").outgoing == 5
    assert led.total().outgoing == 10
    assert led.sources == ["Реми.xlsx", "Сити Молл.xlsx"]


def test_файлы_разного_вида_сводятся_вместе(tmp_path):
    wide = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 4, 6))],
                 total=[])
    flat = _flat_book(tmp_path, [("Сити Молл", "1", "Мыло", "шт", 20, 0, 5, 15)])

    led = ledger.read_many([wide, flat])

    assert [store.title for store in led.stores] == ["Реми", "Сити Молл"]
    assert len(led.items) == 1
    assert led.balanced, "итоги обоих файлов сложились"


def test_один_склад_в_двух_файлах_не_складывается(tmp_path):
    """Скорее всего, файл выбран дважды: удвоить расход хуже, чем развести."""
    first = _single_book(tmp_path, "Реми.xlsx", [("1", "Мыло", 10, 0, 4, 6)],
                         selection='Склад Равно "Реми"')
    second = _single_book(tmp_path, "Реми копия.xlsx",
                          [("1", "Мыло", 10, 0, 4, 6)],
                          selection='Склад Равно "Реми"')

    led = ledger.read_many([first, second])

    assert [store.title for store in led.stores] == ["Реми", "Реми (2)"]
    assert ledger.build(led).doubled, "о повторе нужно предупредить"


def test_товар_с_кодом_и_без_кода_сводится_в_один(tmp_path):
    with_code = _book(tmp_path, ["Реми"],
                      [("1", "Мыло", "шт", _line(10, 0, 4, 6), "00-001")],
                      labels=["Артикул", "Номенклатура", "Ед. изм.", "Код"],
                      filename="с кодом.xlsx")
    without = _book(tmp_path, ["Сити Молл"],
                    [("1", "Мыло", "шт", _line(20, 0, 5, 15))],
                    filename="без кода.xlsx")

    led = ledger.read_many([with_code, without])

    assert len(led.items) == 1
    assert led.items[0].code == "00-001"


def test_несошедшийся_файл_назван_в_предупреждении(tmp_path):
    good = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(10, 0, 4, 6))],
                 total=[], filename="хороший.xlsx")
    bad = _book(tmp_path, ["Сити Молл"],
                [("1", "Мыло", "шт", _line(10, 0, 4, 6))],
                total=[], filename="плохой.xlsx")
    from openpyxl import load_workbook
    book = load_workbook(bad)
    book.active.cell(10, 10, 99)
    book.save(bad)

    led = ledger.read_many([good, bad])

    assert not led.balanced
    assert any(note.startswith("плохой.xlsx") for note in led.warnings)
    assert not any(note.startswith("хороший.xlsx") for note in led.warnings)


def test_ошибка_в_одном_из_файлов_называет_файл(tmp_path):
    good = _single_book(tmp_path, "Реми.xlsx", [("1", "Мыло", 10, 0, 4, 6)])
    book = Workbook()
    book.active["A1"] = "просто таблица"
    bad = str(tmp_path / "чужое.xlsx")
    book.save(bad)

    with pytest.raises(parse.ParseError, match="^чужое.xlsx: "):
        ledger.read_many([good, bad])


def test_выгрузка_без_долей_когда_запас_неизвестен(tmp_path):
    from openpyxl import load_workbook

    path = _book(tmp_path, ["Реми"], [("1", "Мыло", "шт", _line(3, 0))],
                 fields=("Расход", "Конечный остаток"))
    report = ledger.build(ledger.read(path))
    target = ledger.save(report, str(tmp_path / "разбор.xlsx"))

    sheet = load_workbook(target)["Главное"]
    values = {sheet.cell(row, 1).value: sheet.cell(row, 2).value
              for row in range(1, sheet.max_row + 1)}
    assert values["Израсходовано от запаса"] is None
    assert values["Вид отчёта"] == parse.LAYOUT_WIDE
