"""Торг-2: код маркировки, излишек и четвёртая страница акта.

Брак и излишек маркируемого товара заводятся с кодом: из него берутся штрих-код
и срок годности, он обязателен, а в акте ТОРГ-2 появляется лист «стр4» с кодами.
Сканер заканчивает код нажатием Enter — окно от этого закрываться не должно.
"""
from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402
from PySide6.QtWidgets import QApplication, QDialog  # noqa: E402

from app.core.claims import (  # noqa: E402
    KIND_TITLES,
    Claim,
    ClaimHint,
    Manager,
    TorgAct,
    TorgItem,
    TorgSettings,
    mark,
    remote,
    torg,
)
from app.ui.widgets import claim_dialogs  # noqa: E402
from app.ui.widgets.claim_dialogs import ClaimDialog  # noqa: E402

# Код с образца склада: GTIN, серийный номер и криптохвост (91, 92), срока нет.
CODE = '01062965463513152158Py"bZOZVsoE91EE1092uj9M9scYxegE/Ws8rA3dhXEvbUBC1rg9YmzG7qIKo4E='
# Тот же товар, но со сроком годности в самом коде (AI 17 — 31.12.2030).
GS = chr(29)
CODE_WITH_EXPIRY = "010629654635131521ABCDEFGH1234" + GS + "17301231" + "91EE10" + "92" + "A" * 44

SETTINGS = TorgSettings(
    organization="ИП Тестов", director_title="ИП Тестов Т. Т.", director_name="Тестов Т. Т.",
    place="г. Тест, ул. Склад 2", chair_name="Иванова И. И.",
    members=[("товаровед", "Петров П."), ("товаровед", "Сидоров С.")])


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


def dialog(app, claim: Claim | None = None, **kwargs) -> ClaimDialog:
    return ClaimDialog(claim, [Manager(7, "Анна")], [], **kwargs)


# --- чтение кода (ядро) -------------------------------------------------------

class TestRead:
    def test_штрих_код_берётся_из_кода(self) -> None:
        read = mark.read(CODE)

        assert read.valid and read.code == CODE
        # У штучного товара GTIN-14 с нулём впереди — привычный EAN-13.
        assert read.gtin == "06296546351315" and read.barcode == "6296546351315"

    def test_срок_из_кода_без_обращения_к_сети(self) -> None:
        read = mark.read(CODE_WITH_EXPIRY)

        assert read.valid
        assert read.expires == date(2030, 12, 31)
        assert not read.needs_lookup and not read.estimated

    def test_если_срока_нет_его_можно_спросить_у_системы(self) -> None:
        read = mark.read(CODE)

        assert read.expires is None and read.needs_lookup

    def test_мусор_не_код(self) -> None:
        read = mark.read("не код вовсе")

        assert not read.valid and read.problem

    def test_пусто_ничего_не_значит(self) -> None:
        read = mark.read("  ")

        assert read.code == "" and not read.valid and read.problems == []

    def test_русская_раскладка_сканера_возвращается_в_латиницу(self) -> None:
        plain = "010629654635131521ABCDEFGH1234" + GS + "91EE10" + "92" + "A" * 44
        # Те же клавиши при русской раскладке: A→Ф, B→И, C→С, D→В, E→У, F→А, G→П, H→Р.
        typed = plain.replace("ABCDEFGH1234", "ФИСВУАПР1234").replace("EE10", "УУ10")

        read = mark.read(typed)

        assert read.valid and read.code == plain

    def test_ответ_системы_дополняет_срок(self) -> None:
        read = mark.read(CODE)
        answer = SimpleNamespace(
            raw={"cisInfo": {"expirationDate": "2029-01-31T00:00:00.000Z",
                             "productName": "Парфюмерная вода"}},
            product_name="Парфюмерная вода")

        done = mark.complete(read, answer)

        assert done.expires == date(2029, 1, 31)
        assert done.name == "Парфюмерная вода" and not done.needs_lookup

    def test_у_парфюмерии_срок_считается_от_даты_производства(self) -> None:
        read = mark.read(CODE)
        answer = SimpleNamespace(
            raw={"cisInfo": {"productionDate": "2026-01-20T00:00:00.000Z"}},
            product_name="")

        done = mark.complete(read, answer, years=3)

        assert done.expires == date(2029, 1, 19)
        assert done.estimated and done.shelf_years == 3

    def test_код_обязателен_у_брака_и_излишка_но_не_у_недопоставки(self) -> None:
        assert mark.needs_code("defect", True)
        assert mark.needs_code("surplus", True)
        assert not mark.needs_code("shortage", True)
        assert not mark.needs_code("defect", False)


# --- окно заявки --------------------------------------------------------------

class TestDialog:
    def test_излишек_среди_видов_заявки(self, app) -> None:
        assert KIND_TITLES["surplus"] == "Излишек"
        d = dialog(app)
        kinds = [d.kind.itemData(i) for i in range(d.kind.count())]

        assert kinds == ["shortage", "defect", "surplus"]

    def test_поле_кода_появляется_с_галочкой(self, app) -> None:
        d = dialog(app)
        d.show()
        assert d.mark_code.isHidden()

        d.marked.setChecked(True)
        assert not d.mark_code.isHidden()

        d.marked.setChecked(False)
        assert d.mark_code.isHidden()

    @pytest.mark.parametrize("kind", ["defect", "surplus"])
    def test_у_брака_и_излишка_маркируемого_товара_без_кода_не_завести(self, app, kind) -> None:
        d = dialog(app)
        d.kind.setCurrentIndex(d.kind.findData(kind))
        d.brand.setText("Rink")
        d.marked.setChecked(True)

        d._accept()

        assert d.result() != QDialog.DialogCode.Accepted
        assert "отсканируйте код маркировки" in d.error.text()

    def test_у_недопоставки_код_не_обязателен(self, app) -> None:
        d = dialog(app)
        d.kind.setCurrentIndex(d.kind.findData("shortage"))
        d.brand.setText("Rink")
        d.marked.setChecked(True)

        d._accept()

        assert d.result() == QDialog.DialogCode.Accepted

    def test_без_галочки_код_не_нужен(self, app) -> None:
        d = dialog(app)
        d.kind.setCurrentIndex(d.kind.findData("defect"))
        d.brand.setText("Rink")

        d._accept()

        assert d.result() == QDialog.DialogCode.Accepted
        assert d.claims_fields()[0]["mark_code"] == ""

    def test_непонятный_код_не_сохраняется(self, app) -> None:
        d = dialog(app)
        d.kind.setCurrentIndex(d.kind.findData("defect"))
        d.brand.setText("Rink")
        d.marked.setChecked(True)
        d.mark_code.setText("абракадабра")

        d._accept()

        assert d.result() != QDialog.DialogCode.Accepted
        assert "не разобран" in d.error.text()

    def test_скан_подставляет_штрих_код_и_срок(self, app, monkeypatch) -> None:
        monkeypatch.setattr(claim_dialogs, "run_task", lambda *a, **k: None)
        d = dialog(app)
        d.marked.setChecked(True)

        d.mark_code.setText(CODE_WITH_EXPIRY)
        d._code_scanned()

        assert d.barcode.text() == "6296546351315"
        assert d.expiry.value() == date(2030, 12, 31)
        assert "взят из кода" in d.mark_note.text()

    def test_введённый_руками_срок_скан_не_затирает(self, app, monkeypatch) -> None:
        monkeypatch.setattr(claim_dialogs, "run_task", lambda *a, **k: None)
        d = dialog(app)
        d.marked.setChecked(True)
        d.expiry.set_value(date(2028, 5, 1))

        d.mark_code.setText(CODE_WITH_EXPIRY)
        d._code_scanned()

        assert d.expiry.value() == date(2028, 5, 1)

    def test_срок_от_прошлого_скана_заменяется_новым(self, app, monkeypatch) -> None:
        monkeypatch.setattr(claim_dialogs, "run_task", lambda *a, **k: None)
        d = dialog(app)
        d.marked.setChecked(True)
        d.mark_code.setText(CODE_WITH_EXPIRY)
        d._code_scanned()
        other = CODE_WITH_EXPIRY.replace("17301231", "17291130").replace("ABCDEFGH1234", "ZZZZZZZZ9999")

        d.mark_code.setText(other)
        d._code_scanned()

        assert d.expiry.value() == date(2029, 11, 30)

    def test_без_срока_в_коде_без_входа_подсказка_про_вход(self, app, monkeypatch) -> None:
        from app.core.marking import service

        monkeypatch.setattr(service, "signed_in", lambda: False)
        d = dialog(app)
        d.marked.setChecked(True)

        d.mark_code.setText(CODE)
        d._code_scanned()

        assert "войдите по сертификату" in d.mark_note.text()
        assert d.expiry.text() == ""

    def test_без_срока_в_коде_спрашивается_система_и_срок_подставляется(self, app, monkeypatch) -> None:
        from app.core.marking import service

        monkeypatch.setattr(service, "signed_in", lambda: True)
        answer = SimpleNamespace(
            raw={"cisInfo": {"expirationDate": "2029-01-31T00:00:00.000Z"}},
            product_name="Вода парфюмерная")
        monkeypatch.setattr(claim_dialogs, "_ask_gis", lambda code: answer)
        def run(fn, *args, on_result=None, **kwargs):
            # Подсказка по штрих-коду — отдельный запрос, здесь он не нужен.
            if fn is not remote.hint:
                on_result(fn(*args))

        monkeypatch.setattr(claim_dialogs, "run_task", run)
        d = dialog(app)
        d.marked.setChecked(True)

        d.mark_code.setText(CODE)
        d._code_scanned()

        assert d.expiry.value() == date(2029, 1, 31)
        assert d.product.text() == "Вода парфюмерная"

    def test_двойной_сигнал_сканера_спрашивает_систему_один_раз(self, app, monkeypatch) -> None:
        """Enter от сканера и уход фокуса приходят оба."""
        from app.core.marking import service

        monkeypatch.setattr(service, "signed_in", lambda: True)
        asked: list[str] = []
        monkeypatch.setattr(claim_dialogs, "run_task",
                            lambda fn, code, *rest, **k: asked.append(code)
                            if fn is claim_dialogs._ask_gis else None)
        d = dialog(app)
        d.marked.setChecked(True)
        d.mark_code.setText(CODE)

        d._code_scanned()
        d._code_scanned()

        assert asked == [CODE]

    def test_подсказка_по_штрих_коду_ставит_галочку(self, app) -> None:
        d = dialog(app)

        d._apply_hint(ClaimHint(brand="Rink", marked=True))

        assert d.marked.isChecked()

    def test_код_попадает_в_позицию_и_поля_очищаются(self, app, monkeypatch) -> None:
        monkeypatch.setattr(claim_dialogs, "run_task", lambda *a, **k: None)
        d = dialog(app)
        d.kind.setCurrentIndex(d.kind.findData("defect"))
        d.brand.setText("Rink")
        d.marked.setChecked(True)
        d.mark_code.setText(CODE_WITH_EXPIRY)
        d._code_scanned()

        d._add_position()

        assert d.claims_fields()[0]["mark_code"] == CODE_WITH_EXPIRY
        assert d.mark_code.text() == "" and not d.marked.isChecked()

    def test_правка_заведённой_заявки_видит_код(self, app) -> None:
        claim = remote._claim({
            "id": 5, "kind": "defect", "brand": "Rink", "product": "Крем", "spec": "",
            "barcode": "6296546351315", "quantity": "1", "mark_code": CODE,
            "comment": "", "status": "new", "applicant_id": 3, "applicant": "С",
            "can_edit": True})
        d = dialog(app, claim)

        assert d.marked.isChecked() and d.mark_code.text() == CODE
        assert d.changes() == {}

        d.mark_code.setText(CODE_WITH_EXPIRY)
        assert d.changes() == {"mark_code": CODE_WITH_EXPIRY}

    def test_enter_от_сканера_окно_не_закрывает(self, app) -> None:
        """Сканер кончает код нажатием Enter: окно было «кнопкой по умолчанию»
        и закрывалось на середине ввода."""
        d = dialog(app)
        d.show()
        d.brand.setText("Rink")

        for field in (d.barcode, d.brand, d.product, d.spec):
            field.setFocus()
            QTest.keyClick(field, Qt.Key.Key_Return)
        d.marked.setChecked(True)
        d.mark_code.setFocus()
        d.mark_code.setText(CODE)
        QTest.keyClick(d.mark_code, Qt.Key.Key_Return)
        QTest.keyClick(d.mark_code, Qt.Key.Key_Enter)

        assert d.result() != QDialog.DialogCode.Accepted
        assert d.isVisible()
        # И случайных позиций Enter тоже не добавляет.
        assert d._positions == []

    def test_ни_одна_кнопка_окна_не_кнопка_по_умолчанию(self, app) -> None:
        from PySide6.QtWidgets import QPushButton

        d = dialog(app)

        assert not any(b.autoDefault() or b.isDefault() for b in d.findChildren(QPushButton))


# --- акт ТОРГ-2 ----------------------------------------------------------------

def titem(number: int = 1, **over) -> TorgItem:
    data = dict(claim_id=number, name=f"Товар {number}", barcode=f"46000000000{number:02d}",
                price=100.0, qty_doc=10)
    data.update(over)
    return TorgItem(**data)


def tact(*items: TorgItem) -> TorgAct:
    return TorgAct(
        number=70, act_date=date(2026, 10, 21), receipt_date=date(2026, 10, 20),
        supplier="ООО Тест", contract_number="1", contract_date=date(2026, 1, 1),
        invoice_number="55", invoice_date=date(2026, 10, 1), description="x", conclusion="y",
        items=list(items))


class TestSurplus:
    def test_излишек_входит_в_поступившее_и_в_итоги(self) -> None:
        extra = titem(qty_surplus=3)

        assert extra.qty_fact == 13
        assert extra.sum_surplus == 300
        totals = tact(extra, titem(2, qty_short=2)).totals()
        assert totals["qty_surplus"] == 3 and totals["sum_surplus"] == 300
        assert totals["qty_fact"] == 13 + 8

    def test_излишек_самостоятельная_причина_для_акта(self) -> None:
        assert not tact(titem(qty_surplus=2)).problems()
        problems = tact(titem()).problems()
        assert any("недостачу, брак или излишек" in p for p in problems)

    def test_товара_которого_нет_в_документах_можно_оформить_излишком(self) -> None:
        assert not tact(titem(qty_doc=0, qty_surplus=2)).problems()
        assert any("по документам" in p for p in tact(titem(qty_doc=0)).problems())

    def test_заявка_излишка_становится_излишком_в_акте(self) -> None:
        claim = Claim(id=4, kind="surplus", brand="Rink", product="Крем", barcode="123",
                      quantity=2, mark_code=CODE)

        built = torg.act_from_claims([claim], 70, date(2026, 10, 21))

        item = built.items[0]
        assert item.qty_surplus == 2 and item.qty_defect == 0 and item.qty_short == 0
        assert item.mark_code == CODE

    def test_излишек_и_код_переживают_сохранение_акта(self) -> None:
        item = titem(qty_surplus=2, mark_code=CODE)

        again = TorgItem.from_dict(item.to_dict())

        assert again.qty_surplus == 2 and again.mark_code == CODE

    def test_текст_по_умолчанию_для_излишка(self) -> None:
        description, conclusion = tact(titem(qty_surplus=2)).suggested_text()

        assert "больше" in description and "излишне поставленную" in conclusion

    def test_излишек_попадает_в_ячейки_бланка(self) -> None:
        book = torg.build_workbook(tact(titem(qty_surplus=3)), SETTINGS)
        second = book["стр2"]

        assert second["GE40"].value == 3 and second["GM40"].value == 300
        assert second["GE41"].value == 3  # итого
        assert second["DV40"].value == 13  # поступило


class TestCodesPage:
    def test_у_акта_без_кодов_четвёртой_страницы_нет(self) -> None:
        book = torg.build_workbook(tact(titem(qty_defect=1)), SETTINGS)

        assert book.sheetnames == ["стр1", "стр2", "стр3"]

    def test_коды_маркировки_на_четвёртой_странице(self) -> None:
        book = torg.build_workbook(tact(
            titem(1, qty_defect=1, mark_code=CODE),
            titem(2, qty_short=2),
            titem(3, qty_surplus=1, mark_code=CODE_WITH_EXPIRY)), SETTINGS)

        sheet = book[torg.CODES_SHEET]
        rows = [[c.value for c in row] for row in sheet.iter_rows()]
        assert rows[0] == ["Номер строки", "Код номенклатуры", "Штрих-код",
                           "Код маркировки", "Брак", "Недостача", "Излишек"]
        # Строка 2 акта — недостача, кода у неё нет: страница её пропускает, а
        # номер в первой колонке остаётся номером позиции в таблице акта.
        assert [r[0] for r in rows[1:]] == ["1", "3"]
        assert rows[1][3] == CODE and rows[1][4] == 1 and rows[1][6] is None
        assert rows[2][6] == 1 and rows[2][4] is None

    def test_разделитель_групп_в_ячейку_не_пишется(self) -> None:
        """GS внутри кода Excel не принимает: страница не должна упасть."""
        book = torg.build_workbook(tact(titem(qty_defect=1, mark_code=CODE_WITH_EXPIRY)),
                                   SETTINGS)

        written = book[torg.CODES_SHEET]["D2"].value

        assert GS not in written and written == CODE_WITH_EXPIRY.replace(GS, "")

    def test_акт_с_кодами_сохраняется_и_открывается(self, tmp_path) -> None:
        import openpyxl

        path = tmp_path / "act.xlsx"
        torg.save(tact(titem(qty_defect=1, mark_code=CODE)), SETTINGS, path)

        again = openpyxl.load_workbook(path)

        assert again.sheetnames[-1] == "стр4"
        assert again["стр4"]["D2"].value == CODE
